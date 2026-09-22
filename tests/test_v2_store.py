import os
import json
import re
import tempfile
import unittest
import zipfile
from datetime import datetime
from unittest.mock import patch

from openpyxl import Workbook

from v2_store import V2Store, classify_buyer_message, extract_store_query, normalize_text


class HandoffDate(datetime):
    @classmethod
    def now(cls, tz=None):
        return cls(2026, 9, 6, 12, 0, tzinfo=tz)


class V2StoreTests(unittest.TestCase):
    def setUp(self):
        # Existing "today" scenarios use the Sunday handoff baseline.
        clock = patch("v2_store.datetime", HandoffDate)
        clock.start()
        self.addCleanup(clock.stop)
        self.temp = tempfile.TemporaryDirectory()
        self.store = V2Store(os.path.join(self.temp.name, "v2.db"))
        self.store.save_v2_product(
            "10001",
            "小江溪125元代金券",
            "售价89元。仅限当前商品，不支持擅自退款。",
        )

    def tearDown(self):
        self.temp.cleanup()

    def create_store_sheet(self):
        path = os.path.join(self.temp.name, "stores.xlsx")
        book = Workbook()
        sheet = book.active
        sheet.append(["店名", "分店名", "省", "市", "区/县", "地址", "电话"])
        sheet.append(["小江溪·江西菜", "华熙五棵松店", "北京", "北京", "海淀", "复兴路69号", "13426310802"])
        sheet.append(["小江溪·江西菜", "西单大悦城店", "北京", "北京", "西城", "西单北大街", "13621240162"])
        sheet.append(["小江溪·江西菜", "东城万象店", "北京", "北京", "东城", "东直门", "10086"])
        sheet.append(["小江溪·江西菜", "西城万象店", "北京", "北京", "西城", "西单", "10010"])
        # duplicate row must not inflate the result
        sheet.append(["小江溪·江西菜", "华熙五棵松店", "北京", "北京", "海淀", "复兴路69号", "13426310802"])
        book.save(path)
        return path

    def create_irregular_store_sheet(self):
        path = os.path.join(self.temp.name, "irregular-stores.xlsx")
        book = Workbook()
        cover = book.active
        cover.title = "说明"
        cover.append(["商户导出文件"])
        sheet = book.create_sheet("门店明细")
        sheet.append(["适用门店列表"])
        sheet.append([])
        sheet.append(["门店名称（可模糊）", "门店地址", "联系电话 / 手机", "营业时间段"])
        sheet.append(["国贸中心店", "建国路88号", "010-1234", "10:00-22:00"])
        book.save(path)

        # Reproduce merchant exports that incorrectly declare the used range as A1.
        rewritten = path + ".tmp"
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(rewritten, "w") as target:
            for entry in source.infolist():
                data = source.read(entry.filename)
                if entry.filename == "xl/worksheets/sheet2.xml":
                    data = re.sub(
                        rb'<dimension ref="[^"]+"(?:></dimension>|/>)',
                        b'<dimension ref="A1"></dimension>',
                        data,
                        count=1,
                    )
                target.writestr(entry, data)
        os.replace(rewritten, path)
        return path

    def create_multi_region_store_sheet(self):
        path = os.path.join(self.temp.name, "multi-region-stores.xlsx")
        book = Workbook()
        sheet = book.active
        sheet.append(["店名", "分店名", "省", "市", "区/县", "地址"])
        sheet.append(["测试品牌", "龙岗万科里店", "广东", "深圳", "龙岗区", "龙岗大道"])
        sheet.append(["测试品牌", "西丽益田假日里店", "广东", "深圳", "南山区", "留仙大道"])
        sheet.append(["测试品牌", "广州万达广场店", "广东", "广州", "天河区", "天河路"])
        sheet.append(["测试品牌", "南京万达广场店", "江苏", "南京", "江宁区", "双龙大道"])
        sheet.append(["测试品牌", "寿光万达广场店", "山东", "寿光", "寿光", "圣城街"])
        book.save(path)
        return path

    def test_raw_text_remains_highest_priority(self):
        self.store.save_ai_summary("10001", "AI摘要", {"summary": "AI摘要", "time_rules": []})
        merged = self.store.build_merged_knowledge("10001", "闲鱼页面资料")
        self.assertLess(merged.index("用户原始资料"), merged.index("AI归纳摘要"))
        self.assertLess(merged.index("AI归纳摘要"), merged.index("闲鱼页面资料"))
        self.assertIn("售价89元", merged)
        self.assertIn("禁止推荐任何其他商品", merged)

    def test_ai_draft_is_inert_until_adopted_and_promotes_matching_profiles(self):
        original = self.store.get_v2_product("10001")
        draft = {
            "sku_profiles": [{
                "sku_key": "sku-100", "sku_name": "100元代金券",
                "sale_price": "69", "face_value": "100", "meal_period": "午餐",
            }],
            "time_rules": [{
                "label": "午餐", "day_type": "any", "start_time": "11:00",
                "end_time": "14:00", "allowed": True, "reply": "100元代金券午餐可用",
            }],
        }
        pending = self.store.save_ai_draft("10001", "AI待采纳草稿", draft)
        self.assertEqual(original["raw_text"], pending["raw_text"])
        self.assertEqual(original["structured"], pending["structured"])
        self.assertEqual([], self.store.list_time_rules("10001"))
        self.assertEqual("AI待采纳草稿", pending["ai_draft_summary"])

        adopted = self.store.adopt_ai_draft("10001", "AI待采纳草稿")
        self.assertEqual("AI待采纳草稿", adopted["raw_text"])
        self.assertEqual("100元代金券", adopted["structured"]["sku_profiles"][0]["sku_name"])
        self.assertEqual(1, len(self.store.list_time_rules("10001")))

    def test_editing_visible_ai_draft_discards_hidden_structure(self):
        self.store.save_ai_draft("10001", "含隐藏旧规则的草稿", {
            "facts": {"不可用日期": "国庆不可用"},
            "time_rules": [{
                "label": "旧规则", "day_type": "weekday", "start_time": "00:00",
                "end_time": "23:59", "allowed": False, "reply": "不可用",
            }],
        })
        adopted = self.store.adopt_ai_draft("10001", "人工删除旧规则后的正文")
        self.assertEqual({}, adopted["structured"])
        self.assertEqual([], self.store.list_time_rules("10001"))

    def test_ai_draft_cannot_fall_back_to_stale_effective_summary(self):
        self.store.save_ai_summary(
            "10001", "页面旧摘要289元", {"facts": {"价格": "289元"}},
        )
        with self.assertRaisesRegex(ValueError, "没有已生成的AI草稿"):
            self.store.adopt_ai_draft("10001", "页面旧摘要289元")

    def test_each_sku_profile_inherits_common_rules_but_keeps_own_difference(self):
        text = self.store._sku_profile_summary({
            "common_rules": {"stack_rule": "同面额可无限叠加", "meal_period": "全天"},
            "sku_profiles": [
                {"sku_key": "a", "sku_name": "工作日午餐", "sale_price": "135", "meal_period": "午餐"},
                {"sku_key": "b", "sku_name": "周末晚餐", "sale_price": "157", "meal_period": "晚餐"},
            ],
        })
        self.assertEqual(2, text.count("叠加规则：同面额可无限叠加"))
        self.assertIn("〔工作日午餐〕\n售价135元\n餐段：午餐", text)
        self.assertIn("〔周末晚餐〕\n售价157元\n餐段：晚餐", text)
        self.assertNotIn("餐段：全天", text)

    def test_question_relevant_knowledge_is_trimmed_before_model_call(self):
        self.store.save_v2_product(
            "10001",
            "测试代金券",
            "【商品规格】\n100元代金券售价67.9元。\n\n"
            "【适用门店】\n武汉万象城店、深圳壹方城店。\n\n"
            "【退款政策】\n非卡券质量问题退款扣5%手续费。",
        )
        self.store.save_ai_summary(
            "10001",
            "【商品规格】100元代金券售价67.9元。\n\n"
            "【使用规则】仅限堂食。\n\n【退款政策】72小时内处理。",
            {"summary": "测试"},
        )
        merged = self.store.build_merged_knowledge(
            "10001", "【页面说明】全国门店很多，页面原价100元。", "100元多少钱"
        )
        self.assertIn("100元代金券售价67.9元", merged)
        self.assertNotIn("武汉万象城店", merged)
        self.assertNotIn("非卡券质量问题退款扣5%", merged)
        self.assertIn("知识裁剪说明", merged)

    def test_refund_question_keeps_refund_policy_but_drops_store_blocks(self):
        self.store.save_v2_product(
            "10001",
            "测试代金券",
            "【适用门店】\n深圳壹方城店。\n\n"
            "【退款政策】\n卡券质量问题可申请仅退款。",
        )
        merged = self.store.build_merged_knowledge("10001", "页面资料", "买了不能用怎么退款")
        self.assertIn("卡券质量问题可申请仅退款", merged)
        self.assertIn("当前商品适用的发货与退款政策", merged)
        self.assertNotIn("深圳壹方城店", merged)

    def test_coupon_type_defaults_to_meituan(self):
        self.assertEqual("meituan", self.store.get_v2_product("10001")["coupon_type"])

    def test_purchase_order_business_questions_are_all_silent(self):
        self.store.save_v2_product("10001", "代买服务", "按实际需求人工报价", coupon_type="purchase_order")
        for message in ("多少钱", "给个报价", "300怎么拍", "可以便宜点吗"):
            result = self.store.resolve_deterministic("10001", message)
            self.assertEqual("purchase_order_other", result["kind"])
            self.assertEqual("silent", result["decision"])
            self.assertEqual("", result["reply"])

    def test_purchase_order_store_fallback_never_answers_business_questions(self):
        self.store.save_v2_product("10001", "代买服务", "付款后发送领取信息", coupon_type="purchase_order")
        greeting = self.store.resolve_deterministic("10001", "你好")
        self.assertEqual("purchase_order_other", greeting["kind"])
        self.assertEqual("silent", greeting["decision"])
        usage = self.store.resolve_deterministic("10001", "付款后怎么领取和核销")
        self.assertEqual("purchase_order_other", usage["kind"])
        self.assertEqual("silent", usage["decision"])

    def test_excel_import_fuzzy_match_and_unavailable(self):
        result = self.store.import_store_list(self.create_store_sheet(), "北京门店", ["10001"])
        self.assertEqual(result["store_count"], 4)
        match = self.store.search_store("10001", "五棵松")
        self.assertEqual(match["status"], "available")
        self.assertEqual(match["matches"][0]["branch"], "华熙五棵松店")
        missing = self.store.search_store("10001", "不存在的门店")
        self.assertEqual(missing["status"], "unavailable")

    def test_sku_store_rules_default_to_inherit_and_custom_is_isolated(self):
        self.store.save_ai_summary("10001", "两种套餐", {
            "sale_options": [
                {"name": "单人套餐", "sale_price": "88", "people_counts": [1]},
                {"name": "双人套餐", "sale_price": "168", "people_counts": [2]},
            ]
        })
        default_list = self.store.import_store_text(
            "【湖北省】\n【武汉】武昌万象城店", "默认门店", ["10001"]
        )
        special_list = self.store.import_store_text(
            "【湖北省】\n【武汉】武汉天地店", "双人专属", []
        )
        skus = self.store.get_v2_product("10001")["skus"]
        self.assertEqual(2, len(skus))
        self.assertTrue(all(sku["mode"] == "inherit" for sku in skus))
        single = next(sku for sku in skus if "单人" in sku["sku_name"])
        double = next(sku for sku in skus if "双人" in sku["sku_name"])
        self.store.set_sku_store_rule(
            "10001", double["sku_key"], double["sku_name"], "custom", [special_list["id"]]
        )
        self.assertEqual("available", self.store.search_store(
            "10001", "万象城", sku_key=single["sku_key"]
        )["status"])
        self.assertEqual("unavailable", self.store.search_store(
            "10001", "万象城", sku_key=double["sku_key"]
        )["status"])
        self.assertEqual("available", self.store.search_store(
            "10001", "武汉天地", sku_key=double["sku_key"]
        )["status"])
        self.assertEqual([default_list["id"]], self.store.effective_store_list_ids(
            "10001", single["sku_key"]
        )[0])

    def test_structured_vouchers_without_prices_remain_available_for_store_binding(self):
        raw_text = (
            "美团200元代金券最多叠加2张。\n"
            "抖音200元代金券最多叠加4张。"
        )
        self.store.save_v2_product("10001", "双平台代金券", raw_text)
        self.store.save_ai_summary("10001", "双平台规格", {
            "products": [
                {
                    "name": "美团200元代金券", "option_type": "代金券",
                    "face_value": "200", "sale_price": "", "max_stack": 2,
                },
                {
                    "name": "抖音200元代金券", "option_type": "代金券",
                    "face_value": "200", "sale_price": "", "max_stack": 4,
                },
            ]
        })

        skus = self.store.get_v2_product("10001")["skus"]

        self.assertEqual(["美团200元代金券", "抖音200元代金券"], [
            sku["sku_name"] for sku in skus
        ])
        self.assertTrue(all(sku["price_pending"] for sku in skus))
        self.assertTrue(all(sku["sellable"] for sku in skus))
        self.assertEqual(2, len({sku["sku_key"] for sku in skus}))
        matched = self.store.match_message_skus(
            "10001", "美团200能用吗", self.store.get_v2_product("10001"),
        )
        self.assertEqual(["美团200元代金券"], [sku["sku_name"] for sku in matched])

        updated = self.store.set_sku_store_rule(
            "10001", skus[0]["sku_key"], skus[0]["sku_name"], "inherit", [], "119.8",
        )
        self.assertEqual("119.8", updated["sale_price"])
        self.assertEqual("119.8", updated["sale_price_override"])
        self.assertFalse(updated["price_pending"])

    def test_store_query_reverse_recommends_only_supported_sku_with_real_price(self):
        self.store.save_ai_summary("10001", "两种套餐", {
            "sale_options": [
                {"name": "单人套餐", "sale_price": "88", "people_counts": [1]},
                {"name": "双人套餐", "sale_price": "168", "people_counts": [2]},
            ]
        })
        first = self.store.import_store_text("【湖北省】\n【武汉】光谷店", "单人门店", [])
        second = self.store.import_store_text("【湖北省】\n【武汉】天地店", "双人门店", [])
        skus = self.store.get_v2_product("10001")["skus"]
        single = next(sku for sku in skus if "单人" in sku["sku_name"])
        double = next(sku for sku in skus if "双人" in sku["sku_name"])
        self.store.set_sku_store_rule("10001", single["sku_key"], single["sku_name"], "custom", [first["id"]])
        self.store.set_sku_store_rule("10001", double["sku_key"], double["sku_name"], "custom", [second["id"]])
        result = self.store.resolve_deterministic("10001", "武汉天地店买哪个")
        self.assertEqual("stores_sku_recommendation", result["kind"])
        self.assertIn("双人套餐", result["reply"])
        self.assertIn("售价168元", result["reply"])
        self.assertNotIn("单人套餐", result["reply"])

    def test_store_query_strips_grounded_food_category_and_price_followup_uses_store_skus(self):
        self.store.save_v2_product(
            "10001", "蚂蚁洞烤肉代金券",
            "美团200元代金券：售价119.8元\n抖音300元代金券：售价205元",
        )
        one = self.store.import_store_text(
            "【广东省】\n【深圳】深圳壹方城店", "美团门店", [],
        )
        other = self.store.import_store_text(
            "【广东省】\n【深圳】深圳万象前海店", "抖音门店", [],
        )
        for sku in self.store.get_v2_product("10001")["skus"]:
            list_id = one["id"] if "美团" in sku["sku_name"] else other["id"]
            self.store.set_sku_store_rule(
                "10001", sku["sku_key"], sku["sku_name"], "custom", [list_id],
            )

        initial = self.store.resolve_deterministic("10001", "深圳壹方城烤肉")
        self.assertEqual("stores_sku_recommendation", initial["kind"])
        self.assertEqual("available", initial["store_status"])
        self.assertEqual(["深圳壹方城店"], [
            row["branch"] for row in initial["store_matches"]
        ])
        for typo in ("深圳壹方成烤肉", "深圳壹方烤肉", "深圳壹方城城烤肉", "深圳一方城烤肉"):
            with self.subTest(typo=typo):
                fuzzy = self.store.resolve_deterministic("10001", typo)
                self.assertEqual("stores_sku_recommendation", fuzzy["kind"])
                self.assertEqual("available", fuzzy["store_status"])
                self.assertEqual(["深圳壹方城店"], [
                    row["branch"] for row in fuzzy["store_matches"]
                ])
        context = {
            "query": initial["store_query"], "matches": initial["store_matches"],
            "store_sku_matrix": initial["store_sku_matrix"], "status": "available",
        }
        followup = self.store.resolve_deterministic(
            "10001", "什么价格呢", store_context=context,
        )
        self.assertEqual("stores_sku_recommendation", followup["kind"])
        self.assertIn("美团200元代金券（售价119.8元）", followup["reply"])
        self.assertNotIn("抖音300元代金券", followup["reply"])

    def test_store_reply_preserves_meaningful_marketplace_sku_names(self):
        meituan = {
            "sku_key": "meituan-200", "sku_name": "美团200（可叠加2张）",
            "face_value": "200", "sale_price": "119.8", "option_type": "voucher",
        }
        douyin = {
            "sku_key": "douyin-200", "sku_name": "抖音200（可叠加4张）",
            "face_value": "200", "sale_price": "138", "option_type": "voucher",
        }
        matrix = [{
            "store": {"branch": "天河店", "city": "广州"},
            "supported_skus": [meituan, douyin], "unknown_skus": [],
        }]
        reply = self.store._format_store_sku_matrix("天河店", matrix)
        self.assertIn("可用规格为", reply)
        self.assertIn("美团200（可叠加2张）（售价119.8元）", reply)
        self.assertIn("抖音200（可叠加4张）（售价138元）", reply)
        selected_reply = self.store._format_store_sku_matrix(
            "天河店", matrix, [meituan, douyin],
        )
        self.assertIn("美团200（可叠加2张）、抖音200（可叠加4张）可以使用", selected_reply)
        self.assertNotIn("200元、200元代金券", selected_reply)

    def test_store_reply_keeps_all_stores_and_skus_in_one_answer(self):
        first_sku = {
            "sku_key": "first", "sku_name": "美团100（可叠加2张）",
            "face_value": "100", "sale_price": "69", "option_type": "voucher",
        }
        second_sku = {
            "sku_key": "second", "sku_name": "抖音200",
            "face_value": "200", "sale_price": "129", "option_type": "voucher",
        }
        reply = self.store._format_store_sku_matrix("青岛", [
            {
                "store": {"branch": "青岛凯德MALL·新都心店", "city": "青岛"},
                "supported_skus": [first_sku], "unknown_skus": [],
            },
            {
                "store": {"branch": "青岛崂山万象汇店", "city": "青岛"},
                "supported_skus": [second_sku], "unknown_skus": [],
            },
        ])

        self.assertEqual(
            "根据“青岛”查询结果：\n\n"
            "【青岛凯德MALL·新都心店】：可用规格为"
            "美团100（可叠加2张）（售价69元）。\n\n"
            "【青岛崂山万象汇店】：可用规格为抖音200（售价129元）。\n\n"
            "请按对应门店支持的规格拍下。",
            reply,
        )

    def test_multi_sku_specific_store_uses_product_union_without_city_fallback(self):
        self.store.save_v2_product(
            "10001", "多规格代金券",
            "100元代金券：售价64元\n125元代金券：售价74元\n"
            "200元代金券：售价120元\n300元代金券：售价192元",
        )
        first = self.store.import_store_text(
            "【安徽省】\n【合肥】合肥之心城店", "100和300门店", [],
        )
        second = self.store.import_store_text(
            "【安徽省】\n【合肥】合肥银泰店", "125和200门店", [],
        )
        skus = self.store.get_v2_product("10001")["skus"]
        for sku in skus:
            face = sku["face_value"]
            list_id = first["id"] if face in {"100", "300"} else second["id"]
            self.store.set_sku_store_rule(
                "10001", sku["sku_key"], sku["sku_name"], "custom", [list_id],
            )

        result = self.store.resolve_deterministic("10001", "合肥之心城")
        self.assertEqual("stores_sku_recommendation", result["kind"])
        self.assertEqual("合肥之心城店", result["store_matches"][0]["branch"])
        self.assertIn("100元代金券（售价64元）", result["reply"])
        self.assertIn("300元代金券（售价192元）", result["reply"])
        self.assertNotIn("125元代金券", result["reply"])
        self.assertNotIn("200元代金券", result["reply"])
        self.assertNotIn("根据“合肥”", result["reply"])
        self.assertEqual(
            "根据“合肥之心城店”查询结果：\n\n"
            "【合肥之心城店】：可用规格为100元代金券（售价64元）、"
            "300元代金券（售价192元）。\n\n"
            "请按对应门店支持的规格拍下。",
            result["reply"],
        )
        segments = result["reply"].split("\n\n")
        self.assertEqual(3, len(segments))
        self.assertIn("合肥之心城店", segments[1])
        self.assertIn("100元代金券", segments[1])
        self.assertIn("300元代金券", segments[1])

        unavailable_sku = self.store.resolve_deterministic(
            "10001", "合肥之心城200可以用吗",
        )
        self.assertEqual("stores_sku_recommendation", unavailable_sku["kind"])
        self.assertEqual(
            "【合肥之心城店】：200元代金券不能使用；"
            "可用规格为100元代金券（售价64元）、300元代金券（售价192元）。",
            unavailable_sku["reply"],
        )
        available_sku = self.store.resolve_deterministic(
            "10001", "合肥之心城300可以用吗",
        )
        self.assertEqual(
            "【合肥之心城店】：300元代金券可以使用（售价192元）。",
            available_sku["reply"],
        )

    def test_store_followup_amount_lists_every_supported_exact_combination(self):
        self.store.save_v2_product(
            "10001", "双平台代金券",
            "美团200元代金券：售价119.8元，最多叠加2张\n"
            "抖音200元代金券：售价138元，最多叠加4张\n"
            "抖音500元代金券：售价328元，最多叠加4张",
        )
        self.store.save_ai_summary("10001", "双平台真实规格", {"products": [
            {"name": "美团200元代金券", "option_type": "代金券", "face_value": "200", "sale_price": "119.8", "max_stack": 2},
            {"name": "抖音200元代金券", "option_type": "代金券", "face_value": "200", "sale_price": "138", "max_stack": 4},
            {"name": "抖音500元代金券", "option_type": "代金券", "face_value": "500", "sale_price": "328", "max_stack": 4},
        ]})
        meituan_list = self.store.import_store_text(
            "【江西省】\n【南昌】南昌江西首店、南昌红谷滩店", "美团门店", [],
        )
        douyin_list = self.store.import_store_text(
            "【江西省】\n【南昌】南昌江西首店", "抖音门店", [],
        )
        for sku in self.store.get_v2_product("10001")["skus"]:
            list_id = meituan_list["id"] if "美团" in sku["sku_name"] else douyin_list["id"]
            self.store.set_sku_store_rule(
                "10001", sku["sku_key"], sku["sku_name"], "custom", [list_id],
            )
        first = self.store.resolve_deterministic("10001", "南昌江西首店")
        context = {
            "query": first["store_query"], "status": first["store_status"],
            "matches": first["store_matches"], "verified": True,
            "store_sku_matrix": first["store_sku_matrix"],
        }

        result = self.store.resolve_deterministic("10001", "400怎么买", store_context=context)

        self.assertEqual("consumption_plan", result["kind"])
        self.assertIn("可选组合方案", result["reply"])
        self.assertIn("美团200元代金券×2（共支付239.6元，可抵扣400元）", result["reply"])
        self.assertIn("抖音200元代金券×2（共支付276元，可抵扣400元）", result["reply"])
        self.assertNotIn("抖音500元代金券", result["reply"])

    def test_bare_amount_then_generic_price_uses_selected_sku(self):
        self.store.save_v2_product("10001", "抖音代金券", "抖音500元代金券：售价328元")
        self.store.save_ai_summary("10001", "抖音500元代金券售价328元", {
            "products": [{
                "name": "抖音500元代金券", "option_type": "代金券",
                "face_value": "500", "sale_price": "328", "max_stack": 4,
            }],
        })
        first = self.store.resolve_deterministic("10001", "500")
        self.assertIn("328元", first["reply"])
        followup = self.store.resolve_deterministic(
            "10001", "多少钱", store_context=first["query_context_update"],
        )
        self.assertIn("抖音500元代金券（售价328元）", followup["reply"])
        self.assertNotIn("抖音500元代金券（可叠加4张）", followup["reply"])

    def test_city_plus_unique_homophone_store_name_can_auto_match(self):
        self.store.save_v2_product(
            "10001", "多规格代金券",
            "100元代金券：售价64元\n300元代金券：售价192元",
        )
        first = self.store.import_store_text(
            "【山东省】\n【济南】济南弘阳广场店", "100门店", [],
        )
        second = self.store.import_store_text(
            "【山东省】\n【济南】济南弘阳广场店", "300门店", [],
        )
        for sku in self.store.get_v2_product("10001")["skus"]:
            list_id = first["id"] if sku["face_value"] == "100" else second["id"]
            self.store.set_sku_store_rule(
                "10001", sku["sku_key"], sku["sku_name"], "custom", [list_id],
            )

        result = self.store.resolve_deterministic("10001", "济南弘扬")
        self.assertEqual("stores", result["kind"])
        self.assertIn("济南弘阳广场店", result["reply"])
        self.assertNotIn("100元代金券", result["reply"])
        self.assertNotEqual("济南", result["store_query"])

    def test_nested_city_county_and_landmark_uses_smallest_grounded_area(self):
        self.store.save_v2_product(
            "10001", "多规格代金券",
            "100元代金券：售价64元\n300元代金券：售价192元",
        )
        first = self.store.import_store_text(
            "【山东省】\n【寿光】寿光万达广场店\n【潍坊】潍坊万达广场店",
            "100门店", [],
        )
        second = self.store.import_store_text(
            "【山东省】\n【寿光】寿光万达广场店", "300门店", [],
        )
        for sku in self.store.get_v2_product("10001")["skus"]:
            list_id = first["id"] if sku["face_value"] == "100" else second["id"]
            self.store.set_sku_store_rule(
                "10001", sku["sku_key"], sku["sku_name"], "custom", [list_id],
            )

        result = self.store.resolve_deterministic("10001", "山东潍坊寿光万达能用吗")
        self.assertEqual("stores_sku_recommendation", result["kind"])
        self.assertEqual(["寿光万达广场店"], [row["branch"] for row in result["store_matches"]])
        self.assertNotIn("潍坊万达广场店", result["reply"])

    def test_generic_landmark_across_cities_requires_geographic_anchor(self):
        self.store.save_v2_product(
            "10001", "多规格代金券",
            "100元代金券：售价64元\n300元代金券：售价192元",
        )
        first = self.store.import_store_text(
            "【山东省】\n【寿光】寿光万达广场店\n【潍坊】潍坊万达广场店\n【青岛】青岛万达广场店",
            "100门店", [],
        )
        second = self.store.import_store_text(
            "【山东省】\n【寿光】寿光万达广场店", "300门店", [],
        )
        for sku in self.store.get_v2_product("10001")["skus"]:
            list_id = first["id"] if sku["face_value"] == "100" else second["id"]
            self.store.set_sku_store_rule(
                "10001", sku["sku_key"], sku["sku_name"], "custom", [list_id],
            )

        result = self.store.resolve_deterministic("10001", "万达能用吗")
        self.assertEqual("stores_clarify", result["kind"])
        self.assertEqual("missing_area", result["store_status"])
        self.assertIn("请补充城市或区县", result["reply"])
        self.assertNotIn("100元代金券", result["reply"])

        context = {
            "query": result["store_query"],
            "matches": result["store_matches"],
            "status": result["store_status"],
            **result["store_context_update"],
        }
        narrowed = self.store.resolve_deterministic(
            "10001", "寿光", store_context=context,
        )
        self.assertEqual("stores_sku_recommendation", narrowed["kind"])
        self.assertEqual(["寿光万达广场店"], [row["branch"] for row in narrowed["store_matches"]])

    def test_multi_sku_unknown_specific_store_lists_all_city_alternatives(self):
        self.store.save_v2_product(
            "10001", "多规格代金券",
            "100元代金券：售价64元\n200元代金券：售价120元",
        )
        first = self.store.import_store_text(
            "【安徽省】\n【合肥】合肥之心城店", "100门店", [],
        )
        second = self.store.import_store_text(
            "【安徽省】\n【合肥】合肥银泰店", "200门店", [],
        )
        for sku in self.store.get_v2_product("10001")["skus"]:
            list_id = first["id"] if sku["face_value"] == "100" else second["id"]
            self.store.set_sku_store_rule(
                "10001", sku["sku_key"], sku["sku_name"], "custom", [list_id],
            )
        result = self.store.resolve_deterministic("10001", "合肥不存在万达能用吗")
        self.assertEqual("stores_sku_recommendation", result["kind"])
        self.assertEqual("allow", result["decision"])
        self.assertEqual("area_fallback", result["store_status"])
        self.assertIn("合肥不存在万达", result["reply"])
        self.assertIn("以下为合肥的所有可用门店", result["reply"])
        self.assertIn("合肥之心城店", result["reply"])
        self.assertIn("合肥银泰店", result["reply"])

    def test_multi_sku_city_over_three_stores_requires_narrowing(self):
        self.store.save_v2_product(
            "10001", "多规格代金券",
            "100元代金券：售价64元\n200元代金券：售价120元",
        )
        first = self.store.import_store_text(
            "【安徽省】\n【合肥】之心城店\n【合肥】砂之船店", "100门店", [],
        )
        second = self.store.import_store_text(
            "【安徽省】\n【合肥】银泰店\n【合肥】万象城店", "200门店", [],
        )
        for sku in self.store.get_v2_product("10001")["skus"]:
            list_id = first["id"] if sku["face_value"] == "100" else second["id"]
            self.store.set_sku_store_rule(
                "10001", sku["sku_key"], sku["sku_name"], "custom", [list_id],
            )
        result = self.store.resolve_deterministic("10001", "合肥有哪些门店")
        self.assertEqual("stores_clarify", result["kind"])
        self.assertEqual("too_many", result["store_status"])
        self.assertEqual(4, len(result["store_matches"]))
        self.assertIn("请补充区县、商圈", result["reply"])
        self.assertNotIn("100元代金券", result["reply"])

    def test_multi_sku_different_lists_with_same_physical_stores_do_not_narrow(self):
        self.store.save_v2_product(
            "10001", "多规格代金券",
            "100元代金券：售价64元\n200元代金券：售价120元\n300元代金券：售价192元",
        )
        source = (
            "【湖北省】\n【襄阳】襄阳万达店\n【襄阳】襄阳吾悦店\n"
            "【襄阳】襄阳民发店\n【襄阳】襄阳武商店"
        )
        first = self.store.import_store_text(source, "100门店", [])
        second = self.store.import_store_text(source, "200门店", [])
        for sku in self.store.get_v2_product("10001")["skus"]:
            list_id = first["id"] if sku["face_value"] == "100" else second["id"]
            self.store.set_sku_store_rule(
                "10001", sku["sku_key"], sku["sku_name"], "custom", [list_id],
            )
        result = self.store.resolve_deterministic("10001", "襄阳可以用不")
        self.assertEqual("stores", result["kind"])
        self.assertNotEqual("too_many", result.get("store_status"))
        for name in ("襄阳万达店", "襄阳吾悦店", "襄阳民发店", "襄阳武商店"):
            self.assertIn(name, result["reply"])
        self.assertIn("可以用，根据“襄阳”查询到可用门店：", result["reply"])
        context = {
            "query": result["store_query"], "matches": result["store_matches"],
            "status": result["store_status"],
            "store_sku_matrix": result.get("store_sku_matrix") or [],
        }
        for question, price in (("100呢", "64元"), ("那200的呢", "120元"), ("300的呢", "192元")):
            with self.subTest(question=question):
                followup = self.store.resolve_deterministic(
                    "10001", question, store_context=context,
                )
                self.assertNotEqual("stores_sku_recommendation", followup["kind"])
                self.assertIn(price, followup["reply"])

        stale_matrix = [{
            "store": result["store_matches"][0],
            "supported_skus": [self.store.get_v2_product("10001")["skus"][0]],
            "unknown_skus": [],
        }]
        self.assertIsNone(self.store.resolve_store_matrix_followup(
            "10001", self.store.get_v2_product("10001"), "300的呢",
            {**context, "store_sku_matrix": stale_matrix},
        ))

    def test_multi_sku_city_up_to_three_lists_each_store_specs(self):
        self.store.save_v2_product(
            "10001", "多规格代金券",
            "100元代金券：售价64元\n200元代金券：售价120元",
        )
        first = self.store.import_store_text(
            "【河南省】\n【郑州】郑州大卫城店\n【郑州】郑州二七万达店", "100门店", [],
        )
        second = self.store.import_store_text(
            "【河南省】\n【郑州】郑州二七万达店", "200门店", [],
        )
        for sku in self.store.get_v2_product("10001")["skus"]:
            list_id = first["id"] if sku["face_value"] == "100" else second["id"]
            self.store.set_sku_store_rule(
                "10001", sku["sku_key"], sku["sku_name"], "custom", [list_id],
            )
        result = self.store.resolve_deterministic("10001", "郑州")
        self.assertEqual("stores_sku_recommendation", result["kind"])
        self.assertEqual(2, len(result["store_matches"]))
        matrix = {
            row["store"]["branch"]: "、".join(
                sku["sku_name"] for sku in row["supported_skus"]
            )
            for row in result["store_sku_matrix"]
        }
        self.assertIn("100元代金券", matrix["郑州大卫城店"])
        self.assertNotIn("200元代金券", matrix["郑州大卫城店"])
        self.assertIn("100元代金券", matrix["郑州二七万达店"])
        self.assertIn("200元代金券", matrix["郑州二七万达店"])
        self.assertIn("\n\n【郑州大卫城店】：可用规格为", result["reply"])
        self.assertIn("100元代金券（售价64元）", result["reply"])
        self.assertIn("200元代金券（售价120元）", result["reply"])
        self.assertEqual([], result["reply_parts"])
        self.assertNotIn("&#x20;", result["reply"])

    def test_multi_sku_store_matrix_supports_amount_followup(self):
        self.store.save_v2_product(
            "10001", "多规格代金券",
            "100元代金券：售价64元\n200元代金券：售价120元",
        )
        first = self.store.import_store_text(
            "【河南省】\n【郑州】郑州大卫城店\n【郑州】郑州二七万达店", "100门店", [],
        )
        second = self.store.import_store_text(
            "【河南省】\n【郑州】郑州二七万达店", "200门店", [],
        )
        for sku in self.store.get_v2_product("10001")["skus"]:
            list_id = first["id"] if sku["face_value"] == "100" else second["id"]
            self.store.set_sku_store_rule(
                "10001", sku["sku_key"], sku["sku_name"], "custom", [list_id],
            )
        initial = self.store.resolve_deterministic("10001", "郑州有哪些门店")
        context = {
            "query": initial["store_query"], "matches": initial["store_matches"],
            "store_sku_matrix": initial["store_sku_matrix"], "status": "available",
        }
        followup = self.store.resolve_deterministic(
            "10001", "那200的呢", store_context=context,
        )
        self.assertEqual("stores_sku_recommendation", followup["kind"])
        lines = followup["reply"].splitlines()
        david = next(line for line in lines if "大卫城店" in line)
        wanda = next(line for line in lines if "二七万达店" in line)
        self.assertIn("200元代金券", david)
        self.assertIn("不能使用", david)
        self.assertIn("200元代金券", wanda)
        self.assertIn("可以使用", wanda)

    def test_single_store_followup_forms_cover_every_sku_and_keep_store_context(self):
        self.store.save_v2_product(
            "10001", "多规格代金券",
            "100元代金券：售价69元\n200元代金券：售价139元\n300元代金券：售价206元",
        )
        supported = self.store.import_store_text(
            "【浙江省】\n【杭州】杭州测试店", "100和200门店", [],
        )
        unsupported = self.store.import_store_text(
            "【浙江省】\n【宁波】宁波测试店", "300门店", [],
        )
        for sku in self.store.get_v2_product("10001")["skus"]:
            list_id = unsupported["id"] if sku["face_value"] == "300" else supported["id"]
            self.store.set_sku_store_rule(
                "10001", sku["sku_key"], sku["sku_name"], "custom", [list_id],
            )
        initial = self.store.resolve_deterministic("10001", "杭州测试店")
        context = {
            "query": initial["store_query"], "matches": initial["store_matches"],
            "store_sku_matrix": initial["store_sku_matrix"], "status": "available",
        }
        cases = {
            "100呢": ("100元代金券可以使用", False),
            "那200的呢": ("200元代金券可以使用", False),
            "300的呢": ("300元代金券不能使用", True),
            "那300元券能用吗": ("300元代金券不能使用", True),
            "300能不能用": ("300元代金券不能使用", True),
            "300的可以用吗": ("300元代金券不能使用", True),
            "请问300代金券支持吗": ("300元代金券不能使用", True),
            "300有吗": ("300元代金券不能使用", True),
            "206那个不能用吗": ("300元代金券不能使用", True),
            "139那个能用吗": ("200元代金券可以使用", False),
        }
        for question, (expected, should_offer_supported) in cases.items():
            with self.subTest(question=question):
                result = self.store.resolve_deterministic(
                    "10001", question, store_context=context,
                )
                self.assertEqual("stores_sku_recommendation", result["kind"])
                self.assertIn("杭州测试店", result["reply"])
                self.assertIn(expected, result["reply"])
                if should_offer_supported:
                    self.assertIn("100元代金券（售价69元）", result["reply"])
                    self.assertIn("200元代金券（售价139元）", result["reply"])

        for question in ("所有规格都能用吗", "那其他规格呢", "每个规格都可以用吗"):
            with self.subTest(question=question):
                result = self.store.resolve_deterministic(
                    "10001", question, store_context=context,
                )
                self.assertEqual("stores_sku_recommendation", result["kind"])
                self.assertIn("100元、200元代金券可以使用", result["reply"])
                self.assertIn("300元代金券不能使用", result["reply"])

        price = self.store.resolve_deterministic(
            "10001", "300多少钱", store_context=context,
        )
        self.assertNotEqual("stores_sku_recommendation", price["kind"])
        self.assertIn("206元", price["reply"])

        for question in ("300的代金券我拍下来是206元吗", "206是300的券吗"):
            with self.subTest(question=question):
                confirmation = self.store.resolve_deterministic(
                    "10001", question, store_context=context,
                )
                self.assertEqual("voucher_value", confirmation["kind"])
                self.assertIn("300元代金券", confirmation["reply"])
                self.assertIn("售价206元", confirmation["reply"])

    def test_different_store_amount_wording_uses_current_store_skus_and_limits(self):
        self.store.save_v2_product(
            "10001", "多规格代金券",
            "100元代金券：售价69元，最多使用2张\n"
            "200元代金券：售价139元，最多使用1张\n"
            "300元代金券：售价206元，最多使用1张",
        )
        hangzhou = self.store.import_store_text(
            "【浙江省】\n【杭州】杭州测试店", "杭州100和200", [],
        )
        ningbo = self.store.import_store_text(
            "【浙江省】\n【宁波】宁波测试店", "宁波300", [],
        )
        for sku in self.store.get_v2_product("10001")["skus"]:
            list_id = ningbo["id"] if sku["face_value"] == "300" else hangzhou["id"]
            self.store.set_sku_store_rule(
                "10001", sku["sku_key"], sku["sku_name"], "custom", [list_id],
            )
        initial = self.store.resolve_deterministic("10001", "杭州测试店")
        context = {
            "query": initial["store_query"], "matches": initial["store_matches"],
            "store_sku_matrix": initial["store_sku_matrix"], "status": "available",
        }
        questions = (
            "400怎么买", "消费400怎么拍", "账单一共400", "400元买哪种券",
            "400用什么券", "400该拍几张", "预算400给个方案",
            "400怎么最划算", "400元左右怎么配", "怎么买能抵400",
            "我这边一共400", "400的呢", "400多少钱",
        )
        for question in questions:
            with self.subTest(question=question):
                result = self.store.resolve_deterministic(
                    "10001", question, store_context=context,
                )
                self.assertNotEqual("stores_sku_recommendation", result["kind"])
                self.assertIn("2张100元代金券", result["reply"])
                self.assertIn("剩余200元", result["reply"])
                self.assertIn("最多使用2张", result["reply"])
                self.assertNotIn("300元代金券", result["reply"])

        unavailable = self.store.resolve_deterministic(
            "10001", "300的呢", store_context=context,
        )
        self.assertEqual("stores_sku_recommendation", unavailable["kind"])
        self.assertIn("300元代金券不能使用", unavailable["reply"])

        alternative = self.store.resolve_deterministic(
            "10001", "300怎么买", store_context=context,
        )
        self.assertEqual("consumption_plan", alternative["kind"])
        self.assertIn("2张100元代金券", alternative["reply"])
        self.assertIn("剩余100元", alternative["reply"])
        self.assertNotIn("300元代金券", alternative["reply"])

    def test_same_store_amount_followups_recommend_purchase_plan(self):
        self.store.save_v2_product(
            "10001", "多规格代金券",
            "100元代金券：售价69元，最多使用2张\n"
            "200元代金券：售价139元，最多使用2张\n"
            "300元代金券：售价206元，最多使用1张",
        )
        source = "【浙江省】\n【杭州】杭州测试店"
        first = self.store.import_store_text(source, "规格门店A", [])
        second = self.store.import_store_text(source, "规格门店B", [])
        for sku in self.store.get_v2_product("10001")["skus"]:
            list_id = first["id"] if sku["face_value"] == "100" else second["id"]
            self.store.set_sku_store_rule(
                "10001", sku["sku_key"], sku["sku_name"], "custom", [list_id],
            )
        initial = self.store.resolve_deterministic("10001", "杭州测试店")
        context = {
            "query": initial["store_query"], "matches": initial["store_matches"],
            "store_sku_matrix": initial.get("store_sku_matrix") or [],
            "status": "available",
        }
        for question in ("400", "400呢", "那400的呢", "400怎么凑", "账单400"):
            with self.subTest(question=question):
                result = self.store.resolve_deterministic(
                    "10001", question, store_context=context,
                )
                self.assertNotEqual("stores_sku_recommendation", result["kind"])
                self.assertIn("200元代金券", result["reply"])
                self.assertIn("共支付278元", result["reply"])
                self.assertIn("可抵扣400元", result["reply"])

        exact = self.store.resolve_deterministic(
            "10001", "300的呢", store_context=context,
        )
        self.assertNotEqual("stores_sku_recommendation", exact["kind"])
        self.assertIn("300元代金券", exact["reply"])
        self.assertIn("206元", exact["reply"])

        limited = self.store.resolve_deterministic(
            "10001", "500怎么凑", store_context=context,
        )
        self.assertIn("2张200元代金券", limited["reply"])
        self.assertIn("剩余100元", limited["reply"])
        self.assertIn("最多使用2张", limited["reply"])

    def test_multi_store_amount_plan_is_calculated_per_store(self):
        self.store.save_v2_product(
            "10001", "多规格代金券",
            "100元代金券：售价69元，最多使用2张\n"
            "200元代金券：售价120元，最多使用2张",
        )
        first = self.store.import_store_text(
            "【河南省】\n【郑州】郑州大卫城店\n【郑州】郑州二七万达店", "100门店", [],
        )
        second = self.store.import_store_text(
            "【河南省】\n【郑州】郑州二七万达店", "200门店", [],
        )
        for sku in self.store.get_v2_product("10001")["skus"]:
            list_id = first["id"] if sku["face_value"] == "100" else second["id"]
            self.store.set_sku_store_rule(
                "10001", sku["sku_key"], sku["sku_name"], "custom", [list_id],
            )
        initial = self.store.resolve_deterministic("10001", "郑州有哪些门店")
        context = {
            "query": initial["store_query"], "matches": initial["store_matches"],
            "store_sku_matrix": initial["store_sku_matrix"], "status": "available",
        }
        result = self.store.resolve_deterministic(
            "10001", "400怎么买", store_context=context,
        )
        self.assertEqual("consumption_plan", result["kind"])
        self.assertIn("购券方案不同", result["reply"])
        self.assertIn("【郑州大卫城店】", result["reply"])
        self.assertIn("2张100元代金券", result["reply"])
        self.assertIn("剩余200元", result["reply"])
        self.assertIn("【郑州二七万达店】", result["reply"])
        self.assertIn("2张200元代金券", result["reply"])
        self.assertIn("可抵扣400元", result["reply"])

    def test_multi_store_singular_context_reference_requires_branch_name(self):
        self.store.save_v2_product(
            "10001", "多规格代金券",
            "100元代金券：售价64元\n200元代金券：售价120元",
        )
        first = self.store.import_store_text(
            "【河南省】\n【郑州】郑州大卫城店\n【郑州】郑州二七万达店", "100门店", [],
        )
        second = self.store.import_store_text(
            "【河南省】\n【郑州】郑州二七万达店", "200门店", [],
        )
        for sku in self.store.get_v2_product("10001")["skus"]:
            list_id = first["id"] if sku["face_value"] == "100" else second["id"]
            self.store.set_sku_store_rule(
                "10001", sku["sku_key"], sku["sku_name"], "custom", [list_id],
            )
        initial = self.store.resolve_deterministic("10001", "郑州")
        context = {
            "query": initial["store_query"], "matches": initial["store_matches"],
            "store_sku_matrix": initial["store_sku_matrix"], "status": "available",
        }
        followup = self.store.resolve_deterministic(
            "10001", "这家店能用吗", store_context=context,
        )
        self.assertEqual("stores_clarify", followup["kind"])
        self.assertIn("刚才查询到多家门店", followup["reply"])
        self.assertIn("第一家、第二家", followup["reply"])

    def test_too_many_store_context_combines_narrowing_followup(self):
        self.store.save_v2_product(
            "10001", "多规格代金券",
            "100元代金券：售价64元\n200元代金券：售价120元",
        )
        first = self.store.import_store_text(
            "【安徽省】\n【合肥】之心城店\n【合肥】砂之船店", "100门店", [],
        )
        second = self.store.import_store_text(
            "【安徽省】\n【合肥】银泰店\n【合肥】万象城店", "200门店", [],
        )
        for sku in self.store.get_v2_product("10001")["skus"]:
            list_id = first["id"] if sku["face_value"] == "100" else second["id"]
            self.store.set_sku_store_rule(
                "10001", sku["sku_key"], sku["sku_name"], "custom", [list_id],
            )
        initial = self.store.resolve_deterministic("10001", "合肥有哪些门店")
        context = {
            "query": initial["store_query"], "matches": initial["store_matches"],
            "status": initial["store_status"],
            **initial["store_context_update"],
        }
        followup = self.store.resolve_deterministic(
            "10001", "之心城", store_context=context,
        )
        self.assertEqual("stores_sku_recommendation", followup["kind"])
        self.assertIn("合肥之心城", followup["store_query"])
        self.assertEqual(1, len(followup["store_matches"]))
        self.assertIn("100元代金券", followup["reply"])
        self.assertNotIn("200元代金券", followup["reply"])

    def test_semantic_multi_question_is_reanswered_only_from_local_facts(self):
        self.store.save_v2_product(
            "10001", "测试代金券",
            "300元代金券：售价192元，付款后发送电子券码",
        )
        self.store.import_store_text(
            "【山东省】\n【济南】济南万象城店", "济南门店", ["10001"],
        )
        analysis = {
            "questions": [
                {"intent": "store", "evidence": "济南可以用吗",
                 "slots": {"store_query": "济南"}, "confidence": 0.96},
                {"intent": "purchase", "evidence": "300的可以直接拍吗",
                 "slots": {"sku_amount": "300"}, "confidence": 0.93},
            ],
            "needs_clarification": False,
        }
        result = self.store.resolve_semantic_analysis(
            "10001", "济南可以用吗，300的可以直接拍吗", analysis,
            original_result={"kind": "stores", "decision": "allow"},
        )
        self.assertEqual("semantic_multi_intent", result["kind"])
        self.assertEqual(["store", "purchase"], result["resolved_intents"])
        self.assertIn("济南万象城店", result["reply"])
        self.assertIn("300元代金券", result["reply"])
        self.assertIn("售价192元", result["reply"])
        self.assertIn("\n\n2. 购买", result["reply"])

    def test_semantic_stock_and_store_questions_use_existing_local_answers(self):
        self.store.save_v2_product(
            "10001", "测试代金券", "100元代金券：售价66元",
        )
        self.store.import_store_text(
            "【上海市】\n【上海】上海万象城店", "上海门店", ["10001"],
        )
        message = "这个有货吗 上海能用吗"
        analysis = {
            "questions": [
                {"intent": "purchase", "evidence": "这个有货吗",
                 "slots": {}, "confidence": 0.98},
                {"intent": "store", "evidence": "上海能用吗",
                 "slots": {"store_query": "上海"}, "confidence": 0.99},
            ],
            "needs_clarification": False,
        }
        original = self.store.resolve_deterministic("10001", message)
        result = self.store.resolve_semantic_analysis(
            "10001", message, analysis, original_result=original,
        )
        self.assertEqual("semantic_multi_intent", result["kind"])
        self.assertEqual(["purchase", "store"], result["resolved_intents"])
        self.assertIn("有的，当前商品还在售，可以直接拍下。", result["reply"])
        self.assertIn("上海万象城店", result["reply"])
        self.assertIn("\n\n2. 适用门店", result["reply"])

    def test_semantic_single_question_never_replaces_existing_rule_result(self):
        analysis = {
            "questions": [{
                "intent": "usage", "evidence": "这个咋用",
                "slots": {}, "confidence": 0.9,
            }],
            "needs_clarification": False,
        }
        result = self.store.resolve_semantic_analysis(
            "10001", "这个咋用", analysis,
            original_result={"kind": "delivery_usage", "decision": "allow"},
        )
        self.assertIsNone(result)

    def test_explicit_sku_store_query_never_falls_back_to_other_sku(self):
        self.store.save_ai_summary("10001", "两种套餐", {
            "sale_options": [
                {"name": "单人套餐", "sale_price": "88", "people_counts": [1]},
                {"name": "双人套餐", "sale_price": "168", "people_counts": [2]},
            ]
        })
        only_double = self.store.import_store_text("【湖北省】\n【武汉】天地店", "双人门店", [])
        skus = self.store.get_v2_product("10001")["skus"]
        single = next(sku for sku in skus if "单人" in sku["sku_name"])
        double = next(sku for sku in skus if "双人" in sku["sku_name"])
        self.store.set_sku_store_rule("10001", single["sku_key"], single["sku_name"], "custom", [])
        self.store.set_sku_store_rule("10001", double["sku_key"], double["sku_name"], "custom", [only_double["id"]])
        result = self.store.resolve_deterministic("10001", "单人套餐在武汉天地店能用吗")
        self.assertEqual("review", result["decision"])
        self.assertIn("暂未配置适用门店", result["reply"])
        self.assertNotIn("可以", result["reply"])

    def test_excel_import_recovers_bad_dimension_and_loose_columns(self):
        result = self.store.import_store_list(
            self.create_irregular_store_sheet(), "不规则门店表", ["10001"]
        )
        self.assertEqual(result["store_count"], 1)
        match = self.store.search_store("10001", "国贸")
        self.assertEqual(match["status"], "available")
        store = match["matches"][0]
        self.assertEqual(store["address"], "建国路88号")
        self.assertEqual(store["phone"], "010-1234")
        self.assertEqual(store["business_hours"], "10:00-22:00")
        reply = self.store.format_store_reply(store, "国贸中心店的地址电话和营业时间")
        self.assertIn("建国路88号", reply)
        self.assertIn("010-1234", reply)
        self.assertIn("10:00-22:00", reply)

    def test_fuzzy_store_returns_all_matches_grouped(self):
        self.store.import_store_list(self.create_store_sheet(), "北京门店", ["10001"])
        result = self.store.search_store("10001", "万象")
        self.assertEqual(result["status"], "available")
        self.assertGreaterEqual(len(result["matches"]), 2)
        reply = self.store.format_store_matches(result["matches"], "万象")
        self.assertIn("东城万象店", reply)
        self.assertIn("西城万象店", reply)
        self.assertNotIn("等店", reply)
        self.assertNotIn("搜索到以下可用门店", reply)

    def test_exact_single_store_confirmation_uses_short_reply(self):
        self.store.import_store_text(
            "【浙江省】\n【杭州】杭州萧山万象汇店", "杭州门店", ["10001"],
        )
        result = self.store.resolve_deterministic(
            "10001", "杭州萧山万象汇店能用吗",
        )
        self.assertEqual("stores", result["kind"])
        self.assertEqual(
            "可以用，根据“杭州萧山万象汇店”查询到可用门店：杭州萧山万象汇店。",
            result["reply"],
        )

    def test_spoken_store_query_is_cleaned_and_always_gets_a_reply(self):
        self.store.import_store_list(self.create_store_sheet(), "北京门店", ["10001"])
        self.assertEqual("沧州", extract_store_query("老板，沧州可以使用嘛？"))
        result = self.store.resolve_deterministic("10001", "老板，沧州可以使用嘛？")
        self.assertEqual("stores", result["kind"])
        self.assertEqual("deny", result["decision"])
        self.assertEqual(
            "根据“沧州”未查询到可用门店。当前商品在该地区暂无已配置的可用门店。",
            result["reply"],
        )

    def test_generic_store_intent_without_location_only_clarifies(self):
        self.store.import_store_list(self.create_store_sheet(), "北京门店", ["10001"])
        for message in (
            "适用门店问题", "想问一下门店", "哪些店能用", "哪里可以用", "有什么适用门店",
        ):
            with self.subTest(message=message):
                query = extract_store_query(message)
                self.assertFalse(query)
                result = self.store.resolve_deterministic("10001", message)
                self.assertEqual("stores_clarify", result["kind"])
                self.assertIn("具体城市", result["reply"])
                self.assertNotIn("广场店", result["reply"])

    def test_generic_store_words_never_enter_fuzzy_store_matching(self):
        self.store.import_store_list(self.create_store_sheet(), "北京门店", ["10001"])
        for query in ("问题", "门店问题", "适用门店", "查询", "咨询"):
            with self.subTest(query=query):
                result = self.store.search_store("10001", query)
                self.assertEqual("missing_query", result["status"])
                self.assertEqual([], result["matches"])

    def test_store_entity_guard_keeps_existing_valid_location_queries(self):
        self.store.import_store_list(self.create_store_sheet(), "北京门店", ["10001"])
        self.assertEqual("北京", extract_store_query("北京哪些店能用"))
        result = self.store.resolve_deterministic("10001", "北京哪些店能用")
        self.assertEqual("stores", result["kind"])
        self.assertIn("北京", result["reply"])
        self.assertIn("可以用，根据“北京”查询到可用门店：", result["reply"])

    def test_store_negative_confirmation_stays_in_store_intent(self):
        self.store.import_store_text(
            "【河南省】\n【郑州】郑州正弘城店", "郑州门店", ["10001"],
        )
        for message in ("万象城不能用吗", "宁波能不能用", "宁波不可以用吗"):
            with self.subTest(message=message):
                result = self.store.resolve_deterministic("10001", message)
                self.assertEqual("stores", result["kind"])
                self.assertNotIn("退款", result["reply"])
                self.assertNotIn("72小时", result["reply"])

    def test_place_names_are_not_glued_to_polite_prefixes_or_usage_suffixes(self):
        cases = {
            "那个宁波能不能用呀": "宁波",
            "请问一下郑州市可以用吗": "郑州市",
            "想咨询鄞州区不可以用吗": "鄞州区",
            "麻烦问下姜山镇能用吗": "姜山镇",
            "这个横溪乡可不可以使用呢": "横溪乡",
        }
        for message, expected in cases.items():
            with self.subTest(message=message):
                self.assertEqual(expected, extract_store_query(message))

    def test_store_query_strips_only_current_product_brand_noise(self):
        product = {"title": "NEED韩国料理+need品牌美团电子券"}
        self.assertEqual("南昌万象城", extract_store_query("南昌万象城need", product))
        self.assertEqual("南昌万象城", extract_store_query("NEED南昌万象城能用吗", product))
        self.assertEqual(
            "宁波鄞州区印象城",
            extract_store_query("请问宁波鄞州区印象城周末两个人能用吗呀", product),
        )
        self.assertEqual("成都IFS", extract_store_query("成都IFS能用吗", product))
        self.assertEqual(
            "武汉万象城",
            extract_store_query("鱼酷武汉万象城支持吗", product_brand="鱼酷"),
        )

    def test_repeated_title_prefix_provides_a_clean_chinese_brand_boundary(self):
        product = {
            "title": "【全国】鱼酷烤鱼2-3人套餐+鱼酷 活鱼烤鱼 烤鱼券 单条鱼",
            "structured": {},
        }
        self.assertEqual("鱼酷", self.store.extract_brand(product))
        self.assertEqual(
            "任丘悦都汇店",
            extract_store_query(
                "鱼酷 任丘悦都汇店", product=product,
                product_brand=self.store.extract_brand(product),
            ),
        )

    def test_store_negative_followup_uses_previous_unavailable_query(self):
        self.store.import_store_text(
            "【河南省】\n【郑州】郑州正弘城店", "郑州门店", ["10001"],
        )
        result = self.store.resolve_deterministic(
            "10001", "确定不能用吗",
            store_context={"query": "郑州万象城", "matches": [], "status": "unavailable"},
        )
        self.assertEqual("stores", result["kind"])
        self.assertIn("郑州万象城", result["reply"])
        self.assertNotIn("退款", result["reply"])

    def test_actual_paid_coupon_failure_without_refund_intent_uses_troubleshooting(self):
        result = self.store.resolve_deterministic("10001", "我已经付款，券码在万象城核销失败")
        self.assertEqual("coupon_troubleshooting", result["kind"])
        self.assertNotIn("退款", result["reply"])

    def test_unspecified_day_defaults_to_store_business_hours(self):
        self.store.save_v2_product("10001", "测试券", "100元代金券售价68.8元。")
        weekend = self.store.resolve_deterministic("10001", "周末能用吗")
        self.assertEqual("day_use", weekend["kind"])
        self.assertIn("营业时间内可以使用", weekend["reply"])
        self.assertNotIn("72小时", weekend["reply"])
        summary = self.store.build_knowledge_summary(self.store.get_v2_product("10001"))
        self.assertIn("【使用时间】\n适用门店营业时间内可用。", summary)

    def test_raw_workday_use_time_overrides_ai_business_hours_default(self):
        self.store.save_v2_product(
            "10001", "朱富贵代金券",
            "【商品规格】\n"
            "朱富贵100元代金券：售价90元，发100元券1张（工作日可用）\n"
            "朱富贵200元代金券：售价180元，发200元券1张（工作日可用）\n\n"
            "【使用时间】\n（工作日可用）\n\n"
            "【有效期】\n当天买当天用。",
        )
        product = self.store.save_ai_summary(
            "10001", "错误的营业时间默认值",
            {
                "summary": "错误的营业时间默认值",
                "products": [
                    {"name": "朱富贵100元代金券", "face_value": "100", "sale_price": "90"},
                    {"name": "朱富贵200元代金券", "face_value": "200", "sale_price": "180"},
                ],
                "facts": {"使用时间": "适用门店营业时间内可用", "有效期": "当天买当天用"},
                "time_rules": [],
            },
        )
        self.assertIn("【使用时间】\n工作日可用。", product["ai_summary"])
        self.assertNotIn("【使用时间】\n适用门店营业时间内可用。", product["ai_summary"])

    def test_holiday_wording_covers_weekend_without_literal_weekend_sku(self):
        self.store.save_v2_product(
            "10001", "鱼酷烤鱼", "2.6斤单鱼套餐：售价118元。平日节假日营业时间通用。",
        )
        weekend = self.store.resolve_deterministic("10001", "周末能用吗")
        self.assertEqual("day_use", weekend["kind"])
        self.assertIn("可以", weekend["reply"])
        self.assertNotIn("没有适用于周末", weekend["reply"])

    def test_explicit_weekend_restriction_overrides_business_hours_default(self):
        self.store.save_v2_product(
            "10001", "测试券", "100元代金券售价68.8元。仅限工作日使用。",
        )
        result = self.store.resolve_deterministic("10001", "周末能用吗")
        self.assertEqual("day_use", result["kind"])
        self.assertIn("周末不可用", result["reply"])

    def test_specific_unavailable_date_is_kept_with_default_business_hours(self):
        self.store.save_ai_summary("10001", "测试券", {
            "facts": {"使用时间": "2026年9月25日至9月27日（中秋节）不可用"}
        })
        summary = self.store.build_knowledge_summary(self.store.get_v2_product("10001"))
        self.assertIn("2026年9月25日至9月27日", summary)
        self.assertIn("除上述明确不可用日期外，适用门店营业时间内可用", summary)

    def test_ai_summary_preserves_every_manual_rule_when_model_output_is_sparse(self):
        source = (
            "⚠️仅适用于部分胡恰门店，拍前请咨询。\n\n"
            "【商品信息】\n"
            "①200元代金券：138元（单次消费最多可用3张）\n"
            "②100元代金券：72元（单次消费最多可用3张）\n\n"
            "相同面额代金券最多叠加3张，不同面额代金券不能互相叠加\n"
            "----------------\n\n"
            "【使用规则】\n"
            "1. 除中秋节（9.25-9.27）、国庆节（10.1-10.7）外，营业时间内可用。\n"
            "2. 全场通用。\n"
            "3. 仅限堂食。\n"
            "4. 无需预约，高峰期可能需要等位。\n"
            "5. 团购用户不可同时享受商家其他优惠，不可与其它代金券叠加使用。\n"
            "6. 本单发票由商家提供，详情请咨询商家。"
        )
        self.store.save_v2_product("10001", "胡恰代金券", source)
        product = self.store.save_ai_summary("10001", "被过度压缩的AI草稿", {
            "facts": {
                "使用时间": "适用门店营业时间内可用",
                "叠加规则": "仅支持同面额代金券叠加，每次最多使用3张",
            },
            "products": [
                {"name": "100元代金券", "face_value": "100", "sale_price": "72", "max_stack": "3"},
                {"name": "200元代金券", "face_value": "200", "sale_price": "138", "max_stack": "3"},
            ],
            "time_rules": [],
        })
        summary = product["ai_summary"]
        for expected in (
            "仅适用于部分胡恰门店", "拍前请咨询", "中秋节", "9.25-9.27",
            "国庆节", "10.1-10.7", "全场通用", "仅限堂食", "无需预约",
            "高峰期可能需要等位", "商家其他优惠", "其它代金券叠加使用",
            "发票由商家提供", "详情请咨询商家",
        ):
            with self.subTest(expected=expected):
                self.assertIn(expected, summary)

    def test_ai_summary_renders_all_structured_rule_fields(self):
        self.store.save_v2_product("10001", "规则完整性测试", "100元代金券：售价72元。")
        product = self.store.save_ai_summary("10001", "简略草稿", {
            "facts": {
                "适用门店范围": "仅适用于部分门店，拍前请咨询",
                "不可用日期": "中秋节和国庆节不可用",
                "堂食限制": "仅限堂食",
                "预约要求": "无需预约，高峰期可能需要等位",
                "优惠同享": "不可同时享受商家其他优惠",
                "发票规则": "发票由商家提供",
                "下单前提醒": "详情请咨询商家",
            },
            "products": [{"name": "100元代金券", "face_value": "100", "sale_price": "72"}],
        })
        summary = product["ai_summary"]
        self.assertIn("【适用范围】", summary)
        self.assertIn("【使用规则】", summary)
        self.assertIn("【退款与发票】", summary)
        self.assertIn("【提醒】", summary)
        for value in product["structured"]["facts"].values():
            self.assertIn(value, summary)

    def test_standalone_unknown_city_is_a_store_query_without_cross_city_fuzzy_match(self):
        self.store.import_store_text(
            "【浙江省】\n【杭州】武林广场店、西湖银泰店",
            "浙江门店", ["10001"],
        )
        result = self.store.resolve_deterministic("10001", "湖州")
        self.assertEqual("stores", result["kind"])
        self.assertEqual("deny", result["decision"])
        self.assertEqual(
            "根据“湖州”未查询到可用门店。当前商品在该地区暂无已配置的可用门店。",
            result["reply"],
        )
        self.assertNotIn("杭州", result["reply"])

    def test_city_store_query_only_returns_matching_area(self):
        self.store.import_store_text(
            "【河北省】\n【邯郸】邯郸万象汇店\n【石家庄】石家庄万象城店",
            "河北门店",
            ["10001"],
        )
        result = self.store.resolve_deterministic("10001", "河北邯郸能用嘛")
        self.assertEqual("stores", result["kind"])
        self.assertIn("邯郸万象汇店", result["reply"])
        self.assertNotIn("石家庄万象城店", result["reply"])

    def test_city_only_lists_all_city_stores_and_specific_query_keeps_landmark(self):
        self.store.import_store_list(self.create_multi_region_store_sheet(), "多地区门店", ["10001"])
        city = self.store.resolve_deterministic("10001", "深圳有门店吗")
        self.assertEqual("stores", city["kind"])
        self.assertIn("可以用，根据“深圳”查询到可用门店：", city["reply"])
        self.assertIn("龙岗万科里店", city["reply"])
        self.assertIn("西丽益田假日里店", city["reply"])
        self.assertNotIn("广州万达广场店", city["reply"])

        specific = self.store.resolve_deterministic("10001", "深圳龙岗万科里店可以用吗")
        self.assertEqual("stores", specific["kind"])
        self.assertIn("龙岗万科里店", specific["reply"])
        self.assertNotIn("西丽益田假日里店", specific["reply"])

    def test_unknown_specific_store_falls_back_to_all_stores_in_mentioned_city(self):
        self.store.save_v2_product(
            "10001", "鱼酷活鱼烤鱼", "鱼酷活鱼烤鱼单条鱼2.6斤：118元",
        )
        self.store.import_store_text(
            "【湖北省】\n【武汉】武汉江汉路M+店、武昌万象城店\n"
            "【湖北省】\n【襄阳】襄阳万达店",
            "湖北门店", ["10001"],
        )

        result = self.store.resolve_deterministic("10001", "武汉m➕ 118可以用吗")

        self.assertEqual("stores", result["kind"])
        self.assertEqual("allow", result["decision"])
        self.assertEqual("area_fallback", result["store_status"])
        self.assertTrue(result["reply"].startswith(
            "根据“武汉m➕ 118”未查询到可用门店。\n"
            "以下为武汉的所有可用门店："
        ))
        self.assertIn("武汉江汉路M+店", result["reply"])
        self.assertIn("武昌万象城店", result["reply"])
        self.assertNotIn("襄阳万达店", result["reply"])

    def test_unknown_specific_store_prefers_mentioned_district_over_city(self):
        self.store.import_store_list(
            self.create_multi_region_store_sheet(), "多地区门店", ["10001"]
        )

        result = self.store.resolve_deterministic("10001", "深圳南山区不存在商场能用吗")

        self.assertEqual("area_fallback", result["store_status"])
        self.assertIn("以下为深圳南山的所有可用门店", result["reply"])
        self.assertIn("西丽益田假日里店", result["reply"])
        self.assertNotIn("龙岗万科里店", result["reply"])
        self.assertNotIn("广州万达广场店", result["reply"])

    def test_store_query_matches_landmark_in_address_and_short_brand_alias(self):
        self.store.save_v2_product(
            "10001", "胡恰·景德江西菜代金券", "100元代金券：售价72元。",
        )
        path = os.path.join(self.temp.name, "hucha-stores.xlsx")
        book = Workbook()
        sheet = book.active
        sheet.append(["店名", "分店名", "省", "市", "区/县", "地址"])
        sheet.append([
            "胡恰·景德江西菜", "武汉首店", "湖北", "武汉", "武昌区",
            "中南路街道武珞路598号武商梦时代7层B区711号",
        ])
        book.save(path)
        self.store.import_store_list(path, "胡恰门店", ["10001"])

        product = self.store.get_v2_product("10001")
        query = extract_store_query(
            "胡恰武汉梦时代", product=product,
            product_brand=self.store.extract_brand(product),
        )
        self.assertEqual("武汉梦时代", query)
        result = self.store.resolve_deterministic("10001", "胡恰武汉梦时代")
        self.assertEqual("stores", result["kind"])
        self.assertEqual("allow", result["decision"])
        self.assertIn("武汉首店", result["reply"])
        self.assertEqual("address", result["store_matches"][0]["match_quality"])

        evidence = result["store_matches"][0]["match_evidence"]
        self.assertIn("address", {item["type"] for item in evidence})
        self.assertGreaterEqual(result["store_matches"][0]["match_score"], 85)
        self.assertEqual([], result["store_matches"][0]["unresolved_terms"])

    def test_store_query_ignores_neutral_words_but_keeps_grounded_anchors(self):
        path = os.path.join(self.temp.name, "neutral-query-stores.xlsx")
        book = Workbook()
        sheet = book.active
        sheet.append(["店名", "分店名", "省", "市", "区/县", "地址"])
        sheet.append([
            "胡恰·景德江西菜", "武汉首店", "湖北", "武汉", "武昌区",
            "中南路街道武珞路598号武商梦时代7层B区711号",
        ])
        book.save(path)
        self.store.import_store_list(path, "胡恰门店", ["10001"])

        result = self.store.resolve_deterministic(
            "10001", "麻烦帮我看一下梦时代那家能不能用"
        )
        self.assertEqual("stores", result["kind"])
        self.assertEqual("allow", result["decision"])
        self.assertIn("武汉首店", result["reply"])
        self.assertEqual("梦时代", result["store_query"])

    def test_store_relation_query_requires_exact_store_or_address(self):
        self.store.import_store_text(
            "【湖北省】\n【武汉】武昌梦时代店", "武汉门店", ["10001"],
        )
        result = self.store.resolve_deterministic("10001", "武汉梦时代旁边那家能用吗")
        self.assertEqual("stores_clarify", result["kind"])
        self.assertEqual("location_relation_unverified", result["store_status"])
        self.assertNotIn("可以使用", result["reply"])

    def test_store_correction_uses_only_replacement_target(self):
        self.store.import_store_text(
            "【湖北省】\n【武汉】武汉梦时代店、武汉万象城店",
            "武汉门店", ["10001"],
        )
        result = self.store.resolve_deterministic(
            "10001", "不是梦时代，是武汉万象城能用吗"
        )
        self.assertEqual("stores", result["kind"])
        self.assertIn("武汉万象城店", result["reply"])
        self.assertNotIn("武汉梦时代店", result["reply"])

    def test_store_query_with_conflicting_regions_requires_clarification(self):
        self.store.import_store_text(
            "【湖北省】\n【武汉】武汉万象城店\n"
            "【江苏省】\n【南京】南京万象城店",
            "跨地区门店", ["10001"],
        )
        result = self.store.resolve_deterministic("10001", "武汉还是南京万象城能用吗")
        self.assertEqual("stores_clarify", result["kind"])
        self.assertEqual("conflicting_scope", result["store_status"])

    def test_store_street_and_house_number_are_strong_address_evidence(self):
        self.store.import_store_list(self.create_store_sheet(), "北京门店", ["10001"])
        result = self.store.search_store("10001", "复兴路69号")
        self.assertEqual("available", result["status"])
        self.assertEqual("华熙五棵松店", result["matches"][0]["branch"])
        self.assertEqual("address", result["matches"][0]["match_quality"])
        self.assertEqual(90, result["matches"][0]["match_score"])
        resolved = self.store.resolve_deterministic("10001", "复兴路69号")
        self.assertEqual("stores", resolved["kind"])
        self.assertEqual("allow", resolved["decision"])
        self.assertEqual(
            "可以用，根据“复兴路69号”查询到可用门店：华熙五棵松店。",
            resolved["reply"],
        )

    def test_city_with_which_stores_wording_lists_only_that_city(self):
        self.store.import_store_text(
            "【广东省】\n【深圳】南昌品牌深圳店\n【江西省】\n【南昌】万寿宫店、北京东路店",
            "跨地区门店", ["10001"],
        )
        result = self.store.resolve_deterministic("10001", "南昌有哪些可以用")
        self.assertEqual("stores", result["kind"])
        self.assertIn("可以用，根据“南昌”查询到可用门店：", result["reply"])
        self.assertIn("万寿宫店", result["reply"])
        self.assertIn("北京东路店", result["reply"])
        self.assertNotIn("南昌品牌深圳店", result["reply"])

    def test_store_fuzzy_match_preserves_province_scope(self):
        self.store.import_store_list(self.create_multi_region_store_sheet(), "多地区门店", ["10001"])
        result = self.store.resolve_deterministic("10001", "广东万达有吗")
        self.assertEqual("stores", result["kind"])
        self.assertIn("广州万达广场店", result["reply"])
        self.assertNotIn("南京万达广场店", result["reply"])

    def test_region_plus_brand_lists_branches_in_that_region(self):
        self.store.import_store_list(self.create_multi_region_store_sheet(), "多地区门店", ["10001"])
        result = self.store.resolve_deterministic("10001", "深圳测试品牌有吗")
        self.assertEqual("stores", result["kind"])
        self.assertIn("龙岗万科里店", result["reply"])
        self.assertIn("西丽益田假日里店", result["reply"])
        self.assertNotIn("广州万达广场店", result["reply"])

    def test_store_name_typo_and_unavailable_template(self):
        self.store.import_store_list(self.create_multi_region_store_sheet(), "多地区门店", ["10001"])
        fuzzy = self.store.resolve_deterministic("10001", "南山西丽益田假日店有吗")
        self.assertIn("西丽益田假日里店", fuzzy["reply"])
        missing = self.store.resolve_deterministic("10001", "焦作万达能用吗")
        self.assertEqual(
            "根据“焦作万达”未查询到可用门店。\n"
            "请补充所在城市或区县，我为您查询该地区的所有可用门店。",
            missing["reply"],
        )

    def test_standalone_region_name_is_a_store_query(self):
        self.store.import_store_list(self.create_multi_region_store_sheet(), "多地区门店", ["10001"])
        result = self.store.resolve_deterministic("10001", "寿光可以用不")
        self.assertEqual("stores", result["kind"])
        self.assertIn("寿光万达广场店", result["reply"])

    def test_bare_store_city_mall_and_address_are_store_queries(self):
        self.store.import_store_text(
            "【湖北省】\n【武汉】武汉首店、和平店、武汉万达广场店\n"
            "【广东省】\n【汕头】汕头陈店\n【深圳】深圳坂田五和店、腾讯园区总部店",
            "纯文本门店查询", ["10001"],
        )
        expected = {
            "武汉首店": "武汉首店",
            "和平店": "和平店",
            "汕头": "汕头陈店",
            "深圳坂田五和": "深圳坂田五和店",
            "腾讯园区总部": "腾讯园区总部店",
            "万达广场": "武汉万达广场店",
        }
        for query, branch in expected.items():
            with self.subTest(query=query):
                result = self.store.resolve_deterministic("10001", query)
                self.assertEqual("stores", result["kind"])
                self.assertEqual("allow", result["decision"])
                self.assertTrue(result["reply"].startswith("可以用，根据“"))
                self.assertIn("”查询到可用门店：", result["reply"])
                self.assertIn(branch, result["reply"])

    def test_short_followup_inherits_previous_store_result(self):
        self.store.import_store_list(self.create_store_sheet(), "北京门店", ["10001"])
        first = self.store.resolve_deterministic("10001", "华熙五棵松店")
        self.assertEqual("stores", first["kind"])
        context = {
            "query": first["store_query"],
            "matches": first["store_matches"],
        }
        followup = self.store.resolve_deterministic(
            "10001", "可以用吗嘛", store_context=context
        )
        self.assertEqual("stores", followup["kind"])
        self.assertIn("华熙五棵松店可用", followup["reply"])

    def test_product_titles_are_never_returned_as_store_names(self):
        with self.store._connect() as conn:
            list_id = conn.execute(
                "INSERT INTO store_lists(name,source_file,store_count,created_at,updated_at) VALUES(?,?,?,?,?)",
                ("页面门店", "page", 2, self.store._now(), self.store._now()),
            ).lastrowid
            conn.execute("INSERT INTO product_store_lists(item_id,list_id) VALUES(?,?)", ("10001", list_id))
            for branch in ("武汉万象城店", "肥肥虾庄100元代金券 油焖大虾 武汉美食 搜索标签"):
                conn.execute(
                    "INSERT INTO stores(list_id,brand,branch,province,city,district,address,phone,business_hours,normalized) VALUES(?,?,?,?,?,?,?,?,?,?)",
                    (list_id, "肥肥虾庄", branch, "湖北", "武汉", "", "", "", "", normalize_text(branch + "武汉")),
                )
        result = self.store.search_store("10001", "武汉")
        reply = self.store.format_store_matches(result["matches"], "武汉")
        self.assertIn("武汉万象城店", reply)
        self.assertNotIn("搜索标签", reply)
        self.assertNotIn("100元代金券", reply)

    def test_txt_store_import_and_short_query(self):
        text = """【湖北省】
【武汉】武汉新荣天街店、武昌万象城店

【四川省】
【成都】成都IFS店、成都万象城店
"""
        preview = self.store.preview_store_text(text)
        self.assertEqual(4, preview["store_count"])
        self.assertEqual(2, preview["province_count"])
        imported = self.store.import_store_text(text, "文本门店", ["10001"])
        self.assertEqual(4, imported["store_count"])
        deterministic = self.store.resolve_deterministic("10001", "万象城")
        self.assertEqual("stores_clarify", deterministic["kind"])
        self.assertEqual("missing_area", deterministic["store_status"])
        self.assertIn("请补充城市或区县", deterministic["reply"])

    def test_fixed_replies_do_not_use_llm(self):
        greeting = self.store.resolve_deterministic("10001", "你好")
        self.assertEqual("greeting", greeting["kind"])
        self.assertNotIn("小刀", greeting["reply"])
        delivery = self.store.resolve_deterministic("10001", "付款后怎么发，怎么用")
        self.assertIn("券", delivery["reply"])
        self.assertIn("核销", delivery["reply"])
        code = self.store.resolve_deterministic("10001", "券码没发")
        self.assertIn("向上滑动", code["reply"])
        link = self.store.resolve_deterministic("10001", "链接打不开")
        self.assertIn("浏览器地址栏", link["reply"])
        image = self.store.resolve_deterministic("10001", "[图片]")
        self.assertEqual("暂时无法读取图片中的文字，请把完整门店名称或地址发成文字。", image["reply"])

    def test_media_with_dependent_text_never_guesses_store_or_price(self):
        for message in ("[图片]\n可以用么", "[语音]\n多少钱"):
            result = self.store.resolve_deterministic("10001", message)
            self.assertEqual("media", result["kind"])
            expected = (
                "暂时无法读取图片中的文字，请把完整门店名称或地址发成文字。"
                if "图片" in message else "暂时无法识别语音内容，请把问题发成文字。"
            )
            self.assertEqual(expected, result["reply"])

    def test_media_with_complete_text_can_use_the_text_only(self):
        self.store.import_store_list(self.create_multi_region_store_sheet(), "多地区门店", ["10001"])
        result = self.store.resolve_deterministic(
            "10001", "[图片]\n深圳龙岗万科里可以用吗"
        )
        self.assertEqual("stores", result["kind"])
        self.assertIn("龙岗万科里店", result["reply"])

    def test_first_reply_is_generated_and_can_be_manual(self):
        self.store.save_v2_product(
            "10001",
            "【自动发货】全国同仁四季椰子鸡代金券",
            """①100元代金券：54（每桌最多使用2张）
营业时间内可用，2026年9月25日不可用。
仅限堂食、不可用包间、不支持免费打包、不可拆台使用后拼桌、无需预约、高峰期需等位、不设找零、团购用户不可同时享受商家其他优惠。""",
            coupon_type="electronic_code",
        )
        self.store.save_ai_summary("10001", "摘要", {
            "summary": "摘要",
            "facts": {
                "品牌": "同仁四季",
                "使用时间": "营业时间内可用，2026年9月25日不可用",
                "使用规则": "仅限堂食、不可用包间、不支持免费打包、不可拆台使用后拼桌、无需预约、高峰期需等位、不设找零、团购用户不可同时享受商家其他优惠",
            },
            "time_rules": [],
        })
        product = self.store.get_v2_product("10001")
        first_reply = product["first_reply_text"]
        self.assertIn("您好，当前商品为同仁四季品牌电子券码。", first_reply)
        self.assertIn("【商品信息】\n同仁四季100元代金券：售价54元，发100元券1张", first_reply)
        self.assertIn("【使用规则】\n仅支持同面额代金券叠加，每次最多使用2张。", first_reply)
        self.assertIn("2026年9月25日不可用", first_reply)
        self.assertIn("除上述明确不可用日期外，适用门店营业时间内可用。", first_reply)
        self.assertIn("仅限堂食；不可用包间；不支持免费打包", first_reply)
        self.assertIn("【发券方式】\n付款后发电子券码，门店扫码核销。", first_reply)
        self.assertNotIn("请当天购买、当天使用", first_reply)
        self.assertNotIn("\n\n", first_reply)
        self.store.save_v2_product(
            "10001", product["title"], product["raw_text"],
            first_reply_enabled=True, first_reply_text="人工首次回复", first_reply_manual=True,
        )
        self.store.save_ai_summary("10001", "新摘要", {"summary": "新摘要", "facts": {}, "time_rules": []})
        self.assertEqual("人工首次回复", self.store.get_v2_product("10001")["first_reply_text"])

    def test_sku_meaning_and_rendering_are_preserved(self):
        raw = """①100元代金券：67.9（可叠加2张）
②200元：135.8（发两张100）
③300元：212（可叠加2张）
④400元：279.9（300+100）
⑤500元：349（可叠加2张）
不同面额代金券可叠加2张。"""
        self.store.save_v2_product(
            "10001", "匠熙小馆100/200/300/400/500代金券", raw,
            coupon_type="meituan",
        )
        product = self.store.get_v2_product("10001")
        first_reply = product["first_reply_text"]
        self.assertIn("100元代金券：售价67.9元，发100元券1张", first_reply)
        self.assertIn("200元代金券：售价135.8元，发100元券2张", first_reply)
        self.assertIn("400元代金券：售价279.9元，发300元券1张＋100元券1张", first_reply)
        self.assertIn("支持同面额或不同面额代金券叠加，每次最多使用2张", first_reply)
        self.assertNotIn("{'", first_reply)
        self.assertEqual(5, len(self.store.extract_product_options(product)))

    def test_authoritative_prices_override_listing_teaser_price(self):
        self.store.save_v2_product(
            "10001",
            "半秋山西餐厅60代100全国通用",
            """66.8元购100元代金券（工作日可用）
79.5元购100元代金券（全周通用）""",
        )
        with self.store._connect() as conn:
            conn.execute("UPDATE v2_products SET price='54' WHERE item_id='10001'")
        product = self.store.get_v2_product("10001")
        reply = self.store.price_reply(product)
        self.assertIn("66.8元", reply)
        self.assertIn("79.5元", reply)
        self.assertNotIn("售价54元", reply)

    def test_people_count_price_matches_real_package_semantically(self):
        self.store.save_v2_product(
            "10001", "经典自助电子券码",
            "通用单人经典自助（平日节假日全天通用）：售价165元\n"
            "节假日双人经典自助（节假日）：售价330元\n"
            "节假日三人经典自助（节假日）：售价495元",
        )
        result = self.store.resolve_deterministic("10001", "三个人多少钱")
        self.assertEqual("price", result["kind"])
        self.assertEqual("allow", result["decision"])
        self.assertIn("节假日三人经典自助", result["reply"])
        self.assertIn("495元", result["reply"])
        self.assertNotIn("没有“三个人”", result["reply"])
        self.assertNotIn("SKU", result["reply"])

    def test_weekend_price_returns_only_matching_real_coupon_option(self):
        self.store.save_v2_product(
            "10001", "半秋山100元代金券",
            "66.8元购100元代金券（工作日可用）\n"
            "79.5元购100元代金券（全周通用）",
        )
        result = self.store.resolve_deterministic("10001", "周末的多少钱")
        self.assertEqual("price", result["kind"])
        self.assertEqual("allow", result["decision"])
        self.assertIn("79.5元", result["reply"])
        self.assertNotIn("66.8元", result["reply"])
        self.assertIn("周末可用", result["reply"])

    def test_unspecified_coupon_availability_uses_today_weekend_tier(self):
        self.store.save_v2_product(
            "10001", "半秋山100元代金券",
            "66.8元购100元代金券（工作日可用）\n"
            "79.5元购100元代金券（周末可用）",
        )
        with patch.object(V2Store, "_current_day_type", return_value="weekend"):
            result = self.store.resolve_deterministic("10001", "100元代金券有吗")
        self.assertEqual("sku_availability", result["kind"])
        self.assertIn("今天是周末", result["reply"])
        self.assertIn("79.5元", result["reply"])
        self.assertNotIn("66.8元", result["reply"])

    def test_value_confirmation_uses_current_day_tier_before_correcting_price(self):
        self.store.save_v2_product(
            "10001", "半秋山100元代金券",
            "66.8元购100元代金券（工作日可用）\n"
            "79.5元购100元代金券（周末可用）",
        )
        with patch.object(V2Store, "_current_day_type", return_value="weekend"):
            result = self.store.resolve_deterministic("10001", "50抵100对吧")
        self.assertEqual("voucher_value", result["kind"])
        self.assertIn("今天是周末", result["reply"])
        self.assertIn("79.5元", result["reply"])
        self.assertNotIn("66.8元", result["reply"])

    def test_value_confirmation_honors_explicit_workday_and_holiday_tiers(self):
        self.store.save_v2_product(
            "10001", "半秋山100元代金券",
            "66.8元购100元代金券（工作日可用）\n"
            "79.5元购100元代金券（周末可用）\n"
            "88元购100元代金券（节假日可用）",
        )
        workday = self.store.resolve_deterministic("10001", "工作日66.8抵100对吧")
        self.assertIn("66.8元", workday["reply"])
        self.assertNotIn("79.5元", workday["reply"])
        self.assertNotIn("88元", workday["reply"])

        holiday = self.store.resolve_deterministic("10001", "节假日79.5抵100对吧")
        self.assertIn("88元", holiday["reply"])
        self.assertNotIn("66.8元", holiday["reply"])
        self.assertNotIn("79.5元", holiday["reply"])

    def test_today_value_confirmation_never_falls_back_to_workday_price(self):
        self.store.save_v2_product(
            "10001", "半秋山100元代金券",
            "66.8元购100元代金券（工作日可用）\n"
            "79.5元购100元代金券（周末可用）",
        )
        with patch.object(V2Store, "_current_day_type", return_value="weekend"):
            result = self.store.resolve_deterministic(
                "10001", "今天买的话能用79代100吗？",
            )
        self.assertEqual("voucher_value", result["kind"])
        self.assertIn("今天是周末", result["reply"])
        self.assertIn("79.5元", result["reply"])
        self.assertNotIn("66.8元", result["reply"])

        correction = self.store.resolve_deterministic(
            "10001", "我这边不是66呢", store_context=result["query_context_update"],
        )
        self.assertEqual("price", correction["kind"])
        self.assertIn("您说得对", correction["reply"])
        self.assertIn("周末档位", correction["reply"])
        self.assertIn("79.5元", correction["reply"])
        self.assertNotIn("66.8元", correction["reply"])

        stock = self.store.resolve_deterministic(
            "10001", "？是不是卖完了", store_context=correction["query_context_update"],
        )
        self.assertEqual("stock", stock["kind"])
        self.assertIn("没有卖完", stock["reply"])

    def test_generic_price_uses_today_tier_but_catalog_can_show_all_tiers(self):
        self.store.save_v2_product(
            "10001", "半秋山100元代金券",
            "66.8元购100元代金券（工作日可用）\n"
            "79.5元购100元代金券（周末可用）",
        )
        with patch.object(V2Store, "_current_day_type", return_value="weekend"):
            result = self.store.resolve_deterministic("10001", "多少钱")
        self.assertIn("今天是周末", result["reply"])
        self.assertIn("79.5元", result["reply"])
        self.assertNotIn("66.8元", result["reply"])
        full_catalog = self.store.price_reply(self.store.get_v2_product("10001"))
        self.assertIn("66.8元", full_catalog)
        self.assertIn("79.5元", full_catalog)

    def test_explicit_day_or_date_overrides_current_day_price(self):
        self.store.save_v2_product(
            "10001", "半秋山100元代金券",
            "66.8元购100元代金券（工作日可用）\n"
            "79.5元购100元代金券（周末可用）",
        )
        with patch.object(V2Store, "_current_day_type", return_value="weekday"):
            weekend = self.store.resolve_deterministic(
                "10001", "周末100元代金券可以直接买吗",
            )
            dated = self.store.resolve_deterministic("10001", "2026年9月6日多少钱")
        self.assertIn("79.5元", weekend["reply"])
        self.assertNotIn("66.8元", weekend["reply"])
        self.assertIn("79.5元", dated["reply"])
        self.assertNotIn("66.8元", dated["reply"])

    def test_meal_period_price_clarifies_people_then_uses_context(self):
        self.store.save_v2_product(
            "10001", "午餐自助电子券码",
            "单人经典自助（午餐可用）：售价165元\n"
            "双人经典自助（午餐可用）：售价330元\n"
            "三人经典自助（午餐可用）：售价495元",
        )
        first = self.store.resolve_deterministic("10001", "中午的多少钱")
        self.assertEqual("price", first["kind"])
        self.assertIn("几个人用餐", first["reply"])
        followup = self.store.resolve_deterministic(
            "10001", "三个人", store_context=first["query_context_update"],
        )
        self.assertEqual("price", followup["kind"])
        self.assertIn("三人经典自助", followup["reply"])
        self.assertIn("495元", followup["reply"])

    def test_people_count_availability_matches_named_option(self):
        self.store.save_v2_product(
            "10001", "经典自助电子券码",
            "单人经典自助：售价165元\n三人经典自助：售价495元",
        )
        result = self.store.resolve_deterministic("10001", "三个人有吗")
        self.assertEqual("sku_availability", result["kind"])
        self.assertIn("三人经典自助", result["reply"])
        self.assertNotIn("单人经典自助", result["reply"])

    def test_colloquial_people_day_and_meal_phrasings_use_the_same_filters(self):
        self.store.save_v2_product(
            "10001", "条件套餐",
            "工作日三人午餐套餐：售价288元\n"
            "周末三人午餐套餐：售价328元",
        )
        for message in ("3位工作日午饭多少钱", "我们仨周一中餐多钱"):
            with self.subTest(message=message):
                result = self.store.resolve_deterministic("10001", message)
                self.assertEqual("price", result["kind"])
                self.assertIn("288元", result["reply"])
                self.assertNotIn("328元", result["reply"])
        weekend = self.store.resolve_deterministic("10001", "三个人星期天午餐多少钱")
        self.assertIn("328元", weekend["reply"])
        corrected = self.store.resolve_deterministic(
            "10001", "不是周末，是工作日三个人午餐多少钱"
        )
        self.assertIn("288元", corrected["reply"])
        self.assertNotIn("328元", corrected["reply"])

    def test_store_correction_uses_the_last_positive_city(self):
        self.assertEqual("湖州", extract_store_query("不是杭州，是湖州"))

    def test_weekend_amount_query_calculates_with_only_the_weekend_price(self):
        self.store.save_v2_product(
            "10001", "半秋山100元代金券",
            "66.8元购100元代金券（工作日可用，最多叠加4张）\n"
            "79.5元购100元代金券（全周通用，最多叠加4张）",
        )
        result = self.store.resolve_deterministic("10001", "周末200多少钱")
        self.assertEqual("price", result["kind"])
        self.assertEqual("allow", result["decision"])
        self.assertIn("2张100元代金券", result["reply"])
        self.assertIn("159元", result["reply"])
        self.assertNotIn("133.6元", result["reply"])

        followup = self.store.resolve_deterministic(
            "10001", "200多少", store_context=result["query_context_update"],
        )
        self.assertIn("159元", followup["reply"])
        self.assertEqual("allow", followup["decision"])

    def test_numeric_price_shorthand_calculates_without_manual_review(self):
        self.store.save_v2_product(
            "10001", "同仁四季代金券",
            "100元代金券：售价54元，最多叠加3张",
        )
        three_hundred = self.store.resolve_deterministic("10001", "300的多少")
        self.assertEqual("allow", three_hundred["decision"])
        self.assertIn("3张100元代金券", three_hundred["reply"])
        self.assertIn("162元", three_hundred["reply"])
        amount = self.store.resolve_deterministic("10001", "388多少")
        self.assertEqual("allow", amount["decision"])
        self.assertIn("3张100元代金券", amount["reply"])
        self.assertIn("剩余88元到店自行支付", amount["reply"])
        self.assertNotIn("人工", amount["reply"])

    def test_specific_day_price_overrides_the_all_week_option(self):
        self.store.save_v2_product(
            "10001", "半秋山100元代金券",
            "66.8元购100元代金券（工作日可用）\n"
            "79.5元购100元代金券（全周通用）",
        )
        workday = self.store.resolve_deterministic("10001", "工作日多少钱")
        self.assertEqual("allow", workday["decision"])
        self.assertIn("66.8元", workday["reply"])
        self.assertNotIn("79.5元", workday["reply"])
        today = self.store.resolve_deterministic("10001", "今天多少钱")
        self.assertEqual("allow", today["decision"])
        self.assertNotIn("人工", today["reply"])

    def test_generic_price_uses_real_package_options(self):
        self.store.save_v2_product(
            "10001", "经典自助电子券码",
            "单人经典自助：售价165元\n双人经典自助：售价330元",
        )
        result = self.store.resolve_deterministic("10001", "多少钱")
        self.assertEqual("price", result["kind"])
        self.assertEqual("allow", result["decision"])
        self.assertIn("单人经典自助售价165元", result["reply"])
        self.assertIn("双人经典自助售价330元", result["reply"])
        self.assertNotIn("尚未确认", result["reply"])

    def test_missing_product_attribute_does_not_claim_a_manual_transfer(self):
        result = self.store.resolve_deterministic("10001", "有什么口味")
        self.assertEqual("allow", result["decision"])
        self.assertIn("暂未说明具体口味", result["reply"])
        self.assertNotIn("转人工", result["reply"])

    def test_delivery_method_does_not_confuse_takeaway_with_delivery(self):
        self.store.save_v2_product(
            "10001", "测试餐券", "仅限堂食，支持到店外带，不支持免费打包。"
        )
        result = self.store.resolve_deterministic("10001", "可以外卖配送吗")
        self.assertEqual("delivery_method", result["kind"])
        self.assertIn("仅确认支持到店外带", result["reply"])
        self.assertIn("外卖配送暂未确认", result["reply"])

    def test_missing_delivery_and_stacking_facts_do_not_fake_a_manual_transfer(self):
        self.store.save_v2_product("10001", "测试餐券", "仅限堂食。")
        delivery = self.store.resolve_deterministic("10001", "支持外卖配送吗")
        self.assertEqual("allow", delivery["decision"])
        self.assertIn("暂时无法准确确认", delivery["reply"])
        self.assertNotIn("人工", delivery["reply"])

        product = self.store.get_v2_product("10001")
        stacking = self.store.stacking_reply(product, "一次可以用几张")
        self.assertIn("暂时无法准确确认", stacking)
        self.assertNotIn("人工", stacking)

    def test_offline_product_unifies_every_buyer_consultation(self):
        with self.store._connect() as conn:
            conn.execute("UPDATE v2_products SET item_status='offline' WHERE item_id='10001'")
        ordinary = self.store.resolve_deterministic("10001", "现在能买吗")
        self.assertEqual("offline", ordinary["kind"])
        self.assertIn("已下架", ordinary["reply"])
        aftersale = self.store.resolve_deterministic("10001", "券码不能用怎么退款")
        self.assertEqual("offline", aftersale["kind"])
        self.assertEqual(ordinary["reply"], aftersale["reply"])

    def test_offline_purchase_order_is_not_swallowed_by_silence_rule(self):
        self.store.save_v2_product(
            "10001", "代买服务", "按实际需求人工报价", coupon_type="purchase_order",
        )
        with self.store._connect() as conn:
            conn.execute("UPDATE v2_products SET item_status='offline' WHERE item_id='10001'")
        result = self.store.resolve_deterministic("10001", "现在还能买吗")
        self.assertEqual("offline", result["kind"])
        self.assertEqual("allow", result["decision"])
        self.assertIn("已下架", result["reply"])

    def test_store_followup_uses_previous_matches(self):
        self.store.import_store_list(self.create_store_sheet(), "北京门店", ["10001"])
        first = self.store.resolve_deterministic("10001", "华熙五棵松店可以用吗")
        followup = self.store.resolve_deterministic(
            "10001", "这家店能用吗", store_context={
                "query": first["store_query"], "matches": first["store_matches"]
            }
        )
        self.assertEqual("stores", followup["kind"])
        self.assertIn("华熙五棵松店", followup["reply"])
        self.assertNotIn("请发送具体", followup["reply"])

    def test_refund_order_records_and_calculates_95_percent(self):
        record = self.store.upsert_refund_order({
            "order_id": "ORDER-1", "scope_id": "scope-1", "item_id": "10001",
            "paid_amount": "67.90", "reason": "买多了", "status": "待核实",
            "source": "闲鱼订单状态", "order_url": "https://www.goofish.com/order/ORDER-1",
        })
        self.assertEqual("64.51", record["suggested_amount"])
        self.assertEqual("待核实", record["status"])
        self.assertIn("ORDER-1", record["order_url"])
        updated = self.store.update_refund_order_status(record["id"], "处理中")
        self.assertEqual("处理中", updated["status"])

    def test_refund_list_excludes_text_only_consultations(self):
        self.store.upsert_refund_order({
            "scope_id": "scope-consult", "item_id": "10001",
            "reason": "可以退款吗", "source": "买家退款咨询",
        })
        self.assertFalse(any(row["scope_id"] == "scope-consult" for row in self.store.list_refund_orders()))

    def test_product_ai_can_be_switched_individually(self):
        updated = self.store.set_product_enabled("10001", False)
        self.assertFalse(updated["enabled"])
        updated = self.store.set_product_enabled("10001", True)
        self.assertTrue(updated["enabled"])

    def test_coupon_type_and_refund_rules_are_deterministic(self):
        self.store.save_v2_product(
            "10001", "测试券", "100元代金券：67.9",
            coupon_type="douyin",
        )
        platform = self.store.resolve_deterministic("10001", "这是美团还是抖音券")
        self.assertEqual("coupon_type", platform["kind"])
        self.assertIn("抖音卡券", platform["reply"])

        ordinary = self.store.resolve_deterministic(
            "10001", "买多了想退款，订单实付67.9元"
        )
        self.assertEqual("allow", ordinary["decision"])
        self.assertIn("非卡券质量问题", ordinary["reply"])
        self.assertIn("64.51元", ordinary["reply"])
        self.assertIn("当前商品规则", ordinary["reply"])

        quality = self.store.resolve_deterministic("10001", "券码不能用怎么退款")
        self.assertEqual("review", quality["decision"])
        self.assertIn("卡券质量问题", quality["reply"])

    def test_mixed_coupon_delivery_uses_sku_knowledge_without_fixed_platform(self):
        self.store.save_ai_summary("10001", "双渠道规格", {
            "facts": {},
            "common_rules": {},
            "sku_profiles": [
                {
                    "sku_key": "mt-100", "sku_name": "美团100元代金券",
                    "face_value": "100", "sale_price": "69",
                    "delivery_platform": "美团", "delivery_method": "付款后发送美团券领取信息",
                    "claim_method": "在美团账户领取", "redeem_method": "到店出示美团券核销",
                },
                {
                    "sku_key": "dy-200", "sku_name": "抖音200元代金券",
                    "face_value": "200", "sale_price": "138",
                    "delivery_platform": "抖音", "delivery_method": "付款后发送抖音券领取信息",
                    "claim_method": "在抖音账户领取", "redeem_method": "到店出示抖音券核销",
                },
            ],
            "time_rules": [],
        })
        product = self.store.save_v2_product(
            "10001", "测试双渠道代金券", "美团100元69元；抖音200元138元",
            coupon_type="mixed",
        )
        self.assertEqual("混合发卡", self.store.coupon_type_display(product))
        self.assertNotIn("当前商品为美团电子券", product["first_reply_text"])
        self.assertNotIn("当前商品为电子券码", product["first_reply_text"])
        self.assertIn("美团100元代金券：发券平台：美团", product["first_reply_text"])
        self.assertIn("抖音200元代金券：发券平台：抖音", product["first_reply_text"])

        all_channels = self.store.resolve_deterministic("10001", "这是美团还是抖音券")
        self.assertEqual("coupon_type", all_channels["kind"])
        self.assertIn("美团100元代金券", all_channels["reply"])
        self.assertIn("抖音200元代金券", all_channels["reply"])
        self.assertNotIn("当前商品发放的是混合发卡", all_channels["reply"])

        selected = self.store.resolve_deterministic("10001", "100元的是美团券吗")
        self.assertEqual("coupon_type", selected["kind"])
        self.assertIn("美团100元代金券", selected["reply"])
        self.assertNotIn("抖音200元代金券", selected["reply"])

    def test_mixed_coupon_missing_channel_asks_for_sku_without_guessing(self):
        product = self.store.save_v2_product(
            "10001", "测试代金券", "100元代金券售价69元", coupon_type="mixed",
        )
        self.assertNotIn("【发券方式】", product["first_reply_text"])
        result = self.store.resolve_deterministic("10001", "发的什么券")
        self.assertEqual("coupon_type", result["kind"])
        self.assertIn("暂未标明", result["reply"])
        self.assertNotIn("美团", result["reply"])
        self.assertNotIn("电子券码", result["reply"])

    def test_unredeemed_statement_clarifies_instead_of_inventing_expiry(self):
        result = self.store.resolve_deterministic("10001", "我没验")
        self.assertEqual("aftersale_clarify", result["kind"])
        self.assertIn("还没有核销", result["reply"])
        self.assertNotIn("过期", result["reply"])
        self.assertNotIn("不退不补", result["reply"])
        self.assertEqual("aftersale", result["query_context_update"]["intent"])

    def test_refund_request_without_order_state_asks_for_payment_and_reason(self):
        first = self.store.resolve_deterministic("10001", "我没验")
        second = self.store.resolve_deterministic(
            "10001", "给我退了吧", store_context=first["query_context_update"],
        )
        self.assertEqual("aftersale_clarify", second["kind"])
        self.assertIn("是否已经付款", second["reply"])
        self.assertIn("无法核销", second["reply"])
        self.assertNotIn("过期", second["reply"])

        punctuation = self.store.resolve_deterministic(
            "10001", "？", store_context=second["query_context_update"],
        )
        self.assertEqual("aftersale_clarify", punctuation["kind"])
        self.assertEqual(second["reply"], punctuation["reply"])
        self.assertNotIn("门店", punctuation["reply"])

    def test_order_payment_state_controls_refund_next_step(self):
        unpaid = self.store.resolve_deterministic(
            "10001", "给我退了吧", order_context={
                "status": "等待买家付款", "payment_state_verified": True,
            },
        )
        self.assertEqual("unpaid_cancel", unpaid["kind"])
        self.assertIn("尚未付款", unpaid["reply"])
        self.assertNotIn("手续费", unpaid["reply"])

        paid = self.store.resolve_deterministic(
            "10001", "给我退了吧", order_context={"status": "等待卖家发货"},
        )
        self.assertEqual("aftersale_clarify", paid["kind"])
        self.assertIn("退款原因", paid["source"])
        personal = self.store.resolve_deterministic(
            "10001", "买多了，个人原因", store_context=paid["query_context_update"],
        )
        self.assertEqual("refund_process", personal["kind"])
        self.assertIn("非卡券质量问题", personal["reply"])
        self.assertIn("95%", personal["reply"])

    def test_refund_never_uses_unidentified_cached_unpaid_card(self):
        runtime_cache = {"status": "等待买家付款", "observed_at": 1}
        result = self.store.resolve_deterministic(
            "10001", "还没用的可以退吗", order_context=runtime_cache,
        )
        self.assertNotEqual("unpaid_cancel", result["kind"])
        self.assertNotIn("尚未付款", result["reply"])

        paid = self.store.resolve_deterministic(
            "10001", "已经付款了，还没用，可以退吗",
            order_context={"status": "等待买家付款", "order_id": "old-order", "observed_at": 1},
        )
        self.assertNotEqual("unpaid_cancel", paid["kind"])
        self.assertNotIn("尚未付款", paid["reply"])

        fresh_but_unverified = self.store.resolve_deterministic(
            "10001", "能退吗", order_context={
                "status": "等待买家付款", "order_id": "ORDER-1",
                "observed_at": datetime.now().timestamp(),
            },
        )
        self.assertNotEqual("unpaid_cancel", fresh_but_unverified["kind"])
        self.assertNotIn("尚未付款", fresh_but_unverified["reply"])

    def test_aftersale_does_not_inherit_unpaid_state(self):
        result = self.store.resolve_deterministic(
            "10001", "能退吗", store_context={
                "intent": "aftersale", "aftersale_payment_state": "unpaid",
                "aftersale_stage": "payment_and_reason",
                "aftersale_prompt": "请确认是否已经付款。",
            },
        )
        self.assertNotEqual("unpaid_cancel", result["kind"])
        self.assertNotIn("尚未付款", result["reply"])

    def test_refund_policy_consultation_is_gentle_and_does_not_fake_an_order(self):
        result = self.store.resolve_deterministic("10001", "可以退款吗")
        self.assertEqual("refund_process", result["kind"])
        self.assertIn("尚未核销", result["reply"])
        self.assertIn("实际审核结果为准", result["reply"])
        self.assertNotIn("已经为您记录", result["reply"])

    def test_purchase_order_aftersale_language_also_stays_silent(self):
        self.store.save_v2_product(
            "10001", "代买服务", "付款后按订单说明领取", coupon_type="purchase_order",
        )
        greeting = self.store.resolve_deterministic("10001", "你好")
        self.assertEqual("purchase_order_other", greeting["kind"])
        unredeemed = self.store.resolve_deterministic("10001", "我没验")
        self.assertEqual("purchase_order_other", unredeemed["kind"])
        self.assertEqual("silent", unredeemed["decision"])
        refund = self.store.resolve_deterministic("10001", "给我退了吧")
        self.assertEqual("purchase_order_other", refund["kind"])
        self.assertEqual("silent", refund["decision"])

    def test_refund_status_uses_only_known_order_status(self):
        known = self.store.resolve_deterministic(
            "10001", "退款到哪了", order_context={"status": "退款申请"},
        )
        self.assertEqual("refund_status", known["kind"])
        self.assertIn("退款申请", known["reply"])
        unknown = self.store.resolve_deterministic("10001", "退款什么时候到账")
        self.assertEqual("refund_status", unknown["kind"])
        self.assertIn("无法从当前会话确认", unknown["reply"])
        self.assertNotIn("72小时", unknown["reply"])

    def test_expired_personal_case_uses_gentle_policy_wording(self):
        result = self.store.resolve_deterministic("10001", "忘了用，已经过期了")
        self.assertEqual("same_day_use", result["kind"])
        self.assertIn("通常无法办理", result["reply"])
        self.assertIn("敬请理解", result["reply"])
        self.assertNotIn("不退不补", result["reply"])

    def test_same_day_confirmation_uses_polite_paragraphs(self):
        self.store.save_v2_product(
            "10001", "测试电子券", "请当天购买、当天使用，过期不退不补。",
        )
        result = self.store.resolve_deterministic("10001", "不是现在买了就要用吧")
        self.assertEqual("same_day_use", result["kind"])
        self.assertEqual(3, len(result["reply"].split("\n\n")))
        self.assertIn("需要在购买当天使用", result["reply"])
        self.assertIn("还请您理解", result["reply"])
        self.assertNotIn("不退不补", result["reply"])

    def test_store_query_discards_pronouns_and_accepts_city(self):
        self.store.import_store_list(self.create_store_sheet(), "北京门店", ["10001"])
        city = self.store.resolve_deterministic("10001", "北京可以吗")
        self.assertEqual("stores", city["kind"])
        self.assertIn("北京", city["reply"])
        self.assertIn("可以用，根据“北京”查询到可用门店：", city["reply"])
        unclear = self.store.resolve_deterministic("10001", "该门店不可用吗")
        self.assertEqual("stores_clarify", unclear["kind"])
        self.assertNotIn("该 不", unclear["reply"])

    def test_old_automatic_first_reply_is_migrated_but_manual_reply_is_preserved(self):
        with self.store._connect() as conn:
            conn.execute(
                """UPDATE v2_products SET first_reply_text='旧版自动回复',
                   first_reply_manual=0,first_reply_template_version=1 WHERE item_id='10001'"""
            )
        migrated = V2Store(self.store.db_path).get_v2_product("10001")
        self.assertIn("【商品信息】", migrated["first_reply_text"])
        with self.store._connect() as conn:
            conn.execute(
                """UPDATE v2_products SET first_reply_text='我手动写的回复',
                   first_reply_manual=1,first_reply_template_version=1 WHERE item_id='10001'"""
            )
        preserved = V2Store(self.store.db_path).get_v2_product("10001")
        self.assertEqual("我手动写的回复", preserved["first_reply_text"])

    def test_time_rule_uses_china_time_and_current_product_only(self):
        self.store.replace_time_rules("10001", [
            {
                "label": "工作日午餐可用",
                "day_type": "weekday",
                "start_time": "11:00",
                "end_time": "14:00",
                "allowed": True,
                "reply": "这张券工作日午餐时段可用。",
            },
            {
                "label": "周末不可用",
                "day_type": "weekend",
                "start_time": "00:00",
                "end_time": "23:59",
                "allowed": False,
                "reply": "这张券周末不可用。",
                "next_hint": "可在下个工作日使用。",
            },
        ])
        weekday = self.store.evaluate_time("10001", "2026-09-01T12:00:00")
        self.assertEqual(weekday["status"], "allowed")
        self.assertEqual(weekday["meal_period"], "午餐")
        weekend = self.store.evaluate_time("10001", "2026-09-05T12:00:00")
        self.assertEqual(weekend["status"], "blocked")
        deterministic = self.store.resolve_deterministic("10001", "现在这个券能用吗")
        if deterministic:
            self.assertNotIn("其他商品", deterministic["reply"])

    def test_synced_source_never_overwrites_manual_knowledge(self):
        self.store.upsert_synced_product({
            "item_id": "20002",
            "title": "同步商品",
            "platform_summary": "闲鱼文案第一版",
            "thumbnail_url": "https://img.alicdn.com/a.jpg",
            "image_urls": ["https://img.alicdn.com/a.jpg"],
            "price": "88",
        })
        self.store.save_synced_summary("20002", "AI初始知识", {"summary": "AI初始知识", "time_rules": []})
        initial = self.store.get_v2_product("20002")
        self.assertEqual("闲鱼文案第一版", initial["raw_text"])
        self.assertTrue(initial["ai_summary"])
        self.assertIn("适用门店营业时间内可用", initial["ai_summary"])
        self.store.save_v2_product("20002", "同步商品", "我人工修改后的权威知识")
        self.store.upsert_synced_product({
            "item_id": "20002",
            "title": "同步商品新标题",
            "platform_summary": "闲鱼文案第二版",
            "image_urls": [],
            "price": "90",
        })
        self.store.save_synced_summary("20002", "新的AI初始知识", {"summary": "新的AI初始知识", "time_rules": []})
        updated = self.store.get_v2_product("20002")
        self.assertEqual("我人工修改后的权威知识", updated["raw_text"])
        self.assertTrue(updated["ai_summary"])
        self.assertIn("适用门店营业时间内可用", updated["ai_summary"])
        self.assertEqual("source_updated", updated["sync_status"])

    def test_store_lists_are_managed_inside_product(self):
        first = self.store.import_store_list(self.create_store_sheet(), "门店表A", ["10001"])
        second = self.store.import_store_list(self.create_store_sheet(), "门店表B", [])
        self.store.set_product_store_lists("10001", [second["id"]])
        bound = self.store.bound_store_lists("10001")
        self.assertEqual([second["id"]], [item["id"] for item in bound])
        self.assertNotEqual(first["id"], second["id"])

    def test_product_image_asset_match_is_scoped_and_has_cooldown(self):
        image_path = os.path.join(self.temp.name, "menu.png")
        with open(image_path, "wb") as handle:
            handle.write(b"not-real-image-needed-for-store-test")
        asset = self.store.save_image_asset({
            "item_id": "10001",
            "name": "双人套餐图",
            "purpose": "买家咨询套餐内容时发送",
            "trigger_words": ["套餐图", "套餐有什么"],
            "reply_text": "给您发一下双人套餐图。",
            "file_path": image_path,
            "enabled": True,
        })
        matched = self.store.resolve_image_asset("10001", "套餐有什么，发个套餐图看看", "scope-a")
        self.assertEqual("allow", matched["status"])
        self.assertEqual(asset["id"], matched["asset"]["id"])
        self.assertIsNone(self.store.resolve_image_asset("other-item", "发个套餐图", "scope-a"))
        self.store.record_image_send(
            asset["id"], "10001", "scope-a", "chat-a", "buyer-a", "sent"
        )
        cooldown = self.store.resolve_image_asset("10001", "再发一次套餐图", "scope-a")
        self.assertEqual("cooldown", cooldown["status"])

    def test_promotion_mode_routes_purchase_intents_to_link_and_qr(self):
        self.store.save_v2_product(
            "10001", "测试推广商品",
            "推广下单链接：https://example.invalid/promo\n周末可用。",
            coupon_type="promotion",
        )
        image_path = os.path.join(self.temp.name, "promotion-qr.png")
        with open(image_path, "wb") as handle:
            handle.write(b"promotion-qr")
        asset = self.store.save_image_asset({
            "item_id": "10001",
            "name": "推广购买二维码",
            "purpose": "promotion_qr",
            "trigger_words": ["推广购买二维码"],
            "reply_text": "",
            "file_path": image_path,
            "enabled": True,
        })
        product = self.store.refresh_first_reply("10001", force=True)
        self.assertIn("请不要在闲鱼页面下单或付款", product["first_reply_text"])
        self.assertIn("https://example.invalid/promo", product["first_reply_text"])
        self.assertIn(f"{{$图片:{asset['id']}}}", product["first_reply_text"])

        for question in (
            "这个多少钱", "怎么买", "把二维码发我", "发我个购买链接", "二维码呢",
            "我去哪扫", "可以在闲鱼直接下单吗", "能拍吗", "想买三张", "拍哪个",
            "购买链接打不开",
        ):
            result = self.store.resolve_deterministic("10001", question)
            self.assertIsNotNone(result, question)
            self.assertEqual("promotion_purchase", result["kind"], question)
            self.assertEqual(asset["id"], result["image_asset_id"])
            self.assertIn("https://example.invalid/promo", result["reply"])
            self.assertIn("不要在闲鱼页面下单或付款", result["reply"])

        # The reserved QR must not leak into ordinary keyword-image matching.
        self.assertIsNone(self.store.resolve_image_asset(
            "10001", "推广购买二维码", "promotion-scope",
        ))

    def test_promotion_mode_never_invents_a_missing_purchase_entry(self):
        self.store.save_v2_product(
            "10001", "未配置入口的推广商品", "仅限堂食。", coupon_type="promotion",
        )
        result = self.store.resolve_deterministic("10001", "购买链接发我")
        self.assertEqual("promotion_purchase", result["kind"])
        self.assertNotIn("http", result["reply"])
        self.assertIn("尚未配置", result["reply"])

        aftersale = self.store.resolve_deterministic(
            "10001", "已经付款但不想要了，怎么退款",
            order_context={"status": "已付款"},
        )
        self.assertNotEqual("promotion_purchase", aftersale["kind"])

    def test_keyword_rule_can_send_text_without_an_image(self):
        asset = self.store.save_image_asset({
            "item_id": "10001",
            "name": "停车说明",
            "purpose": "买家咨询停车时发送",
            "trigger_words": ["停车", "停车费"],
            "reply_text": "商场提供停车场。{$分段符}停车收费以现场公示为准。",
            "file_path": "",
            "enabled": True,
        })
        self.assertEqual("", asset["file_path"])
        matched = self.store.resolve_image_asset("10001", "停车费怎么算", "scope-text")
        self.assertEqual("allow", matched["status"])
        self.assertEqual(asset["id"], matched["asset"]["id"])

    def test_multiple_image_hits_require_review(self):
        for index in (1, 2):
            image_path = os.path.join(self.temp.name, f"menu-{index}.jpg")
            with open(image_path, "wb") as handle:
                handle.write(b"test")
            self.store.save_image_asset({
                "item_id": "10001",
                "name": f"套餐图{index}",
                "trigger_words": ["菜单图"],
                "file_path": image_path,
            })
        matched = self.store.resolve_image_asset("10001", "发一下菜单图", "scope-b")
        self.assertEqual("review", matched["status"])

    def test_specific_coupon_price_only_returns_requested_denomination(self):
        self.store.save_v2_product(
            "10001", "匠熙小馆代金券",
            "100元代金券：67.9\n200元代金券：135.8（发两张100）\n11元代金券：0元",
        )
        product = self.store.get_v2_product("10001")
        reply = self.store.price_reply(product, "100元代金券多少钱")
        self.assertIn("100元代金券", reply)
        self.assertIn("67.9元", reply)
        self.assertNotIn("200元代金券", reply)
        self.assertNotIn("11元代金券", reply)

        short_reply = self.store.price_reply(product, "200多少钱")
        self.assertIn("200元代金券", short_reply)
        self.assertIn("135.8元", short_reply)
        self.assertNotIn("100元代金券", short_reply)

    def test_stack_defaults_to_same_denomination_and_hides_sku_time(self):
        self.store.save_v2_product(
            "10001", "测试代金券",
            "100元代金券（08:00-次日05:00）：84.8，最多叠加4张",
        )
        reply = self.store.get_v2_product("10001")["first_reply_text"]
        self.assertIn("仅支持同面额代金券叠加，每次最多使用4张", reply)
        self.assertNotIn("08:00-次日05:00）：", reply)

    def test_city_label_never_inherits_previous_province(self):
        parsed = self.store.parse_store_text(
            "【重庆市】\n【重庆】沙坪坝店\n【武汉】武昌万象城店\n【苏州】狮山天街店"
        )
        by_city = {row["city"]: row["province"] for row in parsed["records"]}
        self.assertEqual("湖北省", by_city["武汉"])
        self.assertEqual("江苏省", by_city["苏州"])

    def test_store_import_can_replace_product_bindings(self):
        first = self.store.import_store_text("【湖北省】\n【武汉】A店", "旧表", ["10001"])
        second = self.store.import_store_text(
            "【四川省】\n【成都】B店", "新表", ["10001"],
            replace_item_bindings=True,
        )
        product = self.store.get_v2_product("10001")
        self.assertEqual([second["id"]], [entry["id"] for entry in product["store_lists"]])
        self.assertNotEqual(first["id"], second["id"])

    def test_zero_price_placeholder_coupon_is_not_rendered(self):
        self.store.save_v2_product(
            "10001", "肥肥虾庄代金券",
            "8元代金券：售价0元\n100元代金券：售价69元",
        )
        product = self.store.get_v2_product("10001")
        rendered = product["first_reply_text"] + "\n" + self.store.price_reply(product)
        self.assertNotIn("8元代金券", rendered)
        self.assertIn("100元代金券", rendered)

    def test_quantity_skus_are_collapsed_to_one_denomination(self):
        self.store.save_v2_product(
            "10001", "肥肥虾庄代金券",
            "100元代金券：84.8（最多叠加3张）\n100元代金券×2：169.6\n100元代金券×3：254.4",
        )
        product = self.store.get_v2_product("10001")
        options = self.store.extract_product_options(product)
        self.assertEqual(1, len(options))
        rendered = product["first_reply_text"]
        self.assertIn("售价84.8元", rendered)
        self.assertIn("最多使用3张", rendered)
        self.assertNotIn("169.6元", rendered)
        self.assertNotIn("254.4元", rendered)

    def test_quantity_price_and_stacking_use_effective_knowledge(self):
        self.store.save_v2_product(
            "10001", "鱼酷烤鱼",
            "鱼酷活鱼烤鱼一条鱼2.6斤：118元\n100元代金券72.5元，最多叠加2张\n300元代金券240元",
        )
        product = self.store.get_v2_product("10001")
        self.assertIn("2条共236元", self.store.quantity_price_reply(product, "两条多少钱"))
        reply = self.store.stacking_reply(product, "100和300可以一起用吗")
        self.assertIn("不同面额", reply)
        self.assertIn("不能一起使用", reply)

    def test_coupon_quantity_phrasings_calculate_real_total(self):
        self.store.save_v2_product(
            "10001", "同仁四季代金券", "100元代金券：售价54元，最多叠加3张"
        )
        for message in (
            "两张多少钱", "2张多少钱", "买两张多少", "要两张多少钱",
            "来两张", "两张一共多少", "100的两张多少钱", "两张100多少钱",
            "100券买两张多少钱", "两张啥价", "2x100多少钱", "100×2多钱",
        ):
            with self.subTest(message=message):
                result = self.store.resolve_deterministic("10001", message)
                self.assertEqual("price", result["kind"])
                self.assertIn("2张100元代金券", result["reply"])
                self.assertIn("108元", result["reply"])
                self.assertIn("抵扣200元", result["reply"])
                self.assertNotIn("SKU", result["reply"].upper())

    def test_coupon_quantity_with_multiple_denominations_asks_once_then_uses_context(self):
        self.store.save_v2_product(
            "10001", "测试代金券",
            "100元代金券：售价54元，最多叠加3张\n200元代金券：售价108元，最多叠加2张",
        )
        first = self.store.resolve_deterministic("10001", "两张多少钱")
        self.assertIn("哪种面额", first["reply"])
        self.assertEqual("target_amount", first["query_context_update"]["price_filters"]["awaiting"])
        followup = self.store.resolve_deterministic(
            "10001", "100", store_context=first["query_context_update"],
        )
        self.assertIn("2张100元代金券", followup["reply"])
        self.assertIn("108元", followup["reply"])
        self.assertNotIn("哪种面额", followup["reply"])

    def test_coupon_quantity_without_date_uses_current_day_price(self):
        self.store.save_v2_product(
            "10001", "半秋山代金券",
            "66.8元购100元代金券（工作日可用，最多叠加4张）\n"
            "79.5元购100元代金券（周末可用，最多叠加4张）",
        )
        with patch.object(V2Store, "_current_day_type", return_value="weekend"):
            result = self.store.resolve_deterministic("10001", "两张100多少钱")
        self.assertIn("159元", result["reply"])
        self.assertNotIn("133.6元", result["reply"])

    def test_quantity_over_stack_limit_does_not_claim_combined_redemption(self):
        self.store.save_v2_product(
            "10001", "同仁四季椰子鸡代金券",
            "100元代金券：售价55.8元，最多使用2张\n"
            "300元代金券：售价204.8元，最多使用1张",
        )
        result = self.store.resolve_deterministic("10001", "你好，买2张300券")
        self.assertEqual("price", result["kind"])
        self.assertIn("购买2张300元代金券共409.6元", result["reply"])
        self.assertIn("每张可抵扣300元", result["reply"])
        self.assertIn("每次最多使用1张", result["reply"])
        self.assertIn("2张不能在同一次消费中全部使用", result["reply"])
        self.assertNotIn("可抵扣600元", result["reply"])

    def test_composed_sku_limit_counts_delivered_coupon_quantity(self):
        self.store.save_v2_product(
            "10001", "组合代金券",
            "200元代金券：售价108元，发2张100元券，最多使用2张",
        )
        result = self.store.resolve_deterministic("10001", "买2张200券")
        self.assertEqual("price", result["kind"])
        self.assertIn("共216元", result["reply"])
        self.assertIn("每份该规格可抵扣200元", result["reply"])
        self.assertIn("本次购买所得的4张不能在同一次消费中全部使用", result["reply"])
        self.assertNotIn("可抵扣400元", result["reply"])

    def test_fish_single_item_phrasings_match_the_real_package(self):
        self.store.save_v2_product(
            "10001", "鱼酷烤鱼2-3人餐",
            "鱼酷活鱼烤鱼单条鱼2.6斤：118元（不含米饭和配菜）",
        )
        for message in ("一条鱼多少钱", "单条鱼多钱", "整条怎么卖", "一整条鱼什么价"):
            with self.subTest(message=message):
                result = self.store.resolve_deterministic("10001", message)
                self.assertEqual("price", result["kind"])
                self.assertIn("118元", result["reply"])
                self.assertIn("不含米饭和配菜", result["reply"])
                self.assertNotIn("没有", result["reply"])

    def test_package_numeric_amount_never_returns_an_unrelated_package_price(self):
        self.store.save_v2_product(
            "10001", "鱼酷烤鱼2-3人餐",
            "鱼酷活鱼烤鱼套餐：售价118元（单条鱼2.6斤，不含米饭配菜）",
        )
        option_query = self.store.resolve_deterministic("10001", "200的多少钱")
        self.assertIn("没有200元这一商品选项", option_query["reply"])
        self.assertNotIn("售价118元", option_query["reply"])
        bill_query = self.store.resolve_deterministic("10001", "消费200元怎么买")
        self.assertIn("按套餐售卖", bill_query["reply"])
        self.assertNotIn("2张", bill_query["reply"])

    def test_bare_people_count_fills_pending_question(self):
        self.store.save_v2_product(
            "10001", "午餐自助电子券码",
            "单人经典自助（午餐可用）：售价165元\n"
            "双人经典自助（午餐可用）：售价330元\n"
            "三人经典自助（午餐可用）：售价495元",
        )
        first = self.store.resolve_deterministic("10001", "中午多少钱")
        self.assertIn("几个人用餐", first["reply"])
        self.assertEqual("people_count", first["query_context_update"]["price_filters"]["awaiting"])
        for answer in ("3", "三", "3人", "三位", "我们3个"):
            with self.subTest(answer=answer):
                followup = self.store.resolve_deterministic(
                    "10001", answer, store_context=first["query_context_update"],
                )
                self.assertIn("三人经典自助", followup["reply"])
                self.assertIn("495元", followup["reply"])
                self.assertNotIn("几个人用餐", followup["reply"])

    def test_date_and_meal_without_people_asks_only_for_people_count(self):
        self.store.save_v2_product(
            "10001", "分时自助电子券码",
            "工作日单人经典自助（午餐可用）：售价159元\n"
            "工作日双人经典自助（午餐可用）：售价299元",
        )

        result = self.store.resolve_deterministic("10001", "明天中午")

        self.assertIn("几个人用餐", result["reply"])
        self.assertNotIn("请补充具体想查询的商品", result["reply"])
        self.assertEqual("people_count", result["query_context_update"]["price_filters"]["awaiting"])

    def test_people_and_date_followups_fill_the_requested_slot_in_order(self):
        self.store.save_v2_product(
            "10001", "分时自助电子券码",
            "工作日三人午餐自助：售价288元\n周末三人午餐自助：售价328元",
        )
        with patch.object(V2Store, "_current_day_type", return_value="weekend"):
            first = self.store.resolve_deterministic("10001", "三个人中午多少钱")
        self.assertIn("328元", first["reply"])
        self.assertNotIn("288元", first["reply"])
        weekend = self.store.resolve_deterministic(
            "10001", "礼拜天", store_context=first["query_context_update"],
        )
        self.assertIn("328元", weekend["reply"])
        self.assertNotIn("288元", weekend["reply"])

    def test_price_and_availability_colloquialisms_use_real_denominations(self):
        self.store.save_v2_product(
            "10001", "测试代金券", "100元代金券：售价54元\n300元代金券：售价155元"
        )
        for message in ("300啥价", "300怎么收", "300代多少", "300券几多钱"):
            with self.subTest(message=message):
                result = self.store.resolve_deterministic("10001", message)
                self.assertIn("300元代金券", result["reply"])
                self.assertIn("155元", result["reply"])
        for message in ("300还有吗", "300有货吗", "300能拍吗", "能不能买300券", "卖不卖300券"):
            with self.subTest(message=message):
                result = self.store.resolve_deterministic("10001", message)
                self.assertEqual("sku_availability", result["kind"])
                self.assertIn("300元代金券", result["reply"])
                self.assertIn("155元", result["reply"])

    def test_bare_number_without_pending_context_is_not_guessed_as_money(self):
        self.assertIsNone(self.store.conditional_sale_reply(
            self.store.get_v2_product("10001"), "3", None,
        ))

    def test_city_query_does_not_match_brand_or_other_cities(self):
        imported = self.store.import_store_text(
            "【广东省】\n【深圳】南昌品牌深圳店\n【江西省】\n【南昌】万寿宫店、北京东路店",
            "南昌品牌全国门店", ["10001"],
        )
        result = self.store.search_store("10001", "南昌")
        self.assertEqual("available", result["status"])
        self.assertEqual({"南昌"}, {row["city"] for row in result["matches"]})
        self.assertEqual(2, len(result["matches"]))

    def test_source_sync_is_staged_until_explicit_approval(self):
        self.store.save_v2_product("10001", "测试券", "人工知识", first_reply_text="人工首次回复", first_reply_manual=True)
        old = self.store.import_store_text("【湖北省】\n【武汉】旧店", "旧门店", ["10001"])
        self.store.stage_source_update(
            "10001", "新版知识", {"stores": [{"province": "四川省", "city": "成都", "branch": "新店"}]},
            "【四川省】\n【成都】新店",
        )
        staged = self.store.get_v2_product("10001")
        self.assertEqual("人工知识", staged["raw_text"])
        self.assertEqual("人工首次回复", staged["first_reply_text"])
        self.assertEqual([old["id"]], [row["id"] for row in staged["store_lists"]])
        self.store.apply_source_update("10001", ["knowledge"])
        applied = self.store.get_v2_product("10001")
        self.assertTrue(applied["raw_text"].startswith("新版知识"))
        self.assertIn("四川省", applied["raw_text"])
        self.assertIn("成都", applied["raw_text"])
        self.assertIn("新店", applied["raw_text"])
        self.assertEqual("人工首次回复", applied["first_reply_text"])
        self.assertEqual([old["id"]], [row["id"] for row in applied["store_lists"]])

    def test_platform_store_list_can_replace_and_delete(self):
        structured = {"stores": [
            {"province": "湖北省", "city": "武汉", "branch": "武昌万象城店"},
            {"province": "湖北省", "city": "武汉", "branch": "武商梦时代店"},
        ]}
        result = self.store.sync_platform_store_list("10001", "测试商品", "", structured)
        self.assertEqual(2, result["store_count"])
        self.assertEqual(2, len(self.store.search_store("10001", "武汉")["matches"]))
        deleted = self.store.delete_store_list(result["id"])
        self.assertEqual(1, deleted["affected_products"])

    def test_deleted_online_product_can_be_ignored_on_sync(self):
        deleted = self.store.delete_v2_product("10001", True)
        self.assertTrue(deleted["ignored"])
        self.assertTrue(self.store.is_product_ignored("10001"))
        self.assertIsNone(self.store.get_v2_product("10001"))

    def test_calculates_missing_denomination_from_real_sku(self):
        self.store.save_v2_product(
            "10001", "同仁四季代金券",
            "100元代金券：售价54元，最多叠加3张",
        )
        product = self.store.get_v2_product("10001")
        self.assertIn("没有200元代金券", self.store.price_reply(product, "200多少钱"))
        self.assertIn("3张100元代金券", self.store.resolve_deterministic("10001", "300怎么买")["reply"])

    def test_two_same_face_coupons_are_calculated_without_an_explicit_limit(self):
        self.store.save_v2_product(
            "10001", "同仁四季代金券", "100元代金券：售价54元"
        )
        result = self.store.resolve_deterministic("10001", "200多少钱")
        self.assertEqual("price", result["kind"])
        self.assertEqual("allow", result["decision"])
        self.assertIn("2张100元代金券", result["reply"])
        self.assertIn("共支付108元", result["reply"])
        self.assertIn("可抵扣200元", result["reply"])

    def test_amount_over_stack_limit_requires_store_payment(self):
        self.store.save_v2_product(
            "10001", "同仁四季代金券",
            "100元代金券：售价54元，最多叠加5张",
        )
        result = self.store.resolve_deterministic("10001", "600怎么办")
        self.assertIn("最多使用5张", result["reply"])
        self.assertIn("剩余100元", result["reply"])
        self.assertIn("到店自行支付", result["reply"])

    def test_quality_failure_without_refund_intent_uses_troubleshooting(self):
        result = self.store.resolve_deterministic("10001", "买了不能用怎么办")
        self.assertEqual("coupon_troubleshooting", result["kind"])
        self.assertNotIn("仅退款", result["reply"])

    def test_historical_version_can_be_restored(self):
        self.store.save_v2_product("10001", "测试券", "旧知识", note="旧版本")
        old = self.store.list_versions("10001")[0]
        self.store.save_v2_product("10001", "测试券", "新知识", note="新版本")
        restored = self.store.restore_version("10001", old["id"])
        self.assertEqual("旧知识", restored["raw_text"])
        self.assertIn("启用历史版本", self.store.list_versions("10001")[0]["note"])

    def test_composed_coupon_delivery_is_preserved_from_authoritative_text(self):
        self.store.save_v2_product(
            "10001", "组合代金券",
            "100元代金券：售价65元（最多叠加3张）\n"
            "125元代金券：售价74元（最多叠加4张）\n"
            "250元代金券：售价148元，发125元券2张\n"
            "300元代金券：售价195元，发100元券3张\n"
            "375元代金券：售价222元，发125元券3张\n"
            "500元代金券：售价296元，发125元券4张",
        )
        product = self.store.get_v2_product("10001")
        summary = self.store.build_knowledge_summary(product)
        self.assertIn("250元代金券：售价148元，发125元券2张", summary)
        self.assertIn("300元代金券：售价195元，发100元券3张", summary)
        self.assertIn("375元代金券：售价222元，发125元券3张", summary)
        self.assertIn("500元代金券：售价296元，发125元券4张", summary)

    def test_each_sku_keeps_its_own_stack_limit(self):
        self.store.save_v2_product(
            "10001", "多面额代金券",
            "100元代金券：售价65元（最多叠加3张）\n"
            "125元代金券：售价74元（最多叠加4张）",
        )
        product = self.store.get_v2_product("10001")
        summary = self.store.build_knowledge_summary(product)
        self.assertIn("100元代金券最多使用3张", summary)
        self.assertIn("125元代金券最多使用4张", summary)

        mixed_reply = self.store.stacking_reply(product, "100元和125元可以叠加吗")
        self.assertIn("不同面额", mixed_reply)
        self.assertIn("不能一起使用", mixed_reply)
        self.assertIn("100元代金券最多使用3张", mixed_reply)
        self.assertIn("125元代金券最多使用4张", mixed_reply)

        single_reply = self.store.stacking_reply(product, "125元可以用几张")
        self.assertIn("125元代金券", single_reply)
        self.assertIn("最多使用4张", single_reply)

    def test_purchase_at_store_question_is_not_store_search(self):
        result = self.store.resolve_deterministic("10001", "这个要去店里再买吗")
        self.assertEqual("purchase_flow", result["kind"])
        self.assertIn("当前商品页面", result["reply"])
        self.assertNotIn("未查询到可用门店", result["reply"])

    def test_unknown_current_time_does_not_invent_business_hours(self):
        result = self.store.resolve_deterministic("10001", "现在能用吗")
        self.assertEqual("time", result["kind"])
        self.assertEqual("allow", result["decision"])
        self.assertEqual("门店营业时段可用就可以了。", result["reply"])
        self.assertNotIn("人工", result["reply"])

    def test_hoarding_is_governed_by_aftersale_policy(self):
        result = self.store.resolve_deterministic("10001", "囤一年可以吗")
        self.assertEqual("same_day_use", result["kind"])
        self.assertIn("当天", result["reply"])
        self.assertNotIn("无有效期限制", result["reply"])

    def test_transfer_reply_includes_redemption_steps(self):
        result = self.store.resolve_deterministic("10001", "美团券可以转赠吗")
        self.assertEqual("transfer_usage", result["kind"])
        self.assertIn("不支持转赠", result["reply"])
        self.assertIn("出示券码核销", result["reply"])

    def test_all_store_question_requires_specific_store_lookup(self):
        self.store.save_v2_product("10001", "鱼酷烤鱼2-3人餐", "2.6斤烤鱼：118元")
        result = self.store.resolve_deterministic("10001", "团购券所有店铺通用吗")
        self.assertEqual("stores_scope", result["kind"])
        self.assertIn("鱼酷烤鱼券需在指定门店使用", result["reply"])
        self.assertIn("发送城市或店面名称", result["reply"])

    def test_compact_coverage_question_does_not_go_to_manual_review(self):
        for message in ("全国通用？", "全国能用吗", "通用吗", "都能用吗"):
            with self.subTest(message=message):
                result = self.store.resolve_deterministic("10001", message)
                self.assertEqual("stores_scope", result["kind"])
                self.assertEqual("allow", result["decision"])
                self.assertIn("指定门店", result["reply"])
                self.assertNotIn("全国通用", result["reply"])
                self.assertNotIn("人工", result["reply"])

    def test_consumption_amount_recommends_closest_not_over_target(self):
        self.store.save_v2_product(
            "10001", "同仁四季代金券",
            "100元代金券：售价54元，最多叠加2张\n代金券除酒水饮料、特价菜外全场通用",
        )
        result = self.store.resolve_deterministic("10001", "298怎么拍")
        self.assertEqual("consumption_plan", result["kind"])
        self.assertIn("2张100元代金券", result["reply"])
        self.assertIn("共支付108元", result["reply"])
        self.assertIn("剩余98元", result["reply"])
        self.assertNotIn("全场通用", result["reply"])

    def test_consumption_amount_with_spoken_prefix_returns_coupon_plan(self):
        self.store.save_v2_product(
            "10001", "测试代金券",
            "100元代金券：售价54元，最多叠加2张\n代金券全场通用",
        )
        result = self.store.resolve_deterministic("10001", "吃了260怎么买")
        self.assertEqual("consumption_plan", result["kind"])
        self.assertIn("2张100元代金券", result["reply"])
        self.assertIn("共支付108元", result["reply"])
        self.assertIn("剩余60元", result["reply"])

    def test_coupon_catalog_uses_only_real_skus(self):
        self.store.save_v2_product(
            "10001", "同仁四季代金券",
            "100元代金券：售价54元，最多叠加2张",
        )
        result = self.store.resolve_deterministic("10001", "多少代多少")
        self.assertEqual("coupon_catalog", result["kind"])
        self.assertIn("100元代金券", result["reply"])
        self.assertNotIn("300元代金券", result["reply"])

    def test_package_without_coupon_sku_does_not_invent_coupon(self):
        self.store.save_v2_product("10001", "鱼酷烤鱼2-3人餐", "2.6斤烤鱼套餐：118元")
        result = self.store.resolve_deterministic("10001", "多少代多少")
        self.assertIn("没有代金券选项", result["reply"])
        self.assertIn("鱼酷烤鱼2-3人餐", result["reply"])

    def test_bare_coupon_question_on_package_does_not_hallucinate_skus(self):
        self.store.save_v2_product(
            "10001", "鱼酷烤鱼2-3人餐",
            "鱼酷活鱼烤鱼单条鱼2.6斤：118元（不含米饭配菜）",
        )
        result = self.store.resolve_deterministic("10001", "代金券")
        self.assertEqual("coupon_catalog", result["kind"])
        self.assertIn("没有代金券选项", result["reply"])
        self.assertNotIn("100元", result["reply"])

    def test_any_have_question_is_a_deterministic_sku_lookup(self):
        self.store.save_v2_product(
            "10001", "测试商品",
            "100元代金券：售价54元，最多叠加2张\n双人套餐：售价118元",
        )
        existing = self.store.resolve_deterministic("10001", "100元代金券有吗")
        self.assertEqual("sku_availability", existing["kind"])
        self.assertIn("售价54元", existing["reply"])
        missing = self.store.resolve_deterministic("10001", "300元代金券有吗")
        self.assertEqual("sku_availability", missing["kind"])
        self.assertIn("没有300元代金券", missing["reply"])
        self.assertIn("没有“300元”这一有货规格", missing["reply"])
        self.assertIn("到店咨询", missing["reply"])
        self.assertNotIn("2张100元代金券", missing["reply"])
        self.assertNotIn("SKU", missing["reply"].upper())

    def test_bare_still_available_question_is_stock_not_product_option(self):
        result = self.store.resolve_deterministic("10001", "还有吗")
        self.assertEqual("stock", result["kind"])
        self.assertEqual("有的，当前商品还在售，可以直接拍下。", result["reply"])
        self.assertNotIn("商品规格", result["reply"])

    def test_prefix_have_denomination_matches_real_option(self):
        self.store.save_v2_product(
            "10001", "测试代金券", "100元代金券：售价54元\n300元代金券：售价155元"
        )
        result = self.store.resolve_deterministic("10001", "有300吗")
        self.assertEqual("sku_availability", result["kind"])
        self.assertIn("300元代金券", result["reply"])
        self.assertIn("155元", result["reply"])
        self.assertNotIn("100元代金券", result["reply"])

    def test_consumption_plan_prefers_exact_or_closest_real_option(self):
        self.store.save_v2_product(
            "10001", "闽和南代金券",
            "100元代金券：售价85元\n200元代金券：售价170元\n"
            "300元代金券：售价235元\n400元代金券：售价340元",
        )
        exact = self.store.resolve_deterministic("10001", "300怎么买")
        self.assertEqual("consumption_plan", exact["kind"])
        self.assertIn("300元代金券，售价235元", exact["reply"])
        plan = self.store.resolve_deterministic("10001", "390元怎么买")
        self.assertEqual("consumption_plan", plan["kind"])
        self.assertIn("300元代金券", plan["reply"])
        self.assertIn("剩余90元", plan["reply"])
        self.assertNotIn("400元代金券", plan["reply"])
        punctuated = self.store.resolve_deterministic("10001", "390元，怎么买")
        self.assertEqual("consumption_plan", punctuated["kind"])
        self.assertIn("300元代金券", punctuated["reply"])
        self.assertIn("剩余90元", punctuated["reply"])

    def test_spoken_bill_amount_chooses_best_real_option(self):
        self.store.save_v2_product(
            "10001", "小江溪代金券",
            "100元代金券：售价64元（最多叠加3张）\n"
            "125元代金券：售价74元（最多叠加4张）\n"
            "200元代金券：售价120元（最多叠加1张）",
        )
        result = self.store.resolve_deterministic("10001", "吃216怎么买")
        self.assertEqual("consumption_plan", result["kind"])
        self.assertIn("200元代金券", result["reply"])
        self.assertIn("价格120元", result["reply"])
        self.assertIn("剩余16元", result["reply"])
        self.assertNotIn("100元代金券", result["reply"])

    def test_missing_amount_uses_delivered_coupon_composition(self):
        self.store.save_v2_product(
            "10001", "同仁四季代金券",
            "200元代金券：售价108元，发100元券2张（最多叠加2张）",
        )
        result = self.store.resolve_deterministic("10001", "300多少钱")
        self.assertEqual("price", result["kind"])
        self.assertIn("没有300元代金券", result["reply"])
        self.assertIn("2张100元代金券", result["reply"])
        self.assertIn("共支付108元", result["reply"])
        self.assertIn("剩余100元", result["reply"])

    def test_bundle_amount_and_quantity_followup_use_delivered_coupon_count(self):
        self.store.save_v2_product(
            "10001", "美团代金券", "仅支持同面额代金券叠加，每次最多使用5张。",
        )
        platform = json.dumps({
            "title": "美团代金券",
            "sku": [
                {"priceInCent": 14980, "propertyList": [
                    {"actualValueText": "美团100（100x2）"},
                ]},
                {"priceInCent": 22470, "propertyList": [
                    {"actualValueText": "美团100（100x3）"},
                ]},
            ],
        }, ensure_ascii=False)
        with self.store._connect() as conn:
            conn.execute(
                "UPDATE v2_products SET platform_summary=? WHERE item_id=?",
                (platform, "10001"),
            )

        first = self.store.resolve_deterministic("10001", "300多少钱")
        self.assertIn("美团100（100x3）", first["reply"])
        self.assertIn("售价224.7元", first["reply"])
        self.assertIn("发100元券3张", first["reply"])
        self.assertNotIn("没有300元代金券", first["reply"])

        # Exercise the contextual path as well: a prior 300-yuan target must
        # be replaced by the new delivered-coupon quantity, not multiplied.
        context = {
            "price_filters": {
                "target_amount": "300", "amount_kind": "option", "intent": "price",
                "updated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
            },
        }
        followup = self.store.resolve_deterministic("10001", "3张多少钱", context)
        self.assertIn("美团100（100x3）", followup["reply"])
        self.assertIn("售价224.7元", followup["reply"])
        self.assertIn("发3张100元代金券", followup["reply"])
        self.assertIn("抵扣300元", followup["reply"])
        self.assertNotIn("6张", followup["reply"])
        self.assertNotIn("449.4", followup["reply"])

        skus = self.store.list_product_skus("10001")
        self.assertEqual(2, len({sku["sku_key"] for sku in skus}))

    def test_sku_name_bundle_overrides_single_face_fallback(self):
        product = {
            "title": "NEED韩国料理代金券",
            "structured": {"products": [{
                "name": "NEED韩国料理200（100x2）",
                "option_type": "voucher", "face_value": "200", "sale_price": "132",
            }]},
        }
        options = self.store.extract_product_options(product)
        self.assertEqual("100元券2张", options[0]["composition"])
        reply = self.store.price_reply(product, "200多少钱")
        self.assertIn("NEED韩国料理200（100x2）", reply)
        self.assertIn("售价132元", reply)
        self.assertIn("发100元券2张", reply)
        self.assertNotIn("发200元券1张", reply)

    def test_how_much_offsets_value_queries_the_existing_option(self):
        self.store.save_v2_product("10001", "测试代金券", "300元代金券：售价235元")
        result = self.store.resolve_deterministic("10001", "多少代300")
        self.assertEqual("price", result["kind"])
        self.assertIn("300元代金券", result["reply"])
        self.assertIn("235元", result["reply"])

    def test_purchase_order_flow_and_price_are_both_silent(self):
        self.store.save_v2_product(
            "10001", "代买服务", "人工处理", coupon_type="purchase_order",
            coupon_instructions="买家点餐后发桌码，拍下商品后，改价为实际金额*0.78，买家付款后由卖家代付账单",
        )
        flow = self.store.resolve_deterministic("10001", "怎么买")
        self.assertEqual("purchase_order_other", flow["kind"])
        self.assertEqual("silent", flow["decision"])
        self.assertEqual("", flow["reply"])
        price = self.store.resolve_deterministic("10001", "390元怎么买")
        self.assertEqual("purchase_order_other", price["kind"])
        self.assertEqual("silent", price["decision"])
        self.assertEqual("", price["reply"])

    def test_missing_coupon_denomination_recommends_closest_lower_sku(self):
        self.store.save_v2_product(
            "10001", "测试代金券",
            "100元代金券：售价54元\n200元代金券：售价108元\n300元代金券：售价155元",
        )
        result = self.store.resolve_deterministic("10001", "260多少钱")
        self.assertEqual("price", result["kind"])
        self.assertIn("没有260元代金券", result["reply"])
        self.assertIn("200元代金券", result["reply"])
        self.assertNotIn("300元代金券", result["reply"])

    def test_named_package_price_query_uses_matching_real_sku(self):
        self.store.save_v2_product(
            "10001", "自助餐",
            "单人经典自助：售价165元\n双人经典自助：售价324元\n三人经典自助：售价493元",
        )
        result = self.store.resolve_deterministic("10001", "双人经典自助多少钱")
        self.assertEqual("price", result["kind"])
        self.assertIn("双人经典自助", result["reply"])
        self.assertIn("324元", result["reply"])
        self.assertNotIn("493元", result["reply"])

    def test_social_words_are_removed_before_price_intent_routing(self):
        self.store.save_v2_product(
            "10001", "同仁四季椰子鸡",
            "200元代金券（100x2）：售价115.8元，发100元券2张\n300元代金券（限一）：售价212.8元，发300元券1张",
        )
        for message in (
            "你好，什么价格", "您好，请问一下多少钱",
            "麻烦问下啥价，谢谢", "哈喽，报价，辛苦了",
        ):
            with self.subTest(message=message):
                result = self.store.resolve_deterministic("10001", message)
                self.assertEqual("price", result["kind"])
                self.assertIn("当前可选规格及价格如下", result["reply"])
                self.assertIn("115.8元", result["reply"])
                self.assertIn("212.8元", result["reply"])
                self.assertNotIn("没有“你好”", result["reply"])

        greeting = self.store.resolve_deterministic("10001", "你好")
        self.assertEqual("greeting", greeting["kind"])
        thanks = self.store.resolve_deterministic("10001", "谢谢，辛苦了")
        self.assertEqual("courtesy", thanks["kind"])
        self.assertIn("不客气", thanks["reply"])

    def test_intent_profile_separates_action_entities_and_social_tone(self):
        cases = {
            "你好，什么价格": ("price", "catalog_price"),
            "麻烦问下3张多少钱": ("price", "quantity_price"),
            "108抵200": ("price", "voucher_value_confirmation"),
            "消费680怎么拍": ("price", "consumption_plan"),
            "退款到哪了": ("aftersale", "refund_status"),
            "100和300能一起用吗": ("stacking", "mixed_denomination"),
            "有两百的吗": ("sku", "sku_availability"),
        }
        for message, expected in cases.items():
            with self.subTest(message=message):
                profile = self.store.classify_buyer_intents(message)
                self.assertEqual(expected[0], profile["primary_intent"])
                self.assertEqual(expected[1], profile["primary_subtype"])

        combined = self.store.classify_buyer_intents(
            "您好，深圳壹方城明天晚餐3个成年人多少钱，谢谢"
        )
        self.assertEqual("thanks", combined["social_kind"])
        self.assertEqual("store", combined["primary_intent"])
        self.assertEqual(
            {"store", "price", "time"},
            {row["intent"] for row in combined["intents"]},
        )
        self.assertIn("成人", combined["entities"]["audiences"])
        self.assertIn("晚餐", combined["entities"]["meal_periods"])
        self.assertEqual(["3"], combined["entities"]["people"])

    def test_copula_price_and_chinese_denomination_questions_use_clean_sku_entities(self):
        self.store.save_v2_product(
            "10001", "测试代金券",
            "100元代金券：售价57.9元，最多叠加2张\n"
            "300元代金券：售价212.8元，每次限用1张",
        )
        price = self.store.resolve_deterministic("10001", "请问300的代金券是多少钱")
        self.assertEqual("price", price["kind"])
        self.assertIn("300元代金券", price["reply"])
        self.assertIn("212.8元", price["reply"])
        self.assertNotIn("300的代金券是", price["reply"])

        availability = self.store.resolve_deterministic("10001", "有两百的吗")
        self.assertEqual("sku_availability", availability["kind"])
        self.assertIn("没有200元代金券", availability["reply"])
        self.assertIn("2张100元代金券", availability["reply"])
        self.assertIn("支付115.8元", availability["reply"])

    def test_named_package_availability_uses_real_product_text(self):
        self.store.save_v2_product(
            "10001", "自助餐",
            "单人经典自助：售价165元\n双人经典自助：售价324元",
        )
        existing = self.store.resolve_deterministic("10001", "双人经典自助有吗")
        self.assertEqual("sku_availability", existing["kind"])
        self.assertIn("双人经典自助", existing["reply"])
        self.assertIn("324元", existing["reply"])
        missing = self.store.resolve_deterministic("10001", "三人套餐有吗")
        self.assertEqual("sku_availability", missing["kind"])
        self.assertIn("没有“三人套餐”这一规格", missing["reply"])

    def test_product_attribute_questions_never_fall_into_store_search(self):
        self.store.save_v2_product(
            "10001", "鱼酷烤鱼2-3人餐",
            "鱼酷活鱼烤鱼单条鱼2.6斤：118元（不含米饭配菜）",
        )
        self.store.import_store_list(self.create_multi_region_store_sheet(), "多地区门店", ["10001"])
        weight = self.store.resolve_deterministic("10001", "鱼多重")
        self.assertEqual("product_attribute", weight["kind"])
        self.assertIn("单条鱼约2.6斤", weight["reply"])
        self.assertIn("不含米饭和配菜", weight["reply"])
        self.assertNotIn("可用门店", weight["reply"])
        flavor = self.store.resolve_deterministic("10001", "有什么口味")
        self.assertEqual("product_attribute", flavor["kind"])
        self.assertEqual("allow", flavor["decision"])
        self.assertIn("暂未说明具体口味", flavor["reply"])
        self.assertNotIn("可用门店", flavor["reply"])

    def test_numeric_availability_never_matches_a_store_address(self):
        self.store.save_v2_product(
            "10001", "测试代金券", "100元代金券：售价54元，最多叠加2张"
        )
        self.store.import_store_list(self.create_store_sheet(), "北京门店", ["10001"])
        with self.store._connect() as conn:
            conn.execute(
                "UPDATE stores SET address=?, normalized=normalized || ? WHERE branch=?",
                ("复兴路300号", "复兴路300号", "华熙五棵松店"),
            )
        result = self.store.resolve_deterministic("10001", "300有吗")
        self.assertEqual("sku_availability", result["kind"])
        self.assertIn("没有300元代金券", result["reply"])
        self.assertNotIn("五棵松", result["reply"])

    def test_refresh_request_never_claims_operation_completed(self):
        unclear = self.store.resolve_deterministic("10001", "需要刷新一下")
        self.assertEqual("refresh_clarification", unclear["kind"])
        self.assertNotIn("已刷新", unclear["reply"])
        risky = self.store.resolve_deterministic("10001", "帮我刷新券码")
        self.assertEqual("review", risky["decision"])
        self.assertIn("72小时", risky["reply"])
        self.assertNotIn("已刷新", risky["reply"])

    def test_resend_or_cancelled_previous_item_never_promises_reissue(self):
        for message in ("需要重新发一下", "请再发一个", "刚刚那个撤销了"):
            with self.subTest(message=message):
                result = self.store.resolve_deterministic("10001", message)
                self.assertEqual("review", result["decision"])
                self.assertIn("人工", result["reply"])
                self.assertIn("72小时", result["reply"])
                self.assertNotIn("帮您补发", result["reply"])

    def test_code_anomaly_without_refund_intent_starts_troubleshooting(self):
        for message in ("核销失败", "券码无效", "无法核销"):
            with self.subTest(message=message):
                result = self.store.resolve_deterministic("10001", message)
                self.assertEqual("allow", result["decision"])
                self.assertEqual("coupon_troubleshooting", result["kind"])
                self.assertIn("核销", result["reply"])
                self.assertNotIn("退款", result["reply"])

    def test_mixed_denomination_negation_is_never_reversed(self):
        self.store.save_v2_product(
            "10001", "多面额代金券",
            "100元代金券：售价85元（最多叠加4张）\n"
            "300元代金券：售价235元（最多叠加2张）\n"
            "不同面额代金券不能互相叠加",
        )
        product = self.store.get_v2_product("10001")
        summary = self.store.build_knowledge_summary(product)
        self.assertIn("仅支持同面额代金券叠加", summary)
        self.assertNotIn("支持不同面额代金券叠加", summary)
        reply = self.store.stacking_reply(product, "100元和300元可以一起用吗")
        self.assertIn("不能一起使用", reply)

    def test_mixed_denomination_unlimited_rule_is_preserved(self):
        self.store.save_v2_product(
            "10001", "多面额代金券",
            "100元代金券：售价90元\n200元代金券：售价180元\n"
            "不同面额代金券可互相叠加，可无限叠加，不限制数量",
        )
        summary = self.store.build_knowledge_summary(self.store.get_v2_product("10001"))
        self.assertIn("支持同面额及不同面额代金券互相叠加", summary)
        self.assertIn("不限制使用张数", summary)
        self.assertNotIn("仅支持同面额", summary)

    def test_denomination_specific_stack_limits_override_collapsed_ai_summary(self):
        self.store.save_v2_product(
            "10001", "多面额代金券",
            "100元代金券：售价68元\n"
            "200元代金券：售价142元\n"
            "300元代金券：售价202元\n"
            "仅支持同面额代金券叠加，100代金券（可叠加2张）、"
            "300代金券（仅限一张）、200代金券（可叠加5张）。",
        )
        self.store.save_ai_summary(
            "10001", "错误地统一为最多2张",
            {
                "summary": "错误地统一为最多2张",
                "products": [
                    {"name": "100元代金券", "face_value": "100", "sale_price": "68", "max_stack": "2"},
                    {"name": "200元代金券", "face_value": "200", "sale_price": "142", "max_stack": "2"},
                    {"name": "300元代金券", "face_value": "300", "sale_price": "202", "max_stack": "2"},
                ],
                "facts": {"叠加规则": "仅支持同面额代金券叠加，每次最多使用2张"},
                "time_rules": [],
            },
        )
        product = self.store.get_v2_product("10001")
        options = {option["face_value"]: option["max_stack"] for option in self.store.extract_product_options(product)}
        self.assertEqual({"100": "2", "200": "5", "300": "1"}, options)
        self.assertIn("最多使用2张", self.store.stacking_reply(product, "100元代金券最多可以用几张"))
        self.assertIn("最多使用5张", self.store.stacking_reply(product, "200元代金券最多可以用几张"))
        self.assertIn("最多使用1张", self.store.stacking_reply(product, "300元代金券最多可以用几张"))
        generic = self.store.stacking_reply(product, "代金券最多可以用几张")
        self.assertIn("100元代金券最多使用2张", generic)
        self.assertIn("200元代金券最多使用5张", generic)
        self.assertIn("300元代金券最多使用1张", generic)

    def test_package_price_is_never_treated_as_coupon_denomination(self):
        self.store.save_v2_product(
            "10001", "鱼酷烤鱼2-3人餐",
            "鱼酷活鱼烤鱼套餐：售价118元（单条鱼2.6斤，不含米饭配菜）",
        )
        self.store.save_ai_summary(
            "10001", "错误旧摘要",
            {
                "summary": "错误旧摘要",
                "products": [{
                    "name": "可选=118（北京上海", "face_value": "118",
                    "sale_price": "59", "composition": "",
                }],
                "facts": {}, "time_rules": [],
            },
        )
        product = self.store.get_v2_product("10001")
        self.assertEqual([], self.store.extract_product_options(product))
        result = self.store.resolve_deterministic("10001", "300券有吗")
        self.assertEqual("sku_availability", result["kind"])
        self.assertIn("没有300元代金券", result["reply"])
        self.assertIn("鱼酷活鱼烤鱼套餐", result["reply"])
        self.assertIn("售价118元", result["reply"])
        self.assertIn("不属于代金券商品", result["reply"])
        self.assertNotIn("2张118", result["reply"])

    def test_strict_flavor_extraction_rejects_price_and_region_noise(self):
        self.store.save_v2_product(
            "10001", "鱼酷烤鱼套餐", "口味：可选=118（北京上海。套餐售价118元。"
        )
        result = self.store.resolve_deterministic("10001", "什么口味")
        self.assertEqual("allow", result["decision"])
        self.assertIn("暂未说明具体口味", result["reply"])
        self.assertNotIn("转人工", result["reply"])
        self.assertNotIn("72小时", result["reply"])
        self.assertNotIn("118", result["reply"])

    def test_strict_flavor_extraction_keeps_explicit_complete_flavors(self):
        self.store.save_v2_product(
            "10001", "烤鱼套餐", "可选口味：香辣、蒜香、番茄。"
        )
        result = self.store.resolve_deterministic("10001", "有什么口味")
        self.assertEqual("allow", result["decision"])
        self.assertEqual("当前可选口味有：香辣、蒜香、番茄。", result["reply"])

    def test_fixed_flavor_rule_does_not_invent_choices(self):
        self.store.save_v2_product("10001", "套餐", "本套餐为固定口味，不可选口味。")
        result = self.store.resolve_deterministic("10001", "能选口味吗")
        self.assertEqual("当前商品为固定口味，暂不支持选择口味哦。", result["reply"])

    def test_voucher_and_package_default_to_not_combinable(self):
        self.store.save_v2_product(
            "10001", "测试代金券", "100元代金券：售价54元，最多叠加2张"
        )
        result = self.store.resolve_deterministic("10001", "可以和套餐一起用吗")
        self.assertEqual("benefit_combination", result["kind"])
        self.assertEqual("代金券和套餐不能一起使用哦。", result["reply"])

    def test_explicit_voucher_package_combination_is_honored(self):
        self.store.save_v2_product(
            "10001", "测试代金券",
            "100元代金券：售价54元\n本代金券可以与套餐一起使用。",
        )
        result = self.store.resolve_deterministic("10001", "代金券可以和套餐一起用吗")
        self.assertIn("明确说明", result["reply"])
        self.assertIn("可以和套餐一起使用", result["reply"])

    def test_cross_platform_coupons_are_never_mixed(self):
        self.store.save_v2_product(
            "10001", "测试美团代金券", "100元代金券：售价54元", coupon_type="meituan"
        )
        for message in (
            "可以和抖音券一起用吗", "抖音券和美团券能叠加吗", "能和小程序券一起抵扣吗",
            "可以叠加抖音吗", "能和美团一起用吗", "小程序能同时用吗",
        ):
            with self.subTest(message=message):
                result = self.store.resolve_deterministic("10001", message)
                self.assertEqual("benefit_combination", result["kind"])
                self.assertIn("不能与", result["reply"])
                self.assertIn("优惠券混用", result["reply"])

    def test_external_benefit_phrasings_never_fall_into_denomination_stacking(self):
        self.store.save_v2_product(
            "10001", "测试代金券", "100元代金券：售价54元，最多叠加2张"
        )
        for message in (
            "我有会员券能一起用吗", "生日券可以同时用不", "店内红包还能叠吗",
            "银行卡优惠可以再减吗", "积分能不能一起抵扣", "我买了团购还能用这个吗",
        ):
            with self.subTest(message=message):
                result = self.store.resolve_deterministic("10001", message)
                self.assertEqual("benefit_combination", result["kind"])
                self.assertIn("单独使用", result["reply"])
                self.assertNotIn("同面额", result["reply"])

    def test_buyer_owned_external_coupon_is_not_assumed_combinable(self):
        self.store.save_v2_product(
            "10001", "测试代金券", "100元代金券：售价54元"
        )
        for message in ("我有一张别的券，可以一起用吗", "店里送的优惠券还能用吗"):
            with self.subTest(message=message):
                result = self.store.resolve_deterministic("10001", message)
                self.assertEqual("benefit_combination", result["kind"])
                self.assertIn("请勿混用", result["reply"])

    def test_using_with_friends_is_not_mistaken_for_coupon_mixing(self):
        product = self.store.get_v2_product("10001")
        self.assertEqual("", self.store.benefit_combination_reply(product, "可以和朋友一起用吗"))
        self.assertEqual("", self.store.stacking_reply(product, "可以和朋友一起用吗"))

    def test_same_day_policy_uses_polite_direct_answers(self):
        stockpile = self.store.resolve_deterministic("10001", "可以囤货吗")
        self.assertEqual(
            "不建议提前囤券哦。这款券需要当天购买、当天使用，建议您确定到店当天再下单，避免因未及时使用造成过期。",
            stockpile["reply"],
        )
        self.assertNotIn("不退款", stockpile["reply"])
        later = self.store.resolve_deterministic("10001", "可以过两天用吗")
        self.assertIn("暂不支持购买后过两天再用", later["reply"])
        self.assertIn("使用会更稳妥", later["reply"])
        self.assertNotIn("不补发", later["reply"])

    def test_specific_date_range_is_checked_before_price_or_bargaining(self):
        self.store.save_v2_product(
            "10001", "节日自助券",
            "成人单人全天券：售价88元\n中秋（9月25日至27日）不可用。",
        )
        blocked = self.store.resolve_deterministic("10001", "9.26可以用吗")
        self.assertEqual("date_use", blocked["kind"])
        self.assertIn("不可用日期范围", blocked["reply"])
        self.assertNotIn("议价", blocked["reply"])
        allowed = self.store.resolve_deterministic("10001", "9.28可以用吗")
        self.assertEqual("date_use", allowed["kind"])
        self.assertIn("可以使用", allowed["reply"])

    def test_unspecified_day_or_meal_metadata_is_unrestricted(self):
        self.store.save_v2_product("10001", "双人自助", "双人自助套餐：售价176元")
        for message in ("周末双人多少钱", "工作日双人多少钱", "中午双人多少钱", "晚上双人多少钱"):
            with self.subTest(message=message):
                result = self.store.resolve_deterministic("10001", message)
                self.assertIn("176元", result["reply"])
                self.assertNotIn("没有符合", result["reply"])

    def test_two_people_fall_back_to_two_single_person_options(self):
        self.store.save_v2_product(
            "10001", "奶糖爸爸自助小火锅",
            "单人自助：售价48元，平日周末全天通用",
        )
        result = self.store.resolve_deterministic("10001", "你好，请问双人多少钱")
        self.assertEqual("price", result["kind"])
        self.assertIn("2人需要购买2份", result["reply"])
        self.assertIn("共96元", result["reply"])
        self.assertNotIn("没有符合", result["reply"])

    def test_requested_weekend_face_value_ignores_other_sku_denial(self):
        self.store.save_v2_product(
            "10001", "半秋山代金券",
            "100元代金券：售价66.8元，仅限工作日使用，周末不可用\n"
            "300元代金券：售价238.5元，周末可用",
        )
        result = self.store.resolve_deterministic("10001", "周末可用300的券")
        self.assertEqual("day_use", result["kind"])
        self.assertIn("300元代金券", result["reply"])
        self.assertIn("周末", result["reply"])
        self.assertIn("可以使用", result["reply"])
        self.assertNotIn("不可以", result["reply"])

    def test_city_followup_filters_previous_store_candidates(self):
        self.store.save_v2_product(
            "10001", "鱼酷烤鱼", "鱼酷活鱼烤鱼单条鱼2.6斤：118元",
        )
        self.store.import_store_text(
            "【天津】\n天津凯德MALL天津湾店\n天津空港SM广场店\n"
            "【青岛】\n青岛合美MALL店\n青岛崂山金狮店\n"
            "【武汉】\n武汉江汉路M+店\n"
            "【西安】\n西安MOMOPARK店\n"
            "【南京】\n南京南站喜玛拉雅店\n南京大观天地店",
            "全国门店", ["10001"],
        )
        first = self.store.resolve_deterministic("10001", "m+购物中心店能用吗")
        self.assertGreater(len(first["store_matches"]), 1)
        context = {
            "query": first["store_query"], "matches": first["store_matches"],
            "status": first.get("store_status", "available"),
            **(first.get("store_context_update") or {}),
        }
        followup = self.store.resolve_deterministic(
            "10001", "武汉的", store_context=context,
        )
        self.assertEqual("stores", followup["kind"])
        self.assertEqual("allow", followup["decision"])
        self.assertEqual(["武汉江汉路M+店"], [
            row["branch"] for row in followup["store_matches"]
        ])
        self.assertNotIn("暂未查询到", followup["reply"])

    def test_all_day_and_all_period_options_cover_requested_conditions(self):
        self.store.save_v2_product(
            "10001", "通用自助",
            "全天通用双人券：售价176元\n全时段三人券：售价264元",
        )
        lunch = self.store.resolve_deterministic("10001", "中午双人多少钱")
        self.assertIn("176元", lunch["reply"])
        weekend = self.store.resolve_deterministic("10001", "周末三人多少钱")
        self.assertIn("264元", weekend["reply"])

    def test_weekend_uses_holiday_tier_only_when_no_weekend_tier_exists(self):
        self.store.save_v2_product(
            "10001", "双档自助",
            "平日双人经典自助：售价288元\n节假日双人经典自助：售价330元",
        )
        result = self.store.resolve_deterministic("10001", "周末双人多少钱")
        self.assertIn("330元", result["reply"])
        self.assertNotIn("288元", result["reply"])

    def test_explicit_meal_option_beats_all_day_option(self):
        self.store.save_v2_product(
            "10001", "分时自助",
            "全天双人自助：售价176元\n晚市双人自助：售价162元",
        )
        result = self.store.resolve_deterministic("10001", "晚上双人多少钱")
        self.assertIn("162元", result["reply"])
        self.assertNotIn("176元", result["reply"])

    def test_mixed_adult_child_price_uses_real_tickets_and_constraints(self):
        self.store.save_v2_product(
            "10001", "家庭自助",
            "成人双人券：售价176元\n儿童票：售价68元，限身高1.2米至1.4米儿童",
        )
        result = self.store.resolve_deterministic("10001", "两大一小多少钱")
        self.assertEqual("audience_price", result["kind"])
        self.assertIn("176元", result["reply"])
        self.assertIn("68元", result["reply"])
        self.assertIn("244元", result["reply"])
        self.assertIn("身高", result["reply"])

    def test_missing_child_ticket_quotes_adults_only(self):
        self.store.save_v2_product("10001", "成人自助", "成人双人券：售价176元")
        result = self.store.resolve_deterministic("10001", "两大一小多少钱")
        self.assertIn("176元", result["reply"])
        self.assertIn("\n\n当前商品暂时没有“儿童票”这一规格", result["reply"])
        self.assertIn("到店咨询", result["reply"])
        self.assertNotIn("三人", result["reply"])

    def test_student_senior_and_female_tickets_are_not_conflated(self):
        self.store.save_v2_product(
            "10001", "身份自助",
            "学生票：售价88元，需出示学生证\n老人票：售价68元，限60周岁以上\n女士票：售价99元",
        )
        student = self.store.resolve_deterministic("10001", "一个学生多少钱")
        self.assertIn("88元", student["reply"])
        self.assertIn("学生证", student["reply"])
        senior = self.store.resolve_deterministic("10001", "一个老人多少钱")
        self.assertIn("68元", senior["reply"])
        female = self.store.resolve_deterministic("10001", "一个女士多少钱")
        self.assertIn("99元", female["reply"])

    def test_missing_denomination_is_treated_as_consumption_plan(self):
        self.store.save_v2_product(
            "10001", "100元代金券", "100元代金券：售价54元，最多使用3张",
        )
        two_hundred = self.store.resolve_deterministic("10001", "200的多少钱")
        self.assertIn("2张100元代金券", two_hundred["reply"])
        self.assertIn("108元", two_hundred["reply"])
        six_hundred = self.store.resolve_deterministic("10001", "600怎么卖")
        self.assertIn("最多", six_hundred["reply"])
        self.assertIn("剩余", six_hundred["reply"])

    def test_date_usage_does_not_inherit_previous_amount_context(self):
        self.store.save_v2_product(
            "10001", "100元代金券", "100元代金券：售价54元，最多使用2张",
        )
        first = self.store.resolve_deterministic("10001", "300怎么卖")
        context = first.get("query_context_update") or {}
        followup = self.store.resolve_deterministic(
            "10001", "明天可以用吗", store_context=context,
        )
        self.assertEqual("date_use", followup["kind"])
        self.assertNotIn("300", followup["reply"])
        self.assertNotIn("商品选项", followup["reply"])

    def test_ticket_quantity_and_people_followup_keep_meal_period_separate(self):
        self.store.save_v2_product(
            "10001", "烤肉自助",
            "全天单人自助：售价88元\n全天双人自助：售价176元\n"
            "晚市单人自助：售价80元\n晚市双人自助：售价162元",
        )
        first = self.store.resolve_deterministic("10001", "晚上用，两张多少钱")
        self.assertIn("单人券、双人券", first["reply"])
        filters = first["query_context_update"]["price_filters"]
        self.assertEqual(2, filters["purchase_quantity"])
        self.assertEqual("dinner", filters["meal_period"])
        second = self.store.resolve_deterministic(
            "10001", "2", store_context=first["query_context_update"],
        )
        self.assertIn("2张晚市双人自助", second["reply"])
        self.assertIn("324元", second["reply"])
        self.assertNotIn("全天双人", second["reply"])

    def test_multi_question_separates_city_and_coupon_purchase(self):
        self.store.save_v2_product(
            "10001", "NEED韩国料理代金券",
            "100元代金券：售价66.6元\n"
            "300元代金券：售价199.8元，发100元券3张",
        )
        self.store.import_store_text(
            "【山东省】\n【济南】济南万象城店", "济南门店", ["10001"],
        )
        result = self.store.resolve_deterministic(
            "10001", "你好 济南可以用是吧 可以直接拍代300的那个",
        )
        self.assertEqual("multi_intent", result["kind"])
        self.assertEqual(["store", "purchase"], result["resolved_intents"])
        self.assertEqual("济南", result["store_query"])
        self.assertIn("济南万象城店", result["reply"])
        self.assertIn("300元代金券", result["reply"])
        self.assertIn("售价199.8元", result["reply"])
        self.assertNotIn("济南是吧", result["reply"])
        self.assertNotIn("无法准确回答", result["reply"])

    def test_location_matching_uses_most_specific_normalized_unit(self):
        self.store.import_store_list(
            self.create_multi_region_store_sheet(), "多地区门店", ["10001"],
        )
        district = self.store.resolve_deterministic(
            "10001", "老板，深圳 南山区可以用吗？",
        )
        self.assertEqual("stores", district["kind"])
        self.assertEqual("深圳南山区", district["store_query"])
        self.assertIn("西丽益田假日里店", district["reply"])
        self.assertNotIn("龙岗万科里店", district["reply"])

        street = self.store.resolve_deterministic(
            "10001", "留仙 大道，可以用吗？",
        )
        self.assertEqual("stores", street["kind"])
        self.assertEqual("留仙大道", street["store_query"])
        self.assertIn("西丽益田假日里店", street["reply"])

    def test_multi_question_answers_usage_store_value_and_limits_once(self):
        self.store.save_v2_product(
            "10001", "同仁四季椰子鸡代金券",
            "200元代金券：售价108元，发100元券2张\n"
            "一桌最多代200元；仅限前台验电子二维码",
            coupon_type="meituan",
        )
        self.store.save_ai_summary("10001", "规则摘要", {
            "facts": {
                "使用规则": "仅支持同面额代金券叠加；一桌最多代200元；仅限前台验电子二维码",
            },
        })
        self.store.import_store_text(
            "【广东省】\n【深圳】深圳壹方城店", "深圳门店", ["10001"],
        )
        result = self.store.resolve_deterministic(
            "10001", "怎么用，深圳壹方城店可以用吗，108抵用200，有什么限制条件吗",
        )
        self.assertEqual("multi_intent", result["kind"])
        self.assertEqual(
            ["usage", "store", "voucher_value", "restrictions"],
            result["resolved_intents"],
        )
        self.assertEqual("深圳壹方城店", result["store_query"])
        self.assertIn("1. 使用方式", result["reply"])
        self.assertIn("2. 适用门店", result["reply"])
        self.assertIn("3. 售价与面额", result["reply"])
        self.assertIn("4. 使用限制", result["reply"])
        self.assertIn("查询到可用门店", result["reply"])
        self.assertIn("深圳壹方城店", result["reply"])
        self.assertIn("是的，200元代金券售价108元", result["reply"])
        self.assertIn("2张100元代金券", result["reply"])
        self.assertIn("一桌最多代200元", result["reply"])
        self.assertNotIn("无法准确回答", result["reply"])

    def test_paid_price_face_value_confirmation_is_direct_and_grounded(self):
        self.store.save_v2_product(
            "10001", "同仁四季椰子鸡代金券",
            "200元代金券：售价108元，发100元券2张",
        )
        result = self.store.resolve_deterministic("10001", "108抵200吗")
        self.assertEqual("voucher_value", result["kind"])
        self.assertEqual(
            "是的，200元代金券售价108元，购买后发放2张100元代金券，共可抵扣200元。",
            result["reply"],
        )
        self.assertNotIn("无法", result["reply"])

    def test_composed_coupon_quantity_names_unit_count_and_total_cap(self):
        self.store.save_v2_product(
            "10001", "同仁四季椰子鸡代金券",
            "200元代金券：售价108元，发100元券2张\n"
            "一桌最多代200元；不同面额不能混用",
        )
        result = self.store.resolve_deterministic(
            "10001", "200代金券一次可以用几张",
        )
        self.assertEqual("stacking", result["kind"])
        self.assertIn("200元代金券售价108元", result["reply"])
        self.assertIn("发放2张100元代金券", result["reply"])
        self.assertIn("这2张可以同一次使用，合计抵扣200元", result["reply"])
        self.assertIn("每次最多使用2张100元券", result["reply"])
        self.assertNotIn("200元代金券可以叠加", result["reply"])

    def test_location_landmark_matches_branch_with_unspoken_district_prefix(self):
        self.store.import_store_text(
            "【广东省】\n【深圳】宝安壹方城店\n【广州】天环广场店",
            "广东门店", ["10001"],
        )
        result = self.store.resolve_deterministic(
            "10001", "200的券深圳壹方城可以用不",
        )
        self.assertEqual("multi_intent", result["kind"])
        self.assertEqual(["sku", "store"], result["resolved_intents"])
        self.assertEqual("深圳壹方城", result["store_query"])
        self.assertIn("商品规格", result["reply"])
        self.assertIn("宝安壹方城店", result["reply"])
        self.assertNotIn("广州", result["reply"])

    def test_store_and_unavailable_date_are_answered_as_two_intents(self):
        self.store.save_v2_product(
            "10001", "NEED韩国料理代金券",
            "100元代金券：售价66.6元。除2026年9月25日至9月27日外营业时间可用。",
        )
        self.store.import_store_text(
            "【广东省】\n【深圳】宝安壹方城店", "深圳门店", ["10001"],
        )
        result = self.store.resolve_deterministic(
            "10001", "深圳壹方城9月26日可以用吗",
        )
        self.assertEqual("multi_intent", result["kind"])
        self.assertEqual(["store", "date"], result["resolved_intents"])
        self.assertIn("宝安壹方城店", result["reply"])
        self.assertIn("9月26日", result["reply"])
        self.assertIn("不能使用", result["reply"])

    def test_store_and_price_are_answered_separately(self):
        self.store.save_v2_product(
            "10001", "鱼酷烤鱼2-3人餐", "鱼酷活鱼烤鱼套餐：售价118元",
        )
        self.store.import_store_text(
            "【广东省】\n【广州】天环广场店", "广州门店", ["10001"],
        )
        result = self.store.resolve_deterministic("10001", "天环广场店多少钱")
        self.assertEqual("multi_intent", result["kind"])
        self.assertEqual(["store", "price"], result["resolved_intents"])
        self.assertIn("天环广场店", result["reply"])
        self.assertIn("118元", result["reply"])

    def test_wrong_store_and_delivery_mismatch_are_aftersales_not_store_queries(self):
        wrong_store = self.store.resolve_deterministic("10001", "我看错门店了")
        self.assertEqual("aftersale_clarify", wrong_store["kind"])
        self.assertIn("是否已经付款", wrong_store["reply"])
        mismatch = self.store.resolve_deterministic(
            "10001", "刚刚店里把券退了，现在发了三张500的，发错了吧",
        )
        self.assertEqual("coupon_troubleshooting", mismatch["kind"])
        self.assertEqual("allow", mismatch["decision"])
        self.assertIn("发券", mismatch["reply"])
        self.assertNotIn("可用门店", mismatch["reply"])

    def test_minimal_order_followups_never_claim_backend_verification(self):
        paid = self.store.resolve_deterministic("10001", "已付款")
        self.assertEqual("aftersale_clarify", paid["kind"])
        self.assertIn("您反馈", paid["reply"])
        self.assertNotIn("已查询", paid["reply"])
        applied = self.store.resolve_deterministic("10001", "已申请")
        self.assertEqual("aftersale_clarify", applied["kind"])
        self.assertIn("无法直接核验", applied["reply"])
        manual = self.store.resolve_deterministic("10001", "人工")
        self.assertEqual("manual_handoff", manual["kind"])

    def test_atomic_coupon_facts_ground_composed_paid_and_face_value(self):
        self.store.save_v2_product(
            "10001", "同仁四季代金券",
            "100元代金券：售价54元，最多叠加2张",
        )
        result = self.store.resolve_deterministic("10001", "108抵200吗")
        self.assertEqual("voucher_value", result["kind"])
        self.assertIn("2张100元代金券", result["reply"])
        self.assertIn("共可抵扣200元", result["reply"])

    def test_compact_purchase_value_confirmation_uses_atomic_coupon_facts(self):
        self.store.save_v2_product(
            "10001", "同仁四季代金券",
            "100元代金券：售价54元，最多叠加2张",
        )
        result = self.store.resolve_deterministic("10001", "108拍下直接代200吗")
        self.assertEqual("voucher_value", result["kind"])
        self.assertEqual(
            "是的，100元代金券购买2张共支付108元，购买后发放2张100元代金券，共可抵扣200元。",
            result["reply"],
        )

    def test_non_store_reply_snapshots_remain_byte_for_byte_stable(self):
        self.store.save_v2_product(
            "10001", "同仁四季代金券",
            "100元代金券：售价54元，最多叠加2张。"
            "付款后发送美团电子券，到店出示券码核销。仅限堂食。",
        )
        expected = {
            "你好": (
                "greeting",
                "您好，请问想咨询当前商品的使用规则、适用门店还是发货问题？",
            ),
            "108拍下直接代200吗": (
                "voucher_value",
                "是的，100元代金券购买2张共支付108元，购买后发放2张100元代金券，共可抵扣200元。",
            ),
            "消费235咋拍": (
                "consumption_plan",
                "235元消费的话，可以购买2张100元代金券，共支付108元，可抵扣200元。\n\n"
                "剩余35元到店自行支付。\n\n当前同面额代金券每次最多使用2张。",
            ),
            "怎么发货": (
                "delivery_usage",
                "当前商品发放的是美团电子券，付款后发送领取信息，到店出示券码核销。",
            ),
            "一次可以用几张": (
                "stacking",
                "100元代金券售价54元，购买后发放100元代金券。"
                "每次最多使用2张100元券，共可抵扣200元；"
                "不同面额的券不能混用。",
            ),
            "人工": (
                "manual_handoff",
                "已切换至人工处理，请稍候。后续消息将保留给人工查看。",
            ),
        }
        actual = {
            message: (result["kind"], result["reply"])
            for message in expected
            if (result := self.store.resolve_deterministic("10001", message))
        }
        self.assertEqual(expected, actual)

    def test_colloquial_consumption_amount_builds_grounded_purchase_plan(self):
        self.store.save_v2_product(
            "10001", "同仁四季代金券",
            "100元代金券：售价54元，最多叠加2张",
        )
        result = self.store.resolve_deterministic("10001", "消费235咋拍")
        self.assertEqual("consumption_plan", result["kind"])
        self.assertIn("2张100元代金券", result["reply"])
        self.assertIn("共支付108元", result["reply"])
        self.assertIn("剩余35元到店自行支付", result["reply"])

    def test_bare_amount_and_store_usage_are_resolved_as_two_questions(self):
        self.store.save_v2_product(
            "10001", "同仁四季代金券",
            "100元代金券：售价54元，最多叠加2张",
        )
        self.store.import_store_text(
            "【广东省】\n【深圳】宝安壹方城店", "深圳门店", ["10001"],
        )
        result = self.store.resolve_deterministic(
            "10001", "深圳壹方城200可以用吗",
        )
        self.assertEqual("multi_intent", result["kind"])
        self.assertEqual(["store", "sku"], result["resolved_intents"])
        self.assertIn("宝安壹方城店", result["reply"])
        self.assertIn("2张100元代金券", result["reply"])
        self.assertNotIn("暂不支持议价", result["reply"])

    def test_discount_and_catalog_phrasings_use_real_options(self):
        self.store.save_v2_product(
            "10001", "测试代金券", "100元代金券：售价54元\n200元代金券：售价108元",
        )
        discount = self.store.resolve_deterministic("10001", "多少折")
        self.assertEqual("discount", discount["kind"])
        self.assertIn("约5.4折", discount["reply"])
        catalog = self.store.resolve_deterministic("10001", "代金券有哪些")
        self.assertEqual("coupon_catalog", catalog["kind"])
        self.assertIn("100元代金券", catalog["reply"])
        self.assertIn("200元代金券", catalog["reply"])

    def test_new_audience_query_drops_old_purchase_quantity(self):
        self.store.save_v2_product(
            "10001", "晚市自助",
            "晚市单人自助：售价80元\n晚市双人自助：售价162元",
        )
        first = self.store.resolve_deterministic("10001", "晚上两张多少钱")
        context = first["query_context_update"]
        second = self.store.resolve_deterministic(
            "10001", "晚市双人多少钱", store_context=context,
        )
        self.assertIn("162元", second["reply"])
        self.assertNotIn("324元", second["reply"])

    def test_unheaded_flavor_lines_are_extracted_without_price_noise(self):
        self.store.save_v2_product(
            "10001", "鱼酷烤鱼套餐",
            "贵州凯里酸汤牛肉烤鱼（不辣）\n经典蒜香烤鱼（不辣）\n"
            "青花椒烤鱼（微辣）\n套餐售价118元",
        )
        result = self.store.resolve_deterministic("10001", "有什么口味")
        self.assertEqual("product_attribute", result["kind"])
        self.assertIn("贵州凯里酸汤牛肉烤鱼（不辣）", result["reply"])
        self.assertIn("青花椒烤鱼（微辣）", result["reply"])
        self.assertNotIn("118", result["reply"])

    def test_store_entity_is_extracted_before_offer_intent_cleanup(self):
        self.store.import_store_list(
            self.create_multi_region_store_sheet(), "多地区门店", ["10001"]
        )
        result = self.store.resolve_deterministic("10001", "今天深圳什么优惠")
        self.assertEqual("stores", result["kind"])
        self.assertEqual("深圳", result["store_query"])
        self.assertIn("龙岗万科里店", result["reply"])
        self.assertIn("西丽益田假日里店", result["reply"])
        self.assertNotIn("广州万达广场店", result["reply"])
        self.assertNotIn("未查询到深圳什么优惠", result["reply"])

    def test_unique_store_homophone_is_matched_only_inside_current_product(self):
        self.store.import_store_list(
            self.create_multi_region_store_sheet(), "多地区门店", ["10001"]
        )
        result = self.store.resolve_deterministic("10001", "瘦光万达能用吗")
        self.assertEqual("stores", result["kind"])
        self.assertEqual("allow", result["decision"])
        self.assertIn("寿光万达广场店", result["reply"])
        self.assertTrue(result["reply"].startswith("可以用"))

    def test_composite_store_constraints_are_not_dropped(self):
        self.store.import_store_list(
            self.create_multi_region_store_sheet(), "多地区门店", ["10001"]
        )
        result = self.store.resolve_deterministic("10001", "寿光万达可以用么")
        self.assertEqual("stores", result["kind"])
        self.assertIn("寿光万达广场店", result["reply"])
        self.assertNotIn("广州万达广场店", result["reply"])
        self.assertEqual(
            "可以用，根据“寿光万达广场店”查询到可用门店：寿光万达广场店。",
            result["reply"],
        )

    def test_city_is_preserved_when_only_a_partial_landmark_is_typed(self):
        self.store.import_store_text(
            "【浙江省】\n【杭州】杭州城北万象城店\n"
            "【江苏省】\n【南京】南京万象城店",
            "万象城门店", ["10001"],
        )
        result = self.store.resolve_deterministic("10001", "杭州万象城能用吗")
        self.assertEqual("stores", result["kind"])
        self.assertEqual(["杭州城北万象城店"], [
            row["branch"] for row in result["store_matches"]
        ])
        self.assertNotIn("南京万象城店", result["reply"])

    def test_empty_generic_branch_key_never_matches_an_unrelated_store(self):
        self.store.import_store_text(
            "【河北省】\n【秦皇岛】广场店", "门店", ["10001"],
        )
        result = self.store.resolve_deterministic("10001", "鱼酷 任丘悦都汇店")
        self.assertEqual("stores", result["kind"])
        self.assertEqual([], result.get("store_matches") or [])
        self.assertNotIn("广场店可用", result["reply"])

    def test_unverified_store_qualifier_requires_confirmation(self):
        self.store.import_store_text(
            "【湖北省】\n【武汉】印象城店", "武汉门店", ["10001"],
        )
        result = self.store.resolve_deterministic("10001", "青山印象城店")
        self.assertEqual("stores_clarify", result["kind"])
        self.assertEqual("candidate_confirmation", result["store_status"])
        self.assertNotIn("可以使用", result["reply"])

    def test_city_plural_query_does_not_reuse_previous_single_store(self):
        self.store.import_store_text(
            "【湖北省】\n【武汉】武昌万象城店、汉口万象城店",
            "武汉门店", ["10001"],
        )
        previous = self.store.search_store("10001", "武昌万象城店")["matches"]
        result = self.store.resolve_deterministic(
            "10001", "武汉有那家店能用",
            store_context={
                "query": "武昌万象城店", "matches": previous,
                "status": "available", "verified": True,
            },
        )
        self.assertEqual("stores", result["kind"])
        self.assertIn("武昌万象城店", result["reply"])
        self.assertIn("汉口万象城店", result["reply"])

    def test_city_query_removes_colloquial_location_fillers(self):
        self.store.import_store_text(
            "【江苏省】\n【南京】新街口店、河西店", "南京门店", ["10001"],
        )
        self.assertEqual("南京", extract_store_query("南京这个点可以用吗"))
        result = self.store.resolve_deterministic("10001", "南京这个点可以用吗")
        self.assertEqual("stores", result["kind"])
        self.assertIn("新街口店", result["reply"])
        self.assertIn("河西店", result["reply"])

    def test_missing_voucher_plan_uses_one_fact_per_paragraph(self):
        self.store.save_v2_product(
            "10001", "100元代金券", "100元代金券：售价54.9元，最多叠加2张",
        )
        result = self.store.resolve_deterministic("10001", "有300元代金券吗")
        self.assertEqual(
            "当前商品没有300元代金券。\n\n"
            "300元消费的话，可以购买2张100元代金券，共支付109.8元，可抵扣200元。\n\n"
            "剩余100元到店自行支付。\n\n"
            "当前同面额代金券每次最多使用2张。",
            result["reply"],
        )

    def test_real_two_hundred_option_beats_cross_sku_coupon_combination(self):
        self.store.save_v2_product("10001", "周末代金券", "")
        self.store.save_ai_summary("10001", "周末券", {
            "skus": [{
                "name": "周末200元选项", "face_value": "200", "sale_price": "159",
                "applicable_time": "周末可用", "composition": "发2张100元券",
                "max_stack": 2, "stock": 8, "status": "active",
            }],
            "facts": {}, "time_rules": [],
        })
        result = self.store.resolve_deterministic("10001", "周末多少代200有吗")
        self.assertEqual("sku_availability", result["kind"])
        self.assertIn("周末200元选项售价159元", result["reply"])
        self.assertIn("2张100元代金券", result["reply"])
        self.assertIn("共可抵扣200元", result["reply"])
        self.assertIn("叠加使用", result["reply"])

    def test_explicit_zero_stock_sku_is_never_recommended(self):
        self.store.save_v2_product("10001", "周末代金券", "")
        self.store.save_ai_summary("10001", "周末券", {
            "skus": [{
                "name": "周末200元代金券", "face_value": "200", "sale_price": "159",
                "applicable_time": "周末可用", "composition": "发2张100元券",
                "stock": 0, "status": "sold_out",
            }],
            "facts": {}, "time_rules": [],
        })
        result = self.store.resolve_deterministic("10001", "周末200元代金券有吗")
        self.assertEqual("sku_availability", result["kind"])
        self.assertEqual("deny", result["decision"])
        self.assertIn("没有“200元”这一有货规格", result["reply"])
        self.assertIn("到店咨询", result["reply"])
        self.assertNotIn("售价159元", result["reply"])

    def test_audience_mix_uses_existing_tickets_without_requiring_price_word(self):
        self.store.save_v2_product(
            "10001", "家庭自助",
            "成人双人券：售价176元\n儿童票：售价68元，限身高1.2米至1.4米儿童",
        )
        result = self.store.resolve_deterministic("10001", "2大1小")
        self.assertEqual("audience_price", result["kind"])
        self.assertIn("176元", result["reply"])
        self.assertIn("68元", result["reply"])
        self.assertIn("合计244元", result["reply"])

    def test_exact_mixed_audience_package_wins_over_individual_tickets(self):
        self.store.save_v2_product(
            "10001", "家庭自助",
            "两大一小家庭套餐：售价218元\n成人票：售价99元\n儿童票：售价68元",
        )
        result = self.store.resolve_deterministic("10001", "2大1小多少钱")
        self.assertEqual("audience_price", result["kind"])
        self.assertIn("两大一小家庭套餐", result["reply"])
        self.assertIn("218元", result["reply"])
        self.assertNotIn("266元", result["reply"])

    def test_afternoon_tea_and_identity_are_both_hard_sku_filters(self):
        self.store.save_v2_product(
            "10001", "分时自助",
            "下午茶女士票：售价88元\n晚市女士票：售价118元",
        )
        result = self.store.resolve_deterministic("10001", "1位女士下午茶多少钱")
        self.assertEqual("audience_price", result["kind"])
        self.assertIn("88元", result["reply"])
        self.assertNotIn("118元", result["reply"])

    def test_child_height_tier_never_defaults_to_the_cheapest_sku(self):
        self.store.save_v2_product(
            "10001", "儿童自助",
            "1.2米以下儿童票：售价48元\n1.2米至1.4米儿童票：售价68元",
        )
        clarify = self.store.resolve_deterministic("10001", "儿童票多少钱")
        self.assertIn("请问儿童身高", clarify["reply"])
        matched = self.store.resolve_deterministic("10001", "1.3米儿童多少钱")
        self.assertIn("68元", matched["reply"])
        self.assertNotIn("48元", matched["reply"])

    def test_current_use_and_instant_use_do_not_enter_manual_aftersale(self):
        self.store.save_v2_product(
            "10001", "自助券", "100元代金券：售价79元，付款后自动发券",
        )
        self.store.replace_time_rules("10001", [{
            "label": "全天可用", "day_type": "any", "start_time": "00:00",
            "end_time": "23:59", "allowed": True, "reply": "全天可用。",
        }])
        current = self.store.resolve_deterministic("10001", "现在就能用吗")
        self.assertEqual("time", current["kind"])
        self.assertEqual("allow", current["decision"])
        self.assertNotIn("72小时", current["reply"])

        self.store.save_v2_product(
            "10001", "预约自助券", "100元代金券：售价79元，付款后自动发券，必须预约",
        )
        reserved = self.store.resolve_deterministic("10001", "买了马上能用吗")
        self.assertEqual("time", reserved["kind"])
        self.assertEqual("deny", reserved["decision"])
        self.assertIn("先预约", reserved["reply"])
        self.assertNotIn("72小时", reserved["reply"])


    def test_multi_sku_catalog_labels_time_and_hides_zero_stock(self):
        self.store.save_v2_product("10001", "半秋山100元代金券", "")
        self.store.save_ai_summary("10001", "分时券", {
            "skus": [
                {"name": "100元代金券", "face_value": "100", "sale_price": "66.8",
                 "applicable_time": "工作日可用", "stock": 5, "status": "active"},
                {"name": "100元代金券", "face_value": "100", "sale_price": "66.8",
                 "applicable_time": "周末可用", "stock": 3, "status": "active"},
                {"name": "100元代金券", "face_value": "100", "sale_price": "79.5",
                 "applicable_time": "法定节假日可用", "stock": 0, "status": "sold_out"},
            ], "facts": {}, "time_rules": [],
        })
        result = self.store.resolve_deterministic("10001", "代金券还有嘛")
        self.assertIn(result["kind"], {"coupon_catalog", "sku_availability"})
        self.assertIn("周末可用", result["reply"])
        self.assertNotIn("工作日可用", result["reply"])
        self.assertNotIn("79.5", result["reply"])

    def test_weekend_expansion_lists_only_in_stock_weekend_skus(self):
        self.store.save_v2_product("10001", "半秋山100元代金券", "")
        self.store.save_ai_summary("10001", "分时券", {
            "skus": [
                {"name": "100元代金券", "face_value": "100", "sale_price": "66.8",
                 "applicable_time": "工作日可用", "stock": 5},
                {"name": "100元代金券", "face_value": "100", "sale_price": "79.5",
                 "applicable_time": "周末可用", "stock": 2},
            ], "facts": {}, "time_rules": [],
        })
        result = self.store.resolve_deterministic("10001", "周末通用的还有吗")
        self.assertEqual("sku_availability", result["kind"])
        self.assertIn("周末可用", result["reply"])
        self.assertIn("79.5元", result["reply"])
        self.assertNotIn("66.8元", result["reply"])

    def test_catalog_display_feedback_repeats_time_labels(self):
        self.store.save_v2_product(
            "10001", "分时代金券",
            "100元代金券：售价66.8元，工作日可用\n100元代金券：售价79.5元，周末可用",
        )
        result = self.store.resolve_deterministic(
            "10001", "没看到", store_context={"last_sku_catalog": True},
        )
        self.assertEqual("coupon_catalog", result["kind"])
        self.assertIn("周末可用", result["reply"])
        self.assertNotIn("工作日可用", result["reply"])

    def test_generic_coupon_catalog_recommends_only_todays_sellable_sku(self):
        self.store.save_v2_product("10001", "半秋山100元代金券", "")
        self.store.save_ai_summary("10001", "分时券", {
            "skus": [
                {"name": "100元代金券", "face_value": "100", "sale_price": "66.8",
                 "applicable_time": "工作日可用", "stock": 5},
                {"name": "100元代金券", "face_value": "100", "sale_price": "79.5",
                 "applicable_time": "周末可用", "stock": 2},
                {"name": "100元代金券", "face_value": "100", "sale_price": "88",
                 "applicable_time": "周末可用", "stock": 0},
            ], "facts": {}, "time_rules": [],
        })
        with patch.object(V2Store, "_current_day_type", return_value="weekend"):
            result = self.store.resolve_deterministic("10001", "有什么券")
        self.assertEqual("coupon_catalog", result["kind"])
        self.assertIn("今天是周末", result["reply"])
        self.assertIn("79.5元", result["reply"])
        self.assertNotIn("66.8元", result["reply"])
        self.assertNotIn("88元", result["reply"])

    def test_date_question_names_matching_in_stock_sku_and_price(self):
        self.store.save_v2_product("10001", "半秋山100元代金券", "")
        self.store.save_ai_summary("10001", "分时券", {
            "skus": [
                {"name": "100元代金券", "face_value": "100", "sale_price": "66.8",
                 "applicable_time": "工作日可用", "stock": 5},
                {"name": "100元代金券", "face_value": "100", "sale_price": "79.5",
                 "applicable_time": "周末可用", "stock": 2},
            ], "facts": {}, "time_rules": [],
        })
        result = self.store.resolve_deterministic("10001", "9月6日星期天可以用吗")
        self.assertEqual("date_use", result["kind"])
        self.assertIn("周末可用", result["reply"])
        self.assertIn("79.5元", result["reply"])
        self.assertNotIn("66.8元", result["reply"])

    def test_price_change_operation_is_not_bargaining(self):
        result = self.store.resolve_deterministic("10001", "改个价")
        self.assertEqual("price_change_clarify", result["kind"])
        self.assertIn("改成多少元", result["reply"])

    def test_meal_period_statement_uses_real_matching_option(self):
        self.store.save_v2_product(
            "10001", "分时自助",
            "午餐单人自助：售价88元\n晚餐单人自助：售价118元",
        )
        result = self.store.resolve_deterministic("10001", "中午就餐")
        self.assertEqual("sku_availability", result["kind"])
        self.assertIn("88元", result["reply"])
        self.assertNotIn("118元", result["reply"])

    def test_tomorrow_evening_party_combines_compatible_buffet_skus(self):
        self.store.save_v2_product(
            "10001", "经典自助餐",
            "工作日午市单人经典自助：售价135元\n"
            "全周通用单人经典自助：售价157元\n"
            "工作日双人经典自助：售价297元",
        )
        result = self.store.resolve_deterministic("10001", "明晚三人多少钱")
        self.assertEqual("price", result["kind"])
        self.assertIn("明天晚餐3人购买建议", result["reply"])
        self.assertIn("工作日双人经典自助", result["reply"])
        self.assertIn("全周通用单人经典自助", result["reply"])
        self.assertIn("合计454元", result["reply"])
        self.assertNotIn("午市单人经典自助", result["reply"])
        self.assertNotIn("暂时无法准确报价", result["reply"])

    def test_in_store_purchase_question_is_not_a_store_name(self):
        result = self.store.resolve_deterministic("10001", "现在在门店，能买吗")
        self.assertEqual("stock", result["kind"])
        self.assertIn("当前商品仍在售", result["reply"])

    def test_store_query_strips_instant_words_from_branch(self):
        self.store.import_store_text(
            "【广东省】\n【深圳】宝安总店", "深圳门店", ["10001"],
        )
        self.assertEqual("宝安总店", extract_store_query("宝安总店现在马上能用吗"))
        result = self.store.resolve_deterministic("10001", "宝安总店现在马上能用吗")
        self.assertEqual("stores", result["kind"])
        self.assertIn("宝安总店", result["reply"])

    def test_brand_and_city_store_query_is_split(self):
        path = os.path.join(self.temp.name, "ikea-stores.xlsx")
        book = Workbook()
        sheet = book.active
        sheet.append(["店名", "分店名", "省", "市", "区/县", "地址"])
        sheet.append(["宜家", "武汉商场店", "湖北", "武汉", "硚口区", "张毕湖路"])
        sheet.append(["宜家", "深圳商场店", "广东", "深圳", "南山区", "北环大道"])
        book.save(path)
        self.store.import_store_list(path, "宜家门店", ["10001"])
        result = self.store.resolve_deterministic("10001", "宜家武汉")
        self.assertEqual("stores", result["kind"])
        self.assertIn("武汉商场店", result["reply"])
        self.assertNotIn("深圳商场店", result["reply"])


    def test_holiday_query_uses_holiday_sku_not_product_wide_weekday_denial(self):
        self.store.save_v2_product("10001", "半秋山100元代金券", "仅限工作日可用。")
        self.store.save_ai_summary("10001", "分时券", {
            "skus": [
                {"name": "100元代金券", "face_value": "100", "sale_price": "66.8",
                 "applicable_time": "仅限工作日可用", "stock": 5},
                {"name": "100元代金券", "face_value": "100", "sale_price": "79.5",
                 "applicable_time": "法定节假日可用", "stock": 3},
            ], "facts": {}, "time_rules": [],
        })
        result = self.store.resolve_deterministic("10001", "节假日能用吗")
        self.assertEqual("day_use", result["kind"])
        self.assertIn("可以", result["reply"])
        self.assertIn("法定节假日可用", result["reply"])
        self.assertIn("79.5元", result["reply"])
        self.assertNotIn("66.8元", result["reply"])

    def test_cannot_purchase_amount_is_an_issue_not_positive_availability(self):
        self.store.save_v2_product("10001", "半秋山100元代金券", "")
        self.store.save_ai_summary("10001", "分时券", {
            "skus": [
                {"name": "100元代金券", "face_value": "100", "sale_price": "66.8",
                 "applicable_time": "工作日可用", "stock": 5},
                {"name": "100元代金券", "face_value": "100", "sale_price": "79.5",
                 "applicable_time": "法定节假日可用", "stock": 3},
            ], "facts": {}, "time_rules": [],
        })
        result = self.store.resolve_deterministic("10001", "拍不了100")
        self.assertEqual("purchase_issue", result["kind"])
        self.assertIn("无法下单", result["reply"])
        self.assertIn("工作日可用", result["reply"])
        self.assertIn("法定节假日可用", result["reply"])
        self.assertNotEqual("当前售价66.8元/张，可拍。", result["reply"])


    def test_structured_store_entities_do_not_glue_area_brand_and_branch(self):
        path = os.path.join(self.temp.name, "structured-stores.xlsx")
        book = Workbook()
        sheet = book.active
        sheet.append(["店名", "分店名", "省", "市", "区/县", "地址"])
        sheet.append(["半秋山", "荆州万达店", "湖北", "荆州", "沙市区", "北京中路"])
        sheet.append(["半秋山", "武汉天地店", "湖北", "武汉", "江岸区", "卢沟桥路"])
        sheet.append(["宜家", "武汉商场店", "湖北", "武汉", "硚口区", "张毕湖路"])
        sheet.append(["宜家", "深圳商场店", "广东", "深圳", "南山区", "北环大道"])
        book.save(path)
        self.store.import_store_list(path, "结构化门店", ["10001"])

        self.assertEqual("荆州", extract_store_query("荆州的"))
        for message in ("荆州的", "荆州市的", "湖北省荆州市", "荆州有哪些店"):
            with self.subTest(message=message):
                result = self.store.resolve_deterministic("10001", message)
                self.assertEqual("stores", result["kind"])
                self.assertIn("荆州万达店", result["reply"])
                self.assertNotIn("荆州的", result["reply"])
                self.assertNotIn("武汉天地店", result["reply"])

        for message in ("宜家武汉", "武汉宜家", "武汉的宜家", "宜家在武汉"):
            with self.subTest(message=message):
                result = self.store.resolve_deterministic("10001", message)
                self.assertEqual("stores", result["kind"])
                self.assertIn("武汉商场店", result["reply"])
                self.assertNotIn("深圳商场店", result["reply"])

        branch = self.store.resolve_deterministic("10001", "湖北荆州半秋山万达店能用吗")
        self.assertEqual("stores", branch["kind"])
        self.assertIn("荆州万达店", branch["reply"])

    def test_store_short_name_matches_full_department_store_name(self):
        path = os.path.join(self.temp.name, "department-store.xlsx")
        book = Workbook()
        sheet = book.active
        sheet.append(["店名", "分店名", "省", "市"])
        sheet.append(["半秋山", "汉口大洋百货店", "湖北", "武汉"])
        book.save(path)
        self.store.import_store_list(path, "半秋山门店", ["10001"])
        result = self.store.resolve_deterministic("10001", "半秋山汉口大洋店")
        self.assertEqual("stores", result["kind"])
        self.assertIn("汉口大洋百货店", result["reply"])

    def test_discount_confirmation_uses_today_sellable_sku_only(self):
        self.store.save_v2_product("10001", "半秋山100元代金券", "")
        self.store.save_ai_summary("10001", "分时券", {
            "skus": [
                {"name": "100元代金券", "face_value": "100", "sale_price": "66.8",
                 "applicable_time": "工作日可用", "stock": 5},
                {"name": "100元代金券", "face_value": "100", "sale_price": "79.5",
                 "applicable_time": "周末可用", "stock": 2},
                {"name": "100元代金券", "face_value": "100", "sale_price": "88",
                 "applicable_time": "周末可用", "stock": 0},
            ], "facts": {}, "time_rules": [],
        })
        with patch.object(V2Store, "_current_day_type", return_value="weekend"):
            result = self.store.resolve_deterministic("10001", "今天是7.95折吗")
        self.assertEqual("discount", result["kind"])
        self.assertIn("是的", result["reply"])
        self.assertIn("79.5元", result["reply"])
        self.assertIn("7.95折", result["reply"])
        self.assertNotIn("66.8元", result["reply"])
        self.assertNotIn("88元", result["reply"])

    def test_date_context_is_inherited_by_following_store_name(self):
        self.store.save_v2_product("10001", "半秋山100元代金券", "")
        self.store.save_ai_summary("10001", "分时券", {
            "skus": [
                {"name": "工作日100元代金券", "face_value": "100", "sale_price": "66.8",
                 "applicable_time": "工作日可用", "stock": 5},
                {"name": "周末100元代金券", "face_value": "100", "sale_price": "79.5",
                 "applicable_time": "周末可用", "stock": 2},
            ], "facts": {}, "time_rules": [],
        })
        self.store.import_store_text(
            "【广东省】\n【深圳】世界之窗广场店", "深圳门店", ["10001"],
        )
        first = self.store.resolve_deterministic("10001", "今天可以用吗")
        context = first.get("query_context_update") or {}
        self.assertEqual("weekend", context.get("pending_day_type"))
        second = self.store.resolve_deterministic(
            "10001", "深圳世界之窗广场店", store_context=context,
        )
        self.assertEqual("stores", second["kind"])
        self.assertIn("世界之窗广场店", second["reply"])
        self.assertEqual("周末100元代金券", second["query_context_update"]["selected_sku_name"])

    def test_where_to_buy_uses_current_listing_and_today_option(self):
        self.store.save_v2_product("10001", "半秋山100元代金券", "")
        self.store.save_ai_summary("10001", "分时券", {
            "skus": [
                {"name": "100元代金券", "face_value": "100", "sale_price": "66.8",
                 "applicable_time": "工作日可用", "stock": 5},
                {"name": "100元代金券", "face_value": "100", "sale_price": "79.5",
                 "applicable_time": "周末可用", "stock": 2},
            ], "facts": {}, "time_rules": [],
        })
        with patch.object(V2Store, "_current_day_type", return_value="weekend"):
            result = self.store.resolve_deterministic("10001", "哪里买")
        self.assertEqual("purchase_flow", result["kind"])
        self.assertIn("当前商品页面", result["reply"])
        self.assertIn("79.5元", result["reply"])
        self.assertNotIn("66.8元", result["reply"])

    def test_spoken_face_value_price_uses_today_stock_and_deduplicates(self):
        self.store.save_v2_product(
            "10001", "半秋山100元代金券",
            "100元代金券：售价66.8元，工作日可用\n"
            "100元代金券：售价79.5元，周末可用",
        )
        self.store.save_ai_summary("10001", "分时券", {
            "skus": [
                {"name": "100元代金券", "face_value": "100", "sale_price": "66.8",
                 "applicable_time": "工作日可用", "stock": 0},
                {"name": "100元代金券", "face_value": "100", "sale_price": "79.5",
                 "applicable_time": "周末可用", "stock": 3},
                {"name": "半秋山100元券", "face_value": "100", "sale_price": "79.5",
                 "applicable_time": "周末可用", "stock": 3},
            ], "facts": {}, "time_rules": [],
        })
        with patch.object(V2Store, "_current_day_type", return_value="weekend"):
            result = self.store.resolve_deterministic("10001", "100 的券，多少钱")
        self.assertEqual("price", result["kind"])
        self.assertEqual(1, result["reply"].count("79.5元"))
        self.assertNotIn("66.8元", result["reply"])
        self.assertNotIn("没有“100 的券”", result["reply"])

    def test_zero_stock_structured_sku_is_not_revived_by_raw_text(self):
        self.store.save_v2_product(
            "10001", "半秋山100元代金券",
            "100元代金券：售价66.8元，工作日可用",
        )
        self.store.save_ai_summary("10001", "库存", {
            "skus": [{
                "name": "100元代金券", "face_value": "100", "sale_price": "66.8",
                "applicable_time": "工作日可用", "stock": "0",
            }], "facts": {}, "time_rules": [],
        })
        product = self.store.get_v2_product("10001")
        self.assertEqual([], self.store.extract_product_options(product))

    def test_matching_manual_coupon_text_respects_structured_stock(self):
        for stock in (0, "0", 3):
            with self.subTest(stock=stock):
                self.store.save_v2_product(
                    "10001", "100元代金券",
                    "100元代金券：售价79.5元，周末可用",
                )
                self.store.save_ai_summary("10001", "库存", {
                    "skus": [{
                        "name": "100元代金券", "face_value": "100",
                        "sale_price": "79.5", "applicable_time": "周末",
                        "stock": stock,
                    }], "facts": {}, "time_rules": [],
                })
                with patch.object(V2Store, "_current_day_type", return_value="weekend"):
                    result = self.store.resolve_deterministic("10001", "100元代金券有吗")
                self.assertEqual("sku_availability", result["kind"])
                if str(stock) == "0":
                    self.assertEqual("deny", result["decision"])
                    self.assertNotIn("有的", result["reply"])
                    self.assertNotIn("售价79.5元", result["reply"])
                else:
                    self.assertEqual("allow", result["decision"])
                    self.assertIn("售价79.5元", result["reply"])

    def test_redemption_requests_use_in_stock_stackable_sku(self):
        self.store.save_v2_product("10001", "同仁四季代金券", "仅支持同面额叠加，一桌最多代300元")
        self.store.save_ai_summary("10001", "", {
            "skus": [{"name": "100元代金券", "face_value": "100",
                      "sale_price": "79.5", "stock": 5}],
        })
        for message in ("200", "200元", "有没有200的", "200有吗", "有200元代金券吗", "没有200的了？"):
            with self.subTest(message=message):
                result = self.store.resolve_deterministic("10001", message)
                self.assertEqual("redemption_plan", result["kind"])
                self.assertIn("2张100元代金券", result["reply"])
                self.assertIn("159元", result["reply"])
                self.assertIn("可抵扣200元", result["reply"])
                self.assertNotIn("没有200元代金券", result["reply"])
        blocked = self.store.resolve_deterministic("10001", "400")
        self.assertEqual("deny", blocked["decision"])

    def test_redemption_plan_does_not_treat_missing_legacy_stock_as_zero(self):
        self.store.save_v2_product(
            "10001", "同仁四季代金券",
            "100元代金券：售价79.5元，最多使用2张；仅支持同面额叠加",
        )
        result = self.store.resolve_deterministic("10001", "200")
        self.assertEqual("redemption_plan", result["kind"])
        self.assertEqual("allow", result["decision"])
        self.assertIn("2张100元代金券", result["reply"])
        self.assertIn("159元", result["reply"])

    def test_bare_amount_accepts_composed_sku_when_only_mixed_faces_are_forbidden(self):
        self.store.save_v2_product(
            "10001", "同仁四季代金券",
            "200元代金券：售价111.8元，发100元券2张，最多使用2张。"
            "仅支持同面额代金券叠加，不同面额代金券不能叠加。",
        )
        result = self.store.resolve_deterministic("10001", "200")
        self.assertEqual("redemption_plan", result["kind"])
        self.assertEqual("allow", result["decision"])
        self.assertIn("2张100元代金券", result["reply"])
        self.assertIn("111.8元", result["reply"])
        self.assertIn("可抵扣200元", result["reply"])

    def test_condition_only_people_day_and_meal_defaults_to_sku_price(self):
        self.store.save_v2_product(
            "10001", "奶糖爸爸自助小火锅",
            "单人自助：48.9元。平日周末全天通用，几人拍几张。",
        )
        self.store.save_ai_summary("10001", "单人自助", {
            "sale_options": [{
                "name": "奶糖爸爸自助小火锅单人自助",
                "sale_price": "48.9", "people_counts": [1],
                "applicable_time": "平日周末全天通用",
            }],
        })
        for message in ("三人明天中午", "明天中午三人"):
            with self.subTest(message=message):
                result = self.store.resolve_deterministic("10001", message)
                self.assertEqual("price", result["kind"])
                self.assertIn("3人需要购买3份", result["reply"])
                self.assertIn("单价48.9元", result["reply"])
                self.assertIn("共146.7元", result["reply"])

    def test_location_and_condition_statement_combines_store_and_sku_price(self):
        self.store.save_v2_product(
            "10001", "奶糖爸爸自助小火锅",
            "单人自助：48.9元。平日周末全天通用，几人拍几张。",
        )
        self.store.save_ai_summary("10001", "单人自助", {
            "sale_options": [{
                "name": "奶糖爸爸自助小火锅单人自助",
                "sale_price": "48.9", "people_counts": [1],
                "applicable_time": "平日周末全天通用",
            }],
        })
        self.store.import_store_text(
            "【广东省】\n【深圳】五和店", "深圳门店", ["10001"],
        )
        result = self.store.resolve_deterministic(
            "10001", "深圳五和店三人明天中午",
        )
        self.assertEqual("multi_intent", result["kind"])
        self.assertEqual(["store", "conditions"], result["resolved_intents"])
        self.assertIn("五和店", result["reply"])
        self.assertIn("3人需要购买3份", result["reply"])
        self.assertIn("共146.7元", result["reply"])
        self.assertIn("\n\n", result["reply"])
        self.assertNotIn("1. 适用门店", result["reply"])

    def test_composed_sku_quantity_reply_uses_delivered_coupon_units(self):
        self.store.save_v2_product(
            "10001", "同仁四季代金券",
            "200元代金券：售价111.8元，发100元券2张，最多使用2张。"
            "仅支持同面额代金券叠加，不同面额不能混用。",
        )
        result = self.store.resolve_deterministic("10001", "200可以用几张")
        self.assertEqual("stacking", result["kind"])
        self.assertIn("售价111.8元", result["reply"])
        self.assertIn("发放2张100元代金券", result["reply"])
        self.assertIn("合计抵扣200元", result["reply"])
        self.assertIn("最多使用2张100元券", result["reply"])
        self.assertNotIn("200元代金券可以叠加", result["reply"])

    def test_colloquial_max_quantity_prefers_unit_rule_over_bad_structured_limit(self):
        self.store.save_v2_product(
            "10001", "同仁四季代金券",
            "200元代金券：售价116元，发100元券2张。"
            "100元代金券最多使用2张，仅支持同面额叠加。",
        )
        self.store.save_ai_summary("10001", "结构化规格", {
            "products": [{
                "name": "同仁四季200元代金券", "option_type": "voucher",
                "face_value": "200", "sale_price": "116",
                "composition": "100元券2张", "max_stack": "1", "stock": "10",
            }],
        })

        result = self.store.resolve_deterministic("10001", "200最多用几张")

        self.assertEqual("stacking", result["kind"])
        self.assertIn("发放2张100元代金券", result["reply"])
        self.assertIn("这2张可以同一次使用", result["reply"])
        self.assertIn("最多使用2张100元券", result["reply"])
        self.assertNotIn("最多使用1张", result["reply"])

    def test_usage_subject_reads_product_rule_instead_of_store_index(self):
        self.store.save_v2_product(
            "10001", "火锅代金券",
            "100元代金券：售价76元。当前代金券不可用于锅底和酒水。",
        )
        result = self.store.resolve_deterministic("10001", "锅底能用吗")
        self.assertEqual("usage_scope", result["kind"])
        self.assertIn("不可用于锅底和酒水", result["reply"])
        self.assertNotIn("适用门店资料", result["reply"])

    def test_unavailable_dishes_are_not_misread_as_coupon_exclusion(self):
        self.store.save_v2_product(
            "10001", "餐饮代金券",
            "100元代金券：售价65元。部分菜品因时令、售罄或其他不可抗因素无法提供时，"
            "可以联系商家协商处理。",
        )
        result = self.store.resolve_deterministic("10001", "菜品都能用是吧")
        self.assertEqual("usage_scope", result["kind"])
        self.assertNotIn("不可用于菜品", result["reply"])
        self.assertIn("部分菜品可能因时令、售罄", result["reply"])

    def test_specific_dish_exception_is_not_widened_to_all_dishes(self):
        self.store.save_v2_product(
            "10001", "餐饮代金券",
            "100元代金券：售价65元。除酒水饮料外全场通用。"
            "使用范围例外：郑州市正弘店‘粉丝蒸松叶蟹’菜品不可用。",
        )

        result = self.store.resolve_deterministic("10001", "哪些菜品可以用")

        self.assertEqual("usage_scope", result["kind"])
        self.assertIn("除酒水饮料外，其他菜品可以使用", result["reply"])
        self.assertIn("粉丝蒸松叶蟹", result["reply"])
        self.assertNotIn("当前代金券不可用于菜品", result["reply"])

    def test_coupon_multiplier_in_all_store_question_is_not_a_two_yuan_sku(self):
        self.store.save_v2_product(
            "10001", "同仁四季椰子鸡代金券",
            "100元代金券：售价57.9元，发100元券1张。"
            "300元代金券：售价212.8元，发300元券1张。",
        )

        result = self.store.resolve_deterministic(
            "10001", "你好请问100x2的券全国门店通用吗",
        )

        self.assertEqual("stores_scope", result["kind"])
        self.assertIn("100元代金券（购买2张）", result["reply"])
        self.assertIn("不是全国所有门店通用", result["reply"])
        self.assertNotIn("2元代金券", result["reply"])

    def test_platform_skus_recover_missing_sale_prices_for_store_configuration(self):
        product = {
            "title": "蚂蚁洞烤肉代金券",
            "raw_text": (
                "美团200元代金券（最多叠加2张）\n"
                "抖音50元代金券\n"
                "抖音200元代金券（最多叠加4张）"
            ),
            "structured": {"products": [
                {"name": "美团200元代金券", "face_value": "200", "sale_price": ""},
                {
                    "name": "抖音50元代金券", "face_value": "50",
                    "sale_price": "129", "max_stack": "4",
                },
                {"name": "抖音200元代金券", "face_value": "200", "sale_price": ""},
            ]},
            "platform_summary": json.dumps({
                "title": "蚂蚁洞烤肉代金券",
                "description": "不同卡券适用门店不同",
                "sku": [
                    {"priceInCent": 11980, "quantity": 200, "propertyList": [
                        {"actualValueText": "美团200（可叠加2张）"}
                    ]},
                    {"priceInCent": 12900, "quantity": 200, "propertyList": [
                        {"actualValueText": "抖音50x4"}
                    ]},
                    {"priceInCent": 13800, "quantity": 500, "propertyList": [
                        {"actualValueText": "抖音200（可叠加4张）"}
                    ]},
                    {"priceInCent": 5600, "quantity": 200, "propertyList": [
                        {"actualValueText": "防下架勿拍"}
                    ]},
                ],
            }, ensure_ascii=False),
            "manual_edited": True,
        }

        skus = self.store.list_product_skus("platform-sku", product)
        by_name = {sku["sku_name"]: sku for sku in skus}

        self.assertEqual(3, len(skus))
        self.assertEqual("119.8", by_name["美团200（可叠加2张）"]["sale_price"])
        self.assertEqual("2", by_name["美团200（可叠加2张）"]["max_stack"])
        self.assertEqual("138", by_name["抖音200（可叠加4张）"]["sale_price"])
        self.assertEqual("4", by_name["抖音200（可叠加4张）"]["max_stack"])
        self.assertEqual("50元券4张", by_name["抖音50x4"]["composition"])
        self.assertEqual("", by_name["抖音50x4"]["max_stack"])
        self.assertEqual(
            "抖音50x4（售价129元）",
            self.store._sku_public_label(by_name["抖音50x4"]),
        )
        self.assertNotIn(
            "可叠加4张",
            self.store._sku_public_label(by_name["抖音50x4"]),
        )
        self.assertNotIn("防下架勿拍", by_name)

    def test_platform_and_face_keep_distinct_stack_limits(self):
        product = {
            "title": "双平台代金券",
            "raw_text": (
                "美团200元代金券最多叠加2张。\n"
                "抖音200元代金券最多叠加4张。\n"
                "仅支持同一规格代金券叠加，不同规格不能混用。"
            ),
            "structured": {"products": [
                {
                    "name": "美团200元代金券", "face_value": "200",
                    "sale_price": "119.8", "max_stack": "2",
                },
                {
                    "name": "抖音200元代金券", "face_value": "200",
                    "sale_price": "138", "max_stack": "4",
                },
            ]},
        }
        self.assertEqual("", self.store.benefit_combination_reply(
            product, "美团200元代金券可以叠加吗",
        ))
        meituan = self.store.stacking_reply(product, "美团200元代金券能用几张")
        self.assertIn("美团200元代金券", meituan)
        self.assertIn("最多使用2张", meituan)
        self.assertNotIn("抖音200元代金券", meituan)
        douyin = self.store.stacking_reply(product, "抖音200元代金券能用几张")
        self.assertIn("抖音200元代金券", douyin)
        self.assertIn("最多使用4张", douyin)
        bare = self.store.stacking_reply(product, "200元代金券能用几张")
        self.assertIn("美团200元代金券每次最多使用2张", bare)
        self.assertIn("抖音200元代金券每次最多使用4张", bare)
        self.assertIn("请确认需要哪一种规格", bare)
        summary = self.store.build_knowledge_summary(product)
        self.assertIn("美团200元代金券最多使用2张", summary)
        self.assertIn("抖音200元代金券最多使用4张", summary)
        matches = self.store.match_message_skus("10001", "美团200元", product)
        self.assertEqual(["美团200元代金券"], [row["sku_name"] for row in matches])

    def test_quantity_sku_collapse_does_not_invent_stack_limit(self):
        product = {
            "title": "100元代金券",
            "structured": {"products": [
                {
                    "name": "100元代金券", "face_value": "100",
                    "sale_price": "68", "max_stack": "同面额可叠加",
                },
                {
                    "name": "100元代金券x2", "face_value": "100",
                    "sale_price": "136", "composition": "100元券2张",
                    "max_stack": "同面额可叠加",
                },
            ]},
        }

        options = self.store.extract_product_options(product)

        self.assertEqual(1, len(options))
        self.assertEqual("", options[0]["max_stack"])
        self.assertEqual("4", self.store._normalize_stack_limit("最多叠加4张"))
        self.assertEqual("", self.store._normalize_stack_limit("同面额可叠加"))

    def test_elliptical_amount_followup_uses_consumption_plan(self):
        self.store.save_v2_product(
            "10001", "餐饮代金券",
            "100元代金券：售价65元，发100元券1张。支持同面额代金券叠加，不限制数量。",
        )
        result = self.store.resolve_deterministic("10001", "200的呢")
        self.assertEqual("redemption_plan", result["kind"])
        self.assertIn("购买2张100元代金券", result["reply"])
        self.assertIn("共支付130元", result["reply"])
        self.assertNotIn("没有200元代金券", result["reply"])

    def test_buyer_prefixed_amount_uses_consumption_plan_instead_of_offer(self):
        self.store.save_v2_product(
            "10001", "餐饮代金券",
            "100元代金券：售价65元，发100元券1张。支持同面额代金券叠加，最多叠加5张。",
        )
        result = self.store.resolve_deterministic("10001", "我426元")
        self.assertEqual("consumption_plan", result["kind"])
        self.assertIn("购买4张100元代金券", result["reply"])
        self.assertIn("共支付260元", result["reply"])
        self.assertIn("剩余26元到店自行支付", result["reply"])
        self.assertNotIn("出价", result["reply"])

    def test_amount_over_stack_limit_uses_closest_valid_sku(self):
        self.store.save_v2_product(
            "10001", "多规格代金券",
            "100元代金券：售价85元（可叠加4张）\n"
            "300元代金券：售价235元（可叠加2张）\n"
            "400元代金券：售价340元，发100元券4张\n"
            "仅支持同面额代金券叠加。100元券最多叠加4张；300元券最多叠加2张。",
        )
        result = self.store.resolve_deterministic("10001", "500")
        self.assertEqual("consumption_plan", result["kind"])
        self.assertIn("共支付340元", result["reply"])
        self.assertIn("可抵扣400元", result["reply"])
        self.assertIn("剩余100元到店自行支付", result["reply"])
        self.assertNotIn("无法组成", result["reply"])

    def test_global_scope_allows_only_items_outside_explicit_exclusions(self):
        self.store.save_v2_product(
            "10001", "火锅代金券",
            "100元代金券：售价76元。除锅底外，全场通用。",
        )
        beverage = self.store.resolve_deterministic("10001", "饮料可以用吗")
        self.assertEqual("usage_scope", beverage["kind"])
        self.assertIn("饮料消费可以使用", beverage["reply"])
        self.assertIn("除锅底外全场通用", beverage["reply"])
        pot = self.store.resolve_deterministic("10001", "锅底可以用吗")
        self.assertIn("不可用于锅底", pot["reply"])

    def test_usage_scope_falls_back_to_complete_marketplace_description(self):
        product = {
            "raw_text": "【商品规格】\n100元代金券：售价76元。",
            "ai_summary": "酒水饮料不可用。",
            "platform_summary": json.dumps({
                "description": "除酒水饮料外全场通用，不可使用包间。仅限堂食。无需预约。"
            }, ensure_ascii=False),
            "structured": {"facts": {"酒水限制": "酒水饮料不可用"}},
        }
        pot = self.store.usage_subject_reply(product, "锅底可以用吗")
        self.assertIn("锅底消费可以使用", pot["reply"])
        self.assertIn("除酒水饮料外全场通用", pot["reply"])
        beverage = self.store.usage_subject_reply(product, "饮料可以用吗")
        self.assertIn("不可用于饮料", beverage["reply"])

    def test_elliptical_pot_base_questions_use_current_product_rules(self):
        product = {
            "raw_text": "200元代金券：售价138元。除酒水外全场通用。",
            "ai_summary": "",
            "structured": {},
        }
        for message in ("包括锅底吗", "锅底呢", "锅底吗"):
            with self.subTest(message=message):
                result = self.store.usage_subject_reply(product, message)
                self.assertEqual("usage_scope", result["kind"])
                self.assertEqual("allow", result["decision"])
                self.assertIn("锅底消费可以使用", result["reply"])
                self.assertIn("除酒水外全场通用", result["reply"])
                self.assertNotIn("锅底需另付", result["reply"])

    def test_elliptical_pot_base_question_respects_current_product_exclusion(self):
        product = {
            "raw_text": "100元代金券：售价76元。锅底不可用，其余菜品酒水全场通用。",
            "ai_summary": "",
            "structured": {},
        }
        result = self.store.usage_subject_reply(product, "锅底呢")
        self.assertEqual("usage_scope", result["kind"])
        self.assertEqual("allow", result["decision"])
        self.assertIn("不可用于锅底", result["reply"])

    def test_elliptical_pot_base_question_never_invents_unknown_rule(self):
        product = {
            "raw_text": "100元代金券：售价76元。",
            "ai_summary": "",
            "structured": {},
        }
        result = self.store.usage_subject_reply(product, "锅底呢")
        self.assertEqual("usage_scope", result["kind"])
        self.assertEqual("allow", result["decision"])
        self.assertIn("没有明确说明锅底是否可用", result["reply"])

    def test_incomplete_sync_preserves_last_complete_platform_payload(self):
        complete = json.dumps({
            "title": "双平台券",
            "description": "完整使用规则",
            "price": "119.8",
            "sku": [{
                "priceInCent": 11980,
                "propertyList": [{"actualValueText": "美团200（可叠加2张）"}],
            }],
        }, ensure_ascii=False)
        self.store.upsert_synced_product({
            "item_id": "sync-sku", "title": "双平台券",
            "platform_summary": complete, "image_urls": [], "price": "119.8",
        })

        self.store.upsert_synced_product({
            "item_id": "sync-sku", "title": "双平台券新标题",
            "platform_summary": json.dumps({
                "title": "双平台券新标题", "description": "",
                "price": "120", "stock": None, "sku": [],
            }, ensure_ascii=False),
            "image_urls": [], "price": "120",
        })

        payload = json.loads(self.store.get_v2_product("sync-sku")["platform_summary"])
        self.assertEqual("双平台券新标题", payload["title"])
        self.assertEqual("120", payload["price"])
        self.assertEqual("完整使用规则", payload["description"])
        self.assertEqual(1, len(payload["sku"]))

    def test_synced_summary_preserves_page_rules_omitted_by_ai(self):
        platform_summary = json.dumps({
            "title": "代金券",
            "description": (
                "【使用规则】\n1. 除酒水饮料外全场通用，不可使用包间。\n"
                "2. 仅限堂食。\n3. 无需预约，高峰期可能需要等位。"
            ),
            "sku": [],
        }, ensure_ascii=False)
        self.store.upsert_synced_product({
            "item_id": "sync-rules", "title": "代金券",
            "platform_summary": platform_summary, "image_urls": [], "price": "76",
        })
        product = self.store.save_synced_summary(
            "sync-rules", "AI只保留了售价", {"facts": {}, "time_rules": []},
        )
        for expected in ("除酒水饮料外全场通用", "不可使用包间", "仅限堂食", "无需预约"):
            self.assertIn(expected, product["raw_text"])

    def test_exclusion_without_positive_scope_cannot_prove_other_items_are_usable(self):
        product = {
            "raw_text": "锅底除外。",
            "ai_summary": "", "structured": {},
        }
        result = self.store.usage_subject_reply(product, "饮料可以用吗")
        self.assertIn("没有明确说明", result["reply"])

    def test_usage_scope_does_not_invent_a_missing_sku(self):
        product = {
            "raw_text": "100元代金券：售价76元。除锅底外全场通用。",
            "ai_summary": "", "structured": {},
        }
        result = self.store.usage_subject_reply(product, "有饮料这个规格并且可以用吗")
        self.assertEqual("sku_availability", result["kind"])
        self.assertIn("没有“饮料”这一在售规格", result["reply"])

    def test_package_and_coupon_default_to_no_combination_without_positive_rule(self):
        self.store.save_v2_product(
            "10001", "代金券", "100元代金券：售价76元。除锅底外全场通用。",
        )
        denied = self.store.resolve_deterministic("10001", "套餐可以用代金券吗")
        self.assertEqual("benefit_combination", denied["kind"])
        self.assertIn("不能一起使用", denied["reply"])
        self.store.save_ai_summary("10001", "明确叠加规则", {
            "facts": {"使用规则": "套餐可与代金券叠加使用"},
        })
        allowed = self.store.resolve_deterministic("10001", "套餐可以用代金券吗")
        self.assertIn("明确说明", allowed["reply"])
        self.assertIn("可以和套餐一起使用", allowed["reply"])

    def test_weight_decimal_is_not_a_date_and_holiday_exclusion_is_respected(self):
        self.store.save_v2_product(
            "10001", "鱼酷烤鱼",
            "鱼酷2.6斤回鱼单鱼套餐：售价118元。"
            "除中秋节（9.25-9.27）、国庆节（10.1-10.7）外，营业时间内可用。",
        )
        self.assertIsNone(self.store._query_date("鱼酷2.6斤回鱼单鱼套餐"))
        result = self.store.resolve_deterministic("10001", "中秋可以用吗")
        self.assertEqual("date_use", result["kind"])
        self.assertIn("不能使用", result["reply"])

    def test_holiday_reads_structured_blackout_and_is_not_duplicated_as_store(self):
        self.store.save_v2_product(
            "10001", "大树餐厅代金券", "100元代金券：售价68元。",
        )
        self.store.save_ai_summary("10001", "节日规则", {
            "facts": {"不可用日期": "中秋节（9.25-9.27）、国庆节（10.1-10.7）不可用"},
        })
        for question in ("中秋可以用吗", "中秋期间可以使用吗", "中秋节可以用吗"):
            with self.subTest(question=question):
                result = self.store.resolve_deterministic("10001", question)
                self.assertEqual("date_use", result["kind"])
                self.assertIn("不能使用", result["reply"])
                self.assertNotIn("适用门店：", result["reply"])

    def test_workday_only_sku_does_not_claim_holiday_availability(self):
        self.store.save_v2_product(
            "10001", "工作日代金券",
            "100元代金券：售价66.8元，仅限工作日可用。",
        )
        result = self.store.resolve_deterministic("10001", "中秋节可以用吗")
        self.assertEqual("date_use", result["kind"])
        self.assertIn("没有可用的商品规格", result["reply"])

    def test_bare_number_without_matching_sale_price_becomes_consumption_plan(self):
        self.store.save_v2_product(
            "10001", "测试代金券",
            "200元代金券：售价124.9元，发200元券1张。"
            "200元代金券最多使用2张。",
        )
        result = self.store.resolve_deterministic("10001", "200")
        self.assertEqual("redemption_plan", result["kind"])
        self.assertIn("124.9元", result["reply"])
        self.assertNotIn("售价2元", result["reply"])

    def test_bare_number_matches_sku_sale_price_ignoring_decimal_part(self):
        self.store.save_v2_product(
            "10001", "测试代金券",
            "300元代金券：售价208.8元，发300元券1张。",
        )
        result = self.store.resolve_deterministic("10001", "208")
        self.assertEqual("sku_price", result["kind"])
        self.assertIn("300元代金券售价208.8元", result["reply"])

    def test_bare_number_without_sale_price_match_uses_non_overage_plan(self):
        self.store.save_v2_product(
            "10001", "测试代金券",
            "100元代金券：售价54元，最多叠加3张。",
        )
        result = self.store.resolve_deterministic("10001", "208")
        self.assertEqual("consumption_plan", result["kind"])
        self.assertIn("2张100元代金券", result["reply"])
        self.assertIn("剩余8元", result["reply"])
        self.assertNotIn("3张100元代金券", result["reply"])

    def test_discounted_total_question_uses_non_overage_plan_and_totals(self):
        self.store.save_v2_product(
            "10001", "同仁四季代金券",
            "200元代金券：售价113.8元，发100元券2张。"
            "300元代金券：售价204.8元，发300元券1张。"
            "仅支持同面额代金券叠加。100元代金券最多使用2张；300元代金券最多使用1张。",
        )
        result = self.store.resolve_deterministic("10001", "273优惠完多少")
        self.assertEqual("consumption_plan", result["kind"])
        self.assertEqual("allow", result["decision"])
        self.assertIn("可抵扣200元", result["reply"])
        self.assertIn("剩余73元", result["reply"])
        self.assertIn("合计实际支出186.8元", result["reply"])
        self.assertIn("共优惠86.2元", result["reply"])
        self.assertNotIn("300元代金券", result["reply"])

    def test_mixed_audience_result_is_split_into_paragraphs(self):
        self.store.save_v2_product(
            "10001", "奶糖爸爸自助小火锅",
            "单人自助：48.9元。几人拍几张。儿童价格以门店为准。",
        )
        self.store.save_ai_summary("10001", "单人自助", {
            "sale_options": [{
                "name": "奶糖爸爸自助小火锅单人自助",
                "sale_price": "48.9", "people_counts": [1],
            }],
        })
        result = self.store.resolve_deterministic(
            "10001", "2位成人1位儿童多少钱",
        )
        self.assertEqual("audience_price", result["kind"])
        self.assertIn("2位成人需要购买2张", result["reply"])
        self.assertIn("共97.8元。\n\n当前商品暂时没有“儿童票”", result["reply"])
        self.assertNotIn("您好，本店目前没有", result["reply"])

    def test_colloquial_purchase_timing_and_store_are_answered_separately(self):
        self.store.save_v2_product(
            "10001", "同仁四季代金券",
            "100元代金券：售价79.5元。请当天购买、当天使用。",
        )
        self.store.import_store_text(
            "【广东省】\n【深圳】丹竹头店", "深圳门店", ["10001"],
        )
        result = self.store.resolve_deterministic(
            "10001", "吃完再买对吧 丹竹头店可以用吗",
        )
        self.assertEqual("multi_intent", result["kind"])
        self.assertEqual(["purchase_timing", "store"], result["resolved_intents"])
        self.assertEqual("丹竹头店", result["store_query"])
        self.assertIn("用餐结束后、结账前购买", result["reply"])
        self.assertIn("丹竹头店", result["reply"])

    def test_store_name_with_bare_price_word_resolves_store_and_price(self):
        self.store.save_v2_product(
            "10001", "测试代金券", "100元代金券：售价72元，最多使用2张",
        )
        self.store.import_store_text(
            "【广东省】\n【东莞】常平天虹店", "东莞门店", ["10001"],
        )
        result = self.store.resolve_deterministic("10001", "东莞常平天虹店多少？")
        self.assertEqual("multi_intent", result["kind"])
        self.assertEqual(["store", "price"], result["resolved_intents"])
        self.assertEqual("东莞常平天虹店", result["store_query"])
        self.assertIn("常平天虹店", result["reply"])
        self.assertIn("72元", result["reply"])
        self.assertNotIn("没有“东莞常平天虹店”这一规格", result["reply"])

    def test_store_price_question_lists_only_skus_supported_by_that_store(self):
        self.store.save_v2_product(
            "10001", "多规格代金券",
            "100元代金券：售价64元\n300元代金券：售价192元",
        )
        first = self.store.import_store_text(
            "【广东省】\n【东莞】东莞常平天虹店", "100元门店", [],
        )
        second = self.store.import_store_text(
            "【广东省】\n【东莞】东莞万象汇店", "300元门店", [],
        )
        for sku in self.store.get_v2_product("10001")["skus"]:
            list_id = first["id"] if sku["face_value"] == "100" else second["id"]
            self.store.set_sku_store_rule(
                "10001", sku["sku_key"], sku["sku_name"], "custom", [list_id],
            )

        result = self.store.resolve_deterministic("10001", "东莞常平天虹店多少？")
        self.assertEqual("multi_intent", result["kind"])
        self.assertIn("100元代金券（售价64元）", result["reply"])
        self.assertNotIn("300元代金券", result["reply"])

    def test_image_only_message_reuses_verified_store_context(self):
        context = {
            "query": "深圳总店", "status": "available",
            "matches": [{
                "province": "广东省", "city": "深圳", "branch": "深圳总店",
                "address": "深圳市测试路1号", "match_quality": "exact",
            }],
        }
        result = self.store.resolve_deterministic(
            "10001", "[图片]", store_context=context,
        )
        self.assertEqual("stores", result["kind"])
        self.assertIn("深圳总店", result["reply"])
        self.assertNotIn("不支持语音图片识别", result["reply"])

    def test_redemption_requests_obey_inventory_day_and_stack_limits(self):
        for stock, limit, time in ((0, 3, "周末"), (1, 3, "周末"), (3, 1, "周末"), (3, 3, "工作日")):
            with self.subTest(stock=stock, limit=limit, time=time):
                self.store.save_v2_product("10001", "代金券", "")
                self.store.save_ai_summary("10001", "", {
                    "skus": [{"name": "100元代金券", "face_value": "100",
                              "sale_price": "79.5", "stock": stock, "max_stack": limit,
                              "applicable_time": time}],
                })
                with patch.object(V2Store, "_current_day_type", return_value="weekend"):
                    result = self.store.resolve_deterministic("10001", "200")
                self.assertEqual("deny", result["decision"])
                self.assertNotIn("可以购买", result["reply"])

    def test_buyer_store_statement_is_locally_verified_not_model_fallback(self):
        self.store.import_store_text(
            "【上海市】\n【上海】上海世博园店\n【上海】龙华会店",
            "上海门店", ["10001"],
        )
        result = self.store.resolve_deterministic("10001", "龙华会店也可以")
        self.assertEqual("stores", result["kind"])
        self.assertIn("龙华会店", result["reply"])
        self.assertNotIn("全国", result["reply"])
        self.assertNotIn("上海除外", result["reply"])

    def test_store_import_preserves_province_city_branch_and_address(self):
        path = os.path.join(self.temp.name, "province-city-stores.xlsx")
        book = Workbook()
        sheet = book.active
        sheet.append(["序号", "省份", "城市", "门店名称", "分店名", "地址", "显示营业时间"])
        sheet.append([1, "广东省", "广州市", "山缓缓火锅", "广州番禺天河城店",
                      "番禺区南村镇天河城2层", "10:00-22:00"])
        book.save(path)
        self.store.import_store_list(path, "山缓缓门店", ["10001"])

        result = self.store.search_store("10001", "广州番禺天河城店")
        self.assertEqual("available", result["status"])
        match = result["matches"][0]
        self.assertEqual("广东省", match["province"])
        self.assertEqual("广州市", match["city"])
        self.assertEqual("广州番禺天河城店", match["branch"])
        self.assertEqual("番禺区南村镇天河城2层", match["address"])

    def test_configured_store_alias_resolves_only_to_existing_current_branch(self):
        self.store.import_store_text(
            "【广东省】\n【深圳】龙岗大运天地店", "深圳门店", ["10001"],
        )
        result = self.store.resolve_deterministic("10001", "深圳大运中心可以使用吗")
        self.assertEqual("stores", result["kind"])
        self.assertIn("龙岗大运天地店", result["reply"])
        self.assertIn("根据“深圳大运中心”", result["reply"])

        other = V2Store(os.path.join(self.temp.name, "other.db"))
        other.save_v2_product("20002", "其他代金券", "100元代金券售价80元")
        other.import_store_text("【广东省】\n【深圳】深圳湾店", "其他门店", ["20002"])
        missing = other.search_store("20002", "深圳大运中心")
        self.assertEqual("area_fallback", missing["status"])
        self.assertEqual(["深圳湾店"], [row["branch"] for row in missing["matches"]])

    def test_wuhan_kaide_wusheng_query_does_not_match_short_address_fragment(self):
        path = os.path.join(self.temp.name, "wuhan-stores.xlsx")
        book = Workbook()
        sheet = book.active
        sheet.append(["店名", "分店名", "省", "市", "地址"])
        sheet.append(["鱼酷", "武胜凯德店", "湖北", "武汉", "中山大道238号凯德广场"])
        sheet.append(["鱼酷", "白沙龙湖天街店", "湖北", "武汉", "张家湾烽胜路龙湖天街"])
        book.save(path)
        self.store.import_store_list(path, "鱼酷门店", ["10001"])

        result = self.store.resolve_deterministic("10001", "武汉凯德广场武胜路能用吗")
        self.assertEqual("stores", result["kind"])
        self.assertEqual(["武胜凯德店"], [row["branch"] for row in result["store_matches"]])
        self.assertIn("武胜凯德店", result["reply"])
        self.assertNotIn("白沙龙湖天街店", result["reply"])

    def test_purchase_flow_is_answered_locally(self):
        result = self.store.resolve_deterministic("10001", "怎么下单")
        self.assertEqual("purchase_flow", result["kind"])
        self.assertIn("当前商品页面", result["reply"])
        self.assertIn("完成付款", result["reply"])

    def test_300_coupon_quantity_uses_300_total_cap_not_100_denomination(self):
        self.store.save_v2_product(
            "10001", "同仁四季代金券",
            "100元代金券：售价111.8元，发100元券2张；"
            "300元代金券：售价204.8元，发300元券1张。"
            "单次或每桌限用数量：100元券最多使用2张；300元券最多使用1张；"
            "单次最多抵扣300元。",
        )
        result = self.store.resolve_deterministic("10001", "300可以用几张")
        self.assertIn("每次最多使用1张300元券", result["reply"])
        self.assertIn("共可抵扣300元", result["reply"])
        self.assertNotIn("共可抵扣100元", result["reply"])

    def test_city_unknown_district_and_reordered_mall_name_match_bound_store(self):
        self.store.import_store_text(
            "【山东省】\n【青岛】青岛未来城万科广场店",
            "鱼酷青岛门店", ["10001"],
        )
        result = self.store.resolve_deterministic(
            "10001", "青岛市市北区万科未来城可以用吗",
        )
        self.assertEqual("stores", result["kind"])
        self.assertEqual("available", result["store_status"])
        self.assertEqual(
            ["青岛未来城万科广场店"],
            [row["branch"] for row in result["store_matches"]],
        )
        self.assertIn("可以用，根据“青岛市市北区万科未来城”", result["reply"])

    def test_county_level_city_alias_matches_parent_city_store_branch(self):
        self.store.import_store_text(
            "【浙江省】\n【温州】南虹广场店、五马街店",
            "温州门店", ["10001"],
        )
        result = self.store.resolve_deterministic("10001", "乐清南虹")
        self.assertEqual("stores", result["kind"])
        self.assertEqual("available", result["store_status"])
        self.assertEqual(["南虹广场店"], [
            row["branch"] for row in result["store_matches"]
        ])
        self.assertIn("南虹广场店", result["reply"])

    def test_store_and_today_are_both_answered_for_bare_exact_branch(self):
        self.store.import_store_text(
            "【北京市】\n【北京】西直门店", "北京门店", ["10001"],
        )
        result = self.store.resolve_deterministic("10001", "西直门今天能用么？")
        self.assertEqual("multi_intent", result["kind"])
        self.assertEqual(["store", "date"], result["resolved_intents"])
        self.assertIn("可以用，根据“西直门”查询到可用门店：西直门店。", result["reply"])
        self.assertIn("可以使用", result["reply"])

    def test_store_and_now_are_both_answered_without_precise_time_wording(self):
        self.store.import_store_text(
            "【北京市】\n【北京】工体店", "北京门店", ["10001"],
        )
        result = self.store.resolve_deterministic("10001", "工体店现在可以用吗")
        self.assertEqual("multi_intent", result["kind"])
        self.assertEqual(["store", "current_time"], result["resolved_intents"])
        self.assertIn("工体店", result["reply"])
        self.assertIn("门店营业时段可用就可以了。", result["reply"])
        self.assertNotIn("精确营业时段", result["reply"])

    def test_store_and_current_time_are_both_answered_for_exact_branch(self):
        self.store.import_store_text(
            "【北京市】\n【北京】工体店", "北京门店", ["10001"],
        )

        result = self.store.resolve_deterministic("10001", "工体店现在可以用吗")

        self.assertEqual("multi_intent", result["kind"])
        self.assertEqual(["store", "current_time"], result["resolved_intents"])
        self.assertIn("工体店", result["reply"])
        self.assertIn("门店营业时段可用就可以了。", result["reply"])

    def test_repeated_ability_word_does_not_break_store_and_date_query(self):
        self.store.import_store_text(
            "【北京市】\n【北京】西直门店", "北京门店", ["10001"],
        )
        result = self.store.resolve_deterministic("10001", "西直门今天能能用么？")
        self.assertEqual("multi_intent", result["kind"])
        self.assertEqual(["store", "date"], result["resolved_intents"])
        self.assertIn("西直门店", result["reply"])

    def test_negative_followup_rechecks_available_store_and_confirms_naturally(self):
        self.store.import_store_text(
            "【山东省】\n【青岛】青岛未来城万科广场店",
            "鱼酷青岛门店", ["10001"],
        )
        first = self.store.resolve_deterministic(
            "10001", "青岛市市北区万科未来城可以用吗",
        )
        context = {
            "query": first["store_query"], "matches": first["store_matches"],
            "status": first["store_status"], "verified": True,
        }
        followup = self.store.resolve_deterministic(
            "10001", "不能用吗", store_context=context,
        )
        self.assertEqual("stores", followup["kind"])
        self.assertEqual("available", followup["store_status"])
        self.assertEqual(
            "可以用，青岛未来城万科广场店在当前商品的可用门店范围内。",
            followup["reply"],
        )

    def test_numeric_prefix_branch_is_not_reclassified_as_a_price(self):
        self.store.save_v2_product(
            "10001", "聚宝源100元代金券 老字号涮肉", "100元代金券：售价77元。",
        )
        self.store.import_store_text(
            "【北京市】\n【北京】聚宝源798店", "聚宝源北京门店", ["10001"],
        )
        result = self.store.resolve_deterministic("10001", "798店聚宝源能用吗")
        self.assertEqual("stores", result["kind"])
        self.assertEqual("available", result["store_status"])
        self.assertIn("聚宝源798店", result["reply"])

    def test_number_inside_branch_name_is_not_reclassified_as_sku(self):
        self.store.save_v2_product(
            "10001", "冯校长老火锅代金券", "100元代金券：售价75元。",
        )
        self.store.import_store_text(
            "【广东省】\n【东莞】东城33小镇店", "冯校长东莞门店", ["10001"],
        )
        result = self.store.resolve_deterministic(
            "10001", "东莞东城33小镇能用吗",
        )
        self.assertEqual("stores", result["kind"])
        self.assertEqual("available", result["store_status"])
        self.assertIn("东城33小镇店", result["reply"])
        self.assertNotIn("33元代金券", result["reply"])

    def test_product_category_is_removed_from_store_search_key(self):
        self.store.save_v2_product(
            "10001", "聚宝源100元代金券 老字号涮肉", "100元代金券：售价77元。",
        )
        self.store.import_store_text(
            "【河北省】\n【石家庄】聚宝源东尚店", "聚宝源河北门店", ["10001"],
        )
        result = self.store.resolve_deterministic(
            "10001", "石家庄聚宝源涮肉东尚店",
        )
        self.assertEqual("stores", result["kind"])
        self.assertIn("聚宝源东尚店", result["reply"])

    def test_purchase_followup_keeps_unavailable_store_context(self):
        context = {
            "query": "聚宝源东尚店", "matches": [],
            "status": "unavailable", "verified": True,
        }
        result = self.store.resolve_deterministic(
            "10001", "可以买吗", store_context=context,
        )
        self.assertEqual("stores", result["kind"])
        self.assertEqual("deny", result["decision"])
        self.assertIn("尚未确认“聚宝源东尚店”属于可用门店", result["reply"])
        self.assertNotIn("当天购当天用", result["reply"])

    def test_colloquial_now_purchase_question_uses_business_hours_reply(self):
        result = self.store.resolve_deterministic("10001", "现在团可以吗")
        self.assertEqual("time", result["kind"])
        self.assertEqual("门店营业时段可用就可以了。", result["reply"])
        self.assertNotIn("不参与团购", result["reply"])

    def test_explicit_date_price_never_falls_back_to_today_or_identity_fare(self):
        self.store.save_v2_product(
            "10001", "作作烧肉单人满贯全天自助",
            """单人满贯全天自助：208元（周末节假日通用）
双人满贯全天自助：416元（周末节假日通用）
儿童及老人请到店自购：168元""",
        )
        self.store.import_store_text(
            "【北京市】\n【北京】青年路店", "作作北京门店", ["10001"],
        )
        result = self.store.resolve_deterministic(
            "10001", "北京青年路店，9.24晚上多少钱？",
        )
        self.assertEqual("multi_intent", result["kind"])
        self.assertIn("青年路店", result["reply"])
        self.assertIn("9月24日晚餐", result["reply"])
        self.assertNotIn("今天", result["reply"])
        self.assertNotIn("儿童及老人", result["reply"])

    def test_m89_two_people_uses_two_high_tier_single_skus_not_light_double(self):
        self.store.save_v2_product(
            "10001", "一绪寿喜烧高阶和牛自助",
            """畅吃清蒸梭子蟹_M8-9和牛单人自助：235元
畅吃清蒸梭子蟹_轻享和牛双人自助：438元""",
        )
        first = self.store.resolve_deterministic("10001", "M8 M9梭子蟹 双人")
        self.assertEqual("price", first["kind"])
        self.assertIn("M8-9和牛单人自助", first["reply"])
        self.assertIn("共470元", first["reply"])
        self.assertNotIn("轻享和牛", first["reply"])

        followup = self.store.resolve_deterministic(
            "10001", "高阶和牛", store_context=first["query_context_update"],
        )
        self.assertIn("M8-9和牛单人自助", followup["reply"])
        self.assertIn("共470元", followup["reply"])
        self.assertNotIn("轻享和牛", followup["reply"])

    def test_store_date_meal_and_ticket_quantity_are_resolved_independently(self):
        self.store.save_v2_product(
            "10001", "作作烧肉单人满贯全天自助",
            """作作全天单人自助：208元（周末节假日通用）
作作全天双人自助：416元（周末节假日通用）
作作全天三人自助：624元（周末节假日通用）""",
        )
        self.store.import_store_text(
            "【北京市】\n【北京】双井店\n【北京】中关村大融城店", "作作北京门店", ["10001"],
        )
        result = self.store.resolve_deterministic(
            "10001", "今天晚餐两张，双井店还有吗",
        )
        self.assertEqual("multi_intent", result["kind"])
        self.assertIn("双井店", result["reply"])
        self.assertIn("购买2张", result["reply"])
        self.assertIn("共416元", result["reply"])
        self.assertNotIn("双井店还", result["reply"])
        self.assertNotIn("这一规格", result["reply"])

    def test_store_child_does_not_reparse_selected_sku_into_today_conditions(self):
        self.store.save_v2_product(
            "10001", "作作烧肉单人满贯全天自助",
            """作作全天单人自助：208元（周末节假日通用）
作作全天双人自助：416元（周末节假日通用）
作作全天三人自助：624元（周末节假日通用）""",
        )
        self.store.import_store_text(
            "【北京市】\n【北京】双井店\n【北京】中关村大融城店", "作作北京门店", ["10001"],
        )
        result = self.store.resolve_deterministic(
            "10001", "您好本周末中关村店两个人能用吗",
        )
        self.assertEqual("multi_intent", result["kind"])
        self.assertIn("中关村", result["reply"])
        self.assertIn("周末可用", result["reply"])
        self.assertIn("416元", result["reply"])
        self.assertNotIn("工作日", result["reply"])
        self.assertNotIn("没有符合", result["reply"])

    def test_missing_date_or_meal_restriction_means_unrestricted(self):
        self.store.save_v2_product(
            "10001", "全天单人自助", "全天单人自助：88元。",
        )
        for message in ("工作日晚餐1人能用吗", "周末午餐1人能用吗"):
            with self.subTest(message=message):
                result = self.store.resolve_deterministic("10001", message)
                self.assertEqual("allow", result["decision"])
                self.assertIn("可以", result["reply"])
                self.assertNotIn("没有符合", result["reply"])

    def test_explicit_workday_label_remains_workday_only(self):
        self.store.save_v2_product(
            "10001", "工作日单人自助", "工作日单人自助：88元。",
        )
        allowed = self.store.resolve_deterministic("10001", "工作日1人能用吗")
        blocked = self.store.resolve_deterministic("10001", "周末1人能用吗")
        self.assertIn("可以", allowed["reply"])
        self.assertRegex(blocked["reply"], r"不可以|不能使用|没有符合")

    def test_direct_purchase_with_named_store_does_not_invent_today_condition(self):
        self.store.save_v2_product(
            "10001", "作作烧肉单人满贯全天自助",
            "单人满贯全天自助：208元（仅限工作日）\n双人满贯全天自助：416元（仅限工作日）",
        )
        self.store.import_store_text(
            "【北京市】\n【北京】西单大悦城店", "作作北京门店", ["10001"],
        )
        result = self.store.resolve_deterministic(
            "10001", "您好 可以直拍吧 西单大悦城店 单人满贯全天 😍",
        )
        self.assertEqual("multi_intent", result["kind"])
        self.assertIn("可以直接拍下", result["reply"])
        self.assertIn("西单大悦城店", result["reply"])
        self.assertNotIn("没有符合", result["reply"])
        self.assertNotIn("今天是", result["reply"])

    def test_now_purchase_and_tonight_use_without_store_checks_only_time(self):
        self.store.save_v2_product(
            "10001", "全天单人自助", "全天单人自助：88元。",
        )
        result = self.store.resolve_deterministic("10001", "我现在拍 今晚能用嘛？")
        self.assertEqual("multi_intent", result["kind"])
        self.assertIn("可以现在拍下", result["reply"])
        self.assertIn("门店营业时段可用就可以了", result["reply"])
        self.assertNotIn("补充所在城市", result["reply"])
        self.assertNotIn("未查询到可用门店", result["reply"])
        self.assertNotIn("store", result["resolved_intents"])

    def test_bare_can_use_question_never_becomes_store_query(self):
        self.store.save_v2_product(
            "10001", "全天单人自助", "全天单人自助：88元。",
        )
        for message in ("能用吗", "可以用吗", "能用嘛"):
            with self.subTest(message=message):
                result = self.store.resolve_deterministic("10001", message)
                self.assertFalse(result and result.get("kind") == "stores")
                self.assertNotIn("补充所在城市", str((result or {}).get("reply") or ""))

    def test_tonight_obeys_explicit_weekday_and_weekend_labels(self):
        self.store.save_v2_product(
            "10001", "日期限制自助",
            "工作日单人自助：88元（仅限工作日）\n周末单人自助：99元（仅限周末）",
        )
        result = self.store.resolve_deterministic("10001", "今晚单人能用吗")
        self.assertNotEqual("stores", result["kind"])
        self.assertIn("周末单人自助", result["reply"])
        self.assertNotIn("工作日单人自助", result["reply"])

    def test_buffet_purchase_plan_by_meal_and_audience_totals(self):
        self.store.save_v2_product(
            "10001", "分时段自助餐",
            """工作日午餐成人票：售价88元
工作日午餐老人票：售价68元，限60周岁以上
工作日午餐儿童票：售价48元，限身高1.2米至1.4米
工作日晚餐成人票：售价108元
工作日晚餐老人票：售价78元，限60周岁以上
工作日晚餐儿童票：售价58元，限身高1.2米至1.4米""",
        )
        result = self.store.resolve_deterministic(
            "10001", "工作日午餐2位大人1位老人1位小孩怎么买",
        )
        self.assertEqual("audience_price", result["kind"])
        self.assertIn("工作日午餐购买建议", result["reply"])
        self.assertIn("2位成人需要购买2张工作日午餐成人票", result["reply"])
        self.assertIn("共176元", result["reply"])
        self.assertIn("1位老人需要购买1张工作日午餐老人票，共68元", result["reply"])
        self.assertIn("1位儿童需要购买1张工作日午餐儿童票，共48元", result["reply"])
        self.assertIn("合计292元", result["reply"])
        self.assertNotIn("工作日晚餐成人票", result["reply"])

    def test_buffet_missing_meal_keeps_audience_counts_for_followup(self):
        self.store.save_v2_product(
            "10001", "分时段自助餐",
            """工作日午餐成人票：售价88元
工作日午餐老人票：售价68元
工作日午餐儿童票：售价48元
工作日晚餐成人票：售价108元
工作日晚餐老人票：售价78元
工作日晚餐儿童票：售价58元""",
        )
        first = self.store.resolve_deterministic(
            "10001", "工作日2位成人1位老人1位儿童怎么买",
        )
        self.assertEqual("audience_price", first["kind"])
        self.assertIn("午餐、晚餐", first["reply"])
        self.assertEqual("meal_period", first["query_context_update"]["price_filters"]["awaiting"])

        followup = self.store.resolve_deterministic(
            "10001", "晚餐", store_context=first["query_context_update"],
        )
        self.assertIn("工作日晚餐购买建议", followup["reply"])
        self.assertIn("共216元", followup["reply"])
        self.assertIn("老人票，共78元", followup["reply"])
        self.assertIn("儿童票，共58元", followup["reply"])
        self.assertIn("合计352元", followup["reply"])

    def test_buffet_bare_people_count_defaults_to_adults(self):
        self.store.save_v2_product(
            "10001", "身份票自助餐",
            """工作日午餐成人票：售价88元
工作日午餐老人票：售价68元
工作日午餐儿童票：售价48元""",
        )
        result = self.store.resolve_deterministic(
            "10001", "工作日午餐3人自助怎么买",
        )
        self.assertIn("3位成人需要购买3张工作日午餐成人票", result["reply"])
        self.assertIn("共264元", result["reply"])
        self.assertNotIn("分别有几位", result["reply"])

    def test_unlabelled_buffet_sku_is_an_adult_fare(self):
        self.store.save_v2_product(
            "10001", "一绪寿喜烧自助",
            "轻享和牛单人自助：售价219元\n"
            "轻享和牛双人自助：售价438元\n"
            "儿童票：售价68元",
        )
        result = self.store.resolve_deterministic("10001", "怎么拍？轻享两位")
        self.assertIn("轻享和牛双人自助", result["reply"])
        self.assertIn("438元", result["reply"])
        self.assertNotIn("没有“成人票”", result["reply"])
        self.assertNotIn("儿童票", result["reply"])

    def test_buffet_tier_is_a_hard_filter_for_the_same_party_size(self):
        self.store.save_v2_product(
            "10001", "多档自助餐",
            "经典双人自助：售价398元\n"
            "轻享双人自助：售价438元\n"
            "尊享双人自助：售价528元",
        )
        light = self.store.resolve_deterministic("10001", "轻享两位怎么拍")
        self.assertIn("轻享双人自助", light["reply"])
        self.assertIn("438元", light["reply"])
        self.assertNotIn("经典双人", light["reply"])
        self.assertNotIn("尊享双人", light["reply"])

        premium = self.store.resolve_deterministic("10001", "尊享2人多少钱")
        self.assertIn("尊享双人自助", premium["reply"])
        self.assertIn("528元", premium["reply"])
        self.assertNotIn("398元", premium["reply"])

    def test_unlabelled_buffet_fare_is_all_day_but_special_fares_are_scoped(self):
        self.store.save_v2_product(
            "10001", "分时段自助餐",
            "经典单人自助：售价99元\n"
            "午餐儿童票：售价49元\n"
            "晚餐儿童票：售价59元",
        )
        adult = self.store.resolve_deterministic("10001", "周末晚餐2位怎么拍")
        self.assertIn("需要购买2份经典单人自助", adult["reply"])
        self.assertIn("共198元", adult["reply"])
        self.assertNotIn("儿童票", adult["reply"])

        child = self.store.resolve_deterministic("10001", "晚餐1位儿童多少钱")
        self.assertIn("晚餐儿童票", child["reply"])
        self.assertIn("59元", child["reply"])
        self.assertNotIn("午餐儿童票", child["reply"])

    def test_specific_meal_price_overrides_unrestricted_same_tier(self):
        self.store.save_v2_product(
            "10001", "轻享自助餐",
            "轻享双人自助：售价438元\n"
            "晚餐轻享双人自助：售价458元",
        )
        dinner = self.store.resolve_deterministic("10001", "晚餐轻享2位怎么拍")
        self.assertIn("晚餐轻享双人自助", dinner["reply"])
        self.assertIn("458元", dinner["reply"])

        lunch = self.store.resolve_deterministic("10001", "午餐轻享2位怎么拍")
        self.assertIn("轻享双人自助", lunch["reply"])
        self.assertIn("438元", lunch["reply"])
        self.assertNotIn("458元", lunch["reply"])

    def test_unlabelled_adult_fare_never_answers_a_child_query(self):
        self.store.save_v2_product(
            "10001", "经典自助餐", "经典单人自助：售价99元",
        )
        result = self.store.resolve_deterministic("10001", "1位儿童多少钱")
        self.assertEqual("audience_price", result["kind"])
        self.assertIn("没有“儿童票”", result["reply"])
        self.assertNotIn("需要购买1张经典单人自助", result["reply"])

    def test_family_package_does_not_drop_extra_senior(self):
        self.store.save_v2_product(
            "10001", "家庭自助套餐",
            "午餐两大一小家庭套餐：售价218元\n午餐老人票：售价68元",
        )
        result = self.store.resolve_deterministic("10001", "午餐2大1老1小怎么买")
        self.assertIn("购买1份午餐两大一小家庭套餐，共218元", result["reply"])
        self.assertIn("1位老人需要购买1张午餐老人票，共68元", result["reply"])
        self.assertIn("合计286元", result["reply"])

    def test_store_reply_never_adds_stack_limit_outside_sku_name(self):
        self.store.save_v2_product("10001", "平台代金券", "")
        self.store.save_ai_summary("10001", "平台代金券规格", {
            "products": [
                {"name": "美团100", "option_type": "代金券", "face_value": "100", "sale_price": "69", "max_stack": "2"},
                {"name": "抖音300", "option_type": "代金券", "face_value": "300", "sale_price": "206", "max_stack": "1"},
            ],
        })
        first = self.store.import_store_text(
            "【山东省】\n【青岛】青岛凯德MALL·新都心店", "美团100门店", [],
        )
        second = self.store.import_store_text(
            "【山东省】\n【青岛】青岛万象城店", "抖音300门店", [],
        )
        for sku in self.store.get_v2_product("10001")["skus"]:
            list_id = first["id"] if sku["face_value"] == "100" else second["id"]
            self.store.set_sku_store_rule(
                "10001", sku["sku_key"], sku["sku_name"], "custom", [list_id],
            )
        result = self.store.resolve_deterministic(
            "10001", "青岛凯德MALL新都心店多少？",
        )
        self.assertEqual("multi_intent", result["kind"])
        self.assertIn("美团100（售价69元）", result["reply"])
        self.assertNotIn("美团100（可叠加2张）", result["reply"])
        self.assertNotIn("抖音300", result["reply"])

    def test_store_configuration_reads_synced_platform_skus_without_ai_summary(self):
        platform_summary = json.dumps({
            "title": "COMMUNE幻师自助餐",
            "description": "",
            "price": "88",
            "sku": [
                {
                    "skuId": "lunch-adult", "priceInCent": 8800, "quantity": 12,
                    "propertyList": [{"actualValueText": "工作日午餐成人票"}],
                },
                {
                    "skuId": "dinner-child", "priceInCent": 5800, "quantity": 3,
                    "propertyList": [{"actualValueText": "工作日晚餐儿童票"}],
                },
            ],
        }, ensure_ascii=False)
        self.store.upsert_synced_product({
            "item_id": "direct-platform-skus", "title": "COMMUNE幻师自助餐",
            "platform_summary": platform_summary, "image_urls": [], "price": "88",
        })

        product = self.store.get_v2_product("direct-platform-skus")
        self.assertEqual("", product["ai_summary"])
        self.assertEqual(
            ["工作日午餐成人票", "工作日晚餐儿童票"],
            [sku["sku_name"] for sku in product["skus"]],
        )
        self.assertEqual(["88", "58"], [sku["sale_price"] for sku in product["skus"]])
        self.assertEqual(
            ["source:lunch-adult", "source:dinner-child"],
            [sku["sku_key"] for sku in product["skus"]],
        )

    def test_business_hour_day_words_do_not_restrict_unlabelled_skus(self):
        schedule = (
            "门店营业时间内可用；晚市最后加餐21:30，23:00闭店；"
            "周末节假日最后加餐22:00，23:30闭店"
        )
        self.store.save_v2_product(
            "10001", "作作烧肉满贯全天自助", schedule,
        )
        self.store.save_ai_summary("10001", "营业时间", {
            "packages": [
                {"name": "单人满贯全天自助", "sale_price": "210.8", "applicable_time": schedule},
                {"name": "双人满贯全天自助", "sale_price": "421.6", "applicable_time": schedule},
            ],
            "facts": {"使用时间": schedule}, "time_rules": [],
        })
        single = self.store.resolve_deterministic("10001", "今天一个人")
        double = self.store.resolve_deterministic("10001", "今天两个人")
        tonight = self.store.resolve_deterministic("10001", "今天晚上双人多少钱")
        self.assertIn("单人满贯全天自助", single["reply"])
        self.assertIn("210.8元", single["reply"])
        self.assertIn("双人满贯全天自助", double["reply"])
        self.assertIn("421.6元", double["reply"])
        self.assertIn("双人满贯全天自助", tonight["reply"])
        self.assertNotIn("没有符合", single["reply"] + double["reply"] + tonight["reply"])

    def test_conditional_price_still_obeys_explicit_blackout_date(self):
        self.store.save_v2_product(
            "10001", "全天单人自助",
            "全天单人自助：88元。2026年9月6日不可用。",
        )
        result = self.store.resolve_deterministic("10001", "今天一个人多少钱")
        self.assertEqual("deny", result["decision"])
        self.assertIn("不可用日期范围", result["reply"])

    def test_amount_question_matrix_covers_prefix_suffix_and_chinese_forms(self):
        self.store.save_v2_product(
            "amount-matrix", "代金券",
            "100元代金券：售价57.9元，最多叠加2张\n"
            "300元代金券：售价212.8元，每次限用1张",
        )
        for message in (
            "拍200的", "买200的", "要200的", "来200的", "给我200的",
            "下单200的", "200拍几张", "两百怎么拍",
        ):
            with self.subTest(message=message):
                result = self.store.resolve_deterministic("amount-matrix", message)
                self.assertIsNotNone(result)
                self.assertIn("2张100元代金券", result["reply"])
                self.assertIn("抵扣200元", result["reply"])

    def test_paid_price_and_face_relationship_matrix_uses_full_sku_name(self):
        self.store.save_v2_product(
            "relation-matrix", "平台代金券",
            "美团200（可叠加2张）：售价108元",
        )
        for message in (
            "108是200的券吗", "108可以买200吗", "200的券是108吗",
            "200卖108吗", "108买200", "108得200", "108发200",
            "108换200", "108/200", "108→200", "108能抵200吗",
        ):
            with self.subTest(message=message):
                result = self.store.resolve_deterministic("relation-matrix", message)
                self.assertIsNotNone(result)
                self.assertEqual("voucher_value", result["kind"])
                self.assertIn("美团200元代金券（可叠加2张）售价108元", result["reply"])

    def test_near_price_relationship_corrects_actual_sku_price(self):
        self.store.save_v2_product(
            "near-price", "平台代金券", "美团200元代金券：售价115.8元",
        )
        result = self.store.resolve_deterministic("near-price", "115代200")
        self.assertEqual("voucher_value", result["kind"])
        self.assertIn("不是", result["reply"])
        self.assertIn("115.8元", result["reply"])

    def test_platform_bundle_corrects_listing_price_relationship_question(self):
        platform_summary = json.dumps({
            "title": "同仁四季椰子鸡代金券",
            "description": "100元代金券最多使用2张",
            "price": "108",
            "sku": [
                {
                    "skuId": "dish", "priceInCent": 10800, "quantity": 10,
                    "propertyList": [{"actualValueText": "单品煲仔饭"}],
                },
                {
                    "skuId": "voucher-200", "priceInCent": 11580, "quantity": 10,
                    "propertyList": [{"actualValueText": "100x2"}],
                },
            ],
        }, ensure_ascii=False)
        self.store.upsert_synced_product({
            "item_id": "platform-bundle-price", "title": "同仁四季椰子鸡代金券",
            "platform_summary": platform_summary, "image_urls": [], "price": "108",
        })

        wrong = self.store.resolve_deterministic(
            "platform-bundle-price", "108抵200是这样吗",
        )
        self.assertEqual("voucher_value", wrong["kind"])
        self.assertIn("不是", wrong["reply"])
        self.assertIn("115.8元", wrong["reply"])
        self.assertIn("2张100元代金券", wrong["reply"])

        correct = self.store.resolve_deterministic(
            "platform-bundle-price", "115.8抵200是这样吗",
        )
        self.assertEqual("voucher_value", correct["kind"])
        self.assertIn("是的", correct["reply"])
        self.assertIn("2张100元代金券", correct["reply"])

    def test_bare_positive_stack_rules_mean_unlimited_and_cover_quantity_phrasings(self):
        self.store.save_v2_product(
            "unlimited", "100元代金券",
            "100元代金券：售价66元\n仅支持同面额代金券叠加\n可叠加",
        )
        product = self.store.get_v2_product("unlimited")
        for message in ("一次几张", "最多叠几张", "限几张"):
            with self.subTest(message=message):
                result = self.store.resolve_deterministic("unlimited", message)
                self.assertEqual("stacking", result["kind"])
                self.assertIn("不限制使用张数", result["reply"])
        self.assertIn("不限制使用张数", self.store.build_knowledge_summary(product))

    def test_explicit_stack_limit_answers_requested_quantity_phrasings(self):
        self.store.save_v2_product(
            "limited", "100元代金券", "100元代金券：售价69元，可叠加2张",
        )
        over = self.store.resolve_deterministic("limited", "可以用3张吗")
        same = self.store.resolve_deterministic("limited", "两张同时核销吗")
        self.assertIn("最多使用2张", over["reply"])
        self.assertIn("可以", same["reply"])

    def test_first_reply_does_not_invent_same_day_or_expiry_policy(self):
        self.store.save_v2_product(
            "first-no-invention", "100元代金券", "100元代金券：售价69元",
        )
        reply = self.store.build_first_reply_text(
            self.store.get_v2_product("first-no-invention")
        )
        self.assertNotIn("当天购买", reply)
        self.assertNotIn("过期不退不补", reply)

    def test_buffet_identity_first_counts_and_fare_table(self):
        self.store.save_v2_product(
            "buffet-natural", "自助餐",
            "工作日午餐成人票：售价88元\n"
            "工作日午餐老人票：售价68元\n"
            "工作日午餐儿童票：售价48元",
        )
        plan = self.store.resolve_deterministic(
            "buffet-natural", "工作日午餐成人2个老人1个孩子1个怎么买",
        )
        self.assertIn("合计292元", plan["reply"])
        xplan = self.store.resolve_deterministic(
            "buffet-natural", "工作日午餐成人x2老人x1儿童x1",
        )
        self.assertIn("合计292元", xplan["reply"])
        fares = self.store.resolve_deterministic(
            "buffet-natural", "工作日午餐老人和小孩怎么收费",
        )
        self.assertIn("老人票", fares["reply"])
        self.assertIn("儿童票", fares["reply"])

    def test_package_content_query_keeps_sku_and_item_context_without_inventing_exclusion(self):
        self.store.save_v2_product(
            "package-content", "一绪放题",
            "轻享放题：售价198元\n超值放题：售价169.9元\n尊享放题：售价218元",
        )
        self.store.save_ai_summary("package-content", "分档套餐内容", {
            "sku_profiles": [
                {"name": "轻享放题", "sale_price": "198", "package_contents": "基础菜品"},
                {"name": "超值放题", "sale_price": "169.9", "package_contents": "刺身、烧物"},
                {"name": "尊享放题", "sale_price": "218", "package_contents": "刺身、大闸蟹"},
            ],
        })

        first = self.store.resolve_deterministic("package-content", "有大闸蟹的是哪个？")
        self.assertEqual("package_content", first["kind"])
        self.assertIn("尊享放题包含大闸蟹", first["reply"])
        self.assertIn("暂未明确说明", first["reply"])
        self.assertNotIn("超值放题不包含大闸蟹", first["reply"])

        second = self.store.resolve_deterministic(
            "package-content", "超值没有吗？",
            store_context=first["query_context_update"],
        )
        self.assertEqual("package_content", second["kind"])
        self.assertIn("没有明确说明超值放题是否包含大闸蟹", second["reply"])
        self.assertNotIn("超值放题不包含大闸蟹", second["reply"])

        correction = self.store.resolve_deterministic(
            "package-content", "我看好像是有的哦",
            store_context=second["query_context_update"],
        )
        self.assertEqual("package_content", correction["kind"])
        self.assertIn("没有明确写超值放题不含大闸蟹", correction["reply"])
        self.assertIn("不能仅凭未提及", correction["reply"])

    def test_package_content_only_denies_when_sku_has_explicit_negative_evidence(self):
        product = {
            "structured": {"sku_profiles": [
                {"name": "超值放题", "package_contents": "明确不含大闸蟹"},
                {"name": "尊享放题", "package_contents": "包含大闸蟹"},
            ]},
            "raw_text": "", "ai_summary": "", "platform_summary": "",
        }
        result = self.store.package_content_reply(
            product, "超值没有吗", {"content_item": "大闸蟹"},
        )
        self.assertIn("明确标注超值放题不包含大闸蟹", result["reply"])

    def test_classifier_recognizes_reverse_package_content_query(self):
        profile = classify_buyer_message("请问有大闸蟹的是哪个？")
        self.assertEqual("product_attribute", profile["primary_intent"])
        self.assertEqual("option_by_included_item", profile["primary_subtype"])
        self.assertIn("大闸蟹", profile["entities"]["content_items"])

    def test_knowledge_summary_separates_purchase_combinations_and_rule_sections(self):
        raw_text = (
            "幻师COMMUNE代金券\n"
            "①250元代金券：173.8（最多叠加2张）\n"
            "②300元代金券：211.8（最多叠加2张）\n"
            "③500（250x2）：347.6。\n"
            "④600（300x2）：427.6。\n"
            "2026年9月25日、9月30日及国庆节10月1日至10月7日不可用。\n"
            "仅适用于餐品、酒水。\n"
            "仅限堂食。\n"
            "无需预约，消费高峰期可能需要等位。"
        )
        product = {
            "title": "幻师COMMUNE代金券",
            "raw_text": raw_text,
            "platform_summary": json.dumps({
                "title": "幻师COMMUNE代金券",
                "description": raw_text,
                "sku": [
                    {"skuName": "250元代金券", "priceInCent": 17380},
                    {"skuName": "300元代金券", "priceInCent": 21180},
                ],
            }, ensure_ascii=False),
            "structured": {"facts": {}},
        }

        summary = self.store.build_knowledge_summary(product)

        self.assertIn("【组合购买方案】", summary)
        self.assertIn("500元方案：购买250元代金券×2，共支付347.6元", summary)
        self.assertIn("600元方案：购买300元代金券×2，按当前SKU售价应付423.6元", summary)
        self.assertIn("原文写427.6元，价格冲突，需人工确认", summary)
        self.assertIn("【不可用日期】", summary)
        self.assertIn("【适用范围】", summary)
        self.assertIn("【使用规则】", summary)
        self.assertNotIn("③500", summary)
        self.assertNotIn("④600", summary)


if __name__ == "__main__":
    unittest.main()
